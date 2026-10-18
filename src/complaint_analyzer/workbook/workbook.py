"""Reading the input workbook.

Knows about spreadsheets and nothing else. It does not know what an analysis is,
which fields matter, or how rows are filtered.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterator, Sequence

from openpyxl import load_workbook

HEADER_ROW = 1


class DataError(Exception):
    """Raised when the workbook cannot be read or is missing required columns."""


def cell_text(value: Any) -> str:
    """Render a spreadsheet cell as the text the model will see."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.isoformat(sep=" ")
        return value.date().isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass(frozen=True)
class RawRow:
    """One workbook row, keyed by column header, with its 1-based sheet position."""

    source_row: int
    values: dict[str, str]

    def value(self, column: str) -> str:
        return self.values.get(column, "")


@dataclass(frozen=True)
class Workbook:
    """A sheet that has been read into memory.

    ``columns`` preserves the workbook's own column order so the output workbook can
    reproduce it exactly.
    """

    path: Path
    sheet: str
    columns: tuple[str, ...]
    rows: tuple[RawRow, ...]

    def require_columns(self, names: Sequence[str]) -> None:
        missing = [name for name in names if name not in self.columns]
        if missing:
            raise DataError(
                f"sheet '{self.sheet}' in {self.path} is missing required column(s): "
                + ", ".join(repr(name) for name in missing)
                + ". Available columns: "
                + ", ".join(repr(name) for name in self.columns)
            )

    def __len__(self) -> int:
        return len(self.rows)

    def __iter__(self) -> Iterator[RawRow]:
        return iter(self.rows)


def _read_header(raw_header: Sequence[Any], path: Path, sheet: str) -> list[str]:
    headers: list[str] = []
    seen: set[str] = set()
    for index, value in enumerate(raw_header, start=1):
        name = cell_text(value).strip()
        if not name:
            # Trailing or interior blank header: the column carries no identity, so it
            # is dropped rather than guessed at.
            headers.append("")
            continue
        if name in seen:
            raise DataError(
                f"sheet '{sheet}' in {path} has duplicate column header {name!r} "
                f"at position {index}"
            )
        seen.add(name)
        headers.append(name)
    if not seen:
        raise DataError(f"sheet '{sheet}' in {path} has no column headers in row {HEADER_ROW}")
    return headers


def load_sheet(path: Path, sheet: str) -> Workbook:
    """Read one sheet, using row 1 as the header."""
    if not path.is_file():
        raise DataError(f"input workbook not found: {path}")

    try:
        book = load_workbook(filename=path, read_only=True, data_only=True)
    except Exception as exc:
        raise DataError(f"could not open workbook {path}: {exc}") from exc

    try:
        if sheet not in book.sheetnames:
            raise DataError(
                f"sheet '{sheet}' not found in {path}. "
                f"Available sheets: {', '.join(repr(name) for name in book.sheetnames)}"
            )
        worksheet = book[sheet]

        row_iterator = worksheet.iter_rows(values_only=True)
        try:
            raw_header = next(row_iterator)
        except StopIteration:
            raise DataError(f"sheet '{sheet}' in {path} is empty") from None

        headers = _read_header(raw_header, path, sheet)
        columns = tuple(name for name in headers if name)

        rows: list[RawRow] = []
        for offset, raw in enumerate(row_iterator, start=HEADER_ROW + 1):
            values = {
                name: cell_text(raw[index] if index < len(raw) else None)
                for index, name in enumerate(headers)
                if name
            }
            if not any(values.values()):
                continue
            rows.append(RawRow(source_row=offset, values=values))
    finally:
        book.close()

    return Workbook(path=path, sheet=sheet, columns=columns, rows=tuple(rows))
