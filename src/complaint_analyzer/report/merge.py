"""Merge analysis results into the exported cases workbook.

The output keeps every sheet from the export (``Cases`` plus child-relationship detail
sheets) and appends generated columns to ``Cases``:

* ``<prefix>.<key>`` for each output key of the analysis, in prompt order,
* ``<prefix>.analysis_status`` and ``<prefix>.analysis_error``.

Generated headers and cells get a light-blue fill so they are easy to tell apart from
the Salesforce data, and a ``Legend`` sheet explains the colors and the source run.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from ..run_paths import RESULTS_NAME
from ..salesforce.export import MAIN_SHEET, Table, make_cell, write_table

GENERATED_FILL_COLOR = "DDEBF7"
LEGEND_SHEET = "Legend"
STATUS_KEY = "analysis_status"
ERROR_KEY = "analysis_error"
MISSING_STATUS = "not_run"


class MergeError(Exception):
    """Raised when the inputs to a merge are missing or inconsistent."""


def generated_fill() -> PatternFill:
    return PatternFill(fill_type="solid", start_color=GENERATED_FILL_COLOR, end_color=GENERATED_FILL_COLOR)


def _cell_value(value: Any) -> Any:
    return "" if value is None else value


def _generated_value(value: Any) -> Any:
    """Render a result value for a cell; lists become lines, dicts ``k=v`` pairs."""
    if isinstance(value, (list, tuple)):
        return "\n".join(str(_generated_value(item)) for item in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}={_generated_value(v)}" for k, v in value.items())
    return _cell_value(value)


def read_sheets(path: Path) -> dict[str, Table]:
    """Every sheet in a workbook as a ``Table`` (first row is the header)."""
    if not path.is_file():
        raise MergeError(f"workbook not found: {path}")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        tables: dict[str, Table] = {}
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                tables[sheet.title] = Table(columns=[])
                continue
            columns = [str(name) if name is not None else "" for name in header]
            table = Table(columns=columns)
            for values in rows:
                if values is None or all(value in (None, "") for value in values):
                    continue
                table.rows.append(
                    {name: _cell_value(value) for name, value in zip(columns, values)}
                )
            tables[sheet.title] = table
        return tables
    finally:
        book.close()


def read_results(path: Path) -> dict[str, dict[str, Any]]:
    """``results.jsonl`` keyed by ``record_id``; a later line wins over an earlier one."""
    if not path.is_file():
        raise MergeError(f"results file not found: {path}")
    results: dict[str, dict[str, Any]] = {}
    with path.open(encoding="utf-8") as handle:
        for number, line in enumerate(handle, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                entry = json.loads(line)
            except json.JSONDecodeError as exc:
                raise MergeError(f"{path}:{number}: invalid JSON: {exc}") from exc
            record_id = entry.get("record_id")
            if record_id is not None:
                results[str(record_id)] = entry
    return results


def _error_text(entry: Mapping[str, Any]) -> str:
    parts = []
    if entry.get("error"):
        parts.append(str(entry["error"]))
    parts.extend(str(issue) for issue in entry.get("issues") or [])
    return "\n".join(parts)


@dataclass(frozen=True)
class MergeResult:
    out_path: Path
    analysis: str
    row_count: int
    matched: int
    status_counts: Mapping[str, int]
    generated_columns: tuple[str, ...]
    unmatched_results: int


def merge_results(
    cases: Path,
    results: Path,
    out_path: Path,
    *,
    output_keys: Sequence[str],
    analysis: str,
    id_field: str = "CaseNumber",
    prefix: str | None = None,
    sheet: str = MAIN_SHEET,
    notes: Sequence[str] = (),
    allow_missing_results: bool = False,
) -> MergeResult:
    """Write ``out_path``: ``cases`` with one set of generated columns per record.

    With ``allow_missing_results`` a missing ``results`` file is treated as empty, so
    every row is marked ``not_run`` instead of raising.
    """
    tables = read_sheets(cases)
    if sheet not in tables:
        raise MergeError(f"sheet '{sheet}' not found in {cases}")
    main = tables[sheet]
    if main.rows and id_field not in main.columns:
        raise MergeError(f"id field '{id_field}' is not a column of sheet '{sheet}'")
    entries = read_results(results) if results.is_file() else {}
    if not results.is_file() and main.rows and not allow_missing_results:
        raise MergeError(f"results file not found: {results}")

    prefix = prefix or analysis
    keys = list(output_keys) + [STATUS_KEY, ERROR_KEY]
    generated = [f"{prefix}.{key}" for key in keys]
    clash = sorted(set(generated) & set(main.columns))
    if clash:
        raise MergeError(f"generated columns already exist in '{sheet}': {', '.join(clash)}")

    counts: dict[str, int] = {}
    matched = 0
    seen: set[str] = set()
    for row in main.rows:
        record_id = str(row.get(id_field, ""))
        entry = entries.get(record_id)
        if entry is None:
            status = MISSING_STATUS
            values: Mapping[str, Any] = {}
            error = ""
        else:
            matched += 1
            seen.add(record_id)
            status = str(entry.get("status") or "")
            values = entry.get("values") or {}
            error = _error_text(entry)
        counts[status] = counts.get(status, 0) + 1
        for key, column in zip(output_keys, generated):
            row[column] = _generated_value(values.get(key))
        row[f"{prefix}.{STATUS_KEY}"] = status
        row[f"{prefix}.{ERROR_KEY}"] = error
    main.columns = main.columns + generated

    fill = generated_fill()
    book = OpenpyxlWorkbook(write_only=True)
    write_table(book, sheet, main, fills={column: fill for column in generated})
    for title, table in tables.items():
        if title not in (sheet, LEGEND_SHEET):
            write_table(book, title, table)
    _write_legend(
        book,
        analysis=analysis,
        generated=generated,
        cases=cases,
        results=results,
        counts=counts,
        notes=notes,
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(out_path)

    return MergeResult(
        out_path=out_path,
        analysis=analysis,
        row_count=len(main.rows),
        matched=matched,
        status_counts=dict(sorted(counts.items())),
        generated_columns=tuple(generated),
        unmatched_results=len(set(entries) - seen),
    )


def _write_legend(
    book,
    *,
    analysis: str,
    generated: Sequence[str],
    cases: Path,
    results: Path,
    counts: Mapping[str, int],
    notes: Sequence[str],
) -> None:
    fill = generated_fill()
    rows: list[tuple[str, Any, Any]] = [
        (
            "Generated column",
            f"Light blue ({GENERATED_FILL_COLOR}) columns were produced by the "
            f"'{analysis}' analysis, not by the Salesforce query.",
            fill,
        ),
        ("Query column", "Unfilled columns come straight from Salesforce.", None),
    ]
    rows += [(column, "generated", fill) for column in generated]
    rows += [
        ("Cases source", str(cases), None),
        ("Results source", str(results), None),
    ]
    rows += [(f"status: {status}", count, None) for status, count in sorted(counts.items())]
    rows.append(("Merged at", datetime.now(timezone.utc).isoformat(timespec="seconds"), None))
    rows += [("Note", note, None) for note in notes]

    sheet = book.create_sheet(title=LEGEND_SHEET)
    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 100
    sheet.append([make_cell(sheet, "Item", header=True), make_cell(sheet, "Description", header=True)])
    for item, description, row_fill in rows:
        sheet.append([make_cell(sheet, item, fill=row_fill), make_cell(sheet, description)])


def merge_analysis_run(
    cases: Path,
    run_dir: Path,
    out_path: Path,
    *,
    analysis: str,
    id_field: str = "CaseNumber",
    prefix: str | None = None,
    notes: Sequence[str] = (),
    allow_missing_results: bool = False,
) -> MergeResult:
    """Merge ``<run_dir>/results.jsonl`` for a registered analysis into ``cases``."""
    from ..prompts import get_analysis

    try:
        definition = get_analysis(analysis)
    except Exception as exc:  # unknown analysis name
        raise MergeError(str(exc)) from exc
    return merge_results(
        cases,
        run_dir / RESULTS_NAME,
        out_path,
        output_keys=list(definition.output_columns),
        analysis=analysis,
        id_field=id_field,
        prefix=prefix,
        notes=notes,
        allow_missing_results=allow_missing_results,
    )
