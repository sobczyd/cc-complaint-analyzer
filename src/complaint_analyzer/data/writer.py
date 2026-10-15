"""Writing the analyzed workbook.

Each selected row is reproduced with all of its original columns, followed by the
analysis result columns and a short status block.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .workbook import RawRow

STATUS_COLUMN = "analysis_status"
ISSUES_COLUMN = "analysis_issues"
ERROR_COLUMN = "analysis_error"
STATUS_COLUMNS = (STATUS_COLUMN, ISSUES_COLUMN, ERROR_COLUMN)

OUTPUT_SHEET = "Analysis"
MAX_CELL_LENGTH = 32767  # Hard limit imposed by the xlsx format.


@dataclass(frozen=True)
class OutputRow:
    """One analyzed record, ready to be written."""

    record_id: str
    source_row: int
    status: str
    values: Mapping[str, str]
    issues: Sequence[str] = ()
    error: str = ""


def _clip(text: str) -> str:
    if len(text) <= MAX_CELL_LENGTH:
        return text
    return text[: MAX_CELL_LENGTH - 1] + "\u2026"


def _sanitize(text: str) -> str:
    """Strip control characters openpyxl refuses to write, and clip overlong cells."""
    cleaned = "".join(
        char for char in text if char in "\t\n\r" or ord(char) >= 32
    )
    return _clip(cleaned)


def write_workbook(
    path: Path,
    *,
    source_columns: Sequence[str],
    result_columns: Sequence[str],
    rows: Iterable[OutputRow],
    rows_by_id: Mapping[str, RawRow],
) -> int:
    """Write ``analyzed.xlsx`` and return the number of data rows written."""
    header = [*source_columns, *result_columns, *STATUS_COLUMNS]

    book = OpenpyxlWorkbook(write_only=True)
    sheet = book.create_sheet(title=OUTPUT_SHEET)

    try:
        sheet.append([_header_cell(sheet, name) for name in header])

        written = 0
        for row in rows:
            source = rows_by_id.get(row.record_id)
            line = [
                _text_cell(sheet, source.value(name) if source else "")
                for name in source_columns
            ]
            line.extend(_text_cell(sheet, row.values.get(name, "")) for name in result_columns)
            line.append(_text_cell(sheet, row.status))
            line.append(_text_cell(sheet, "\n".join(row.issues)))
            line.append(_text_cell(sheet, row.error))
            sheet.append(line)
            written += 1

        _set_widths(sheet, header)
        path.parent.mkdir(parents=True, exist_ok=True)
        book.save(path)
    finally:
        book.close()
    return written


def _text_cell(sheet, value: str):
    """A cell openpyxl will never reinterpret as a formula or an error code."""
    cell = WriteOnlyCell(sheet, value=_sanitize(value))
    cell.data_type = "s"
    return cell


def _header_cell(sheet, name: str):
    cell = _text_cell(sheet, name)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def _set_widths(sheet, header: Sequence[str]) -> None:
    for index, name in enumerate(header, start=1):
        width = min(max(len(name) + 2, 14), 60)
        sheet.column_dimensions[get_column_letter(index)].width = width
