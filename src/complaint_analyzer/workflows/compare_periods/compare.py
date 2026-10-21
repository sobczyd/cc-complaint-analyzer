"""Period-over-period comparison of clustered cases.

Time point A is the baseline. Time point B, clustered independently (or by the same run), is mapped
onto A's sub-clusters so both periods use the same labels. The Python side is deterministic: it loads the
two periods, builds the resolution requests, validates the worker answers, applies the map and computes
counts, shares, per-day rates and flags. Reasoning (the sub-cluster mapping) is done by Copilot worker agents
that write candidate JSON, which ``publish_response`` validates before accepting.
"""

from __future__ import annotations

import hashlib
import json
import re
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError

from ...prompts import get_analysis
from ...prompts.cluster_resolve.schema import ResolveResult

MANIFEST = "manifest.json"
REQUESTS_DIR = "requests"
RESPONSES_DIR = "responses"
PERIOD_CASES_FILE = "period_cases.jsonl"
BASELINE_FILE = "baseline_taxonomy.json"
PERIOD_B_FILE = "period_b_taxonomy.json"
MAPPING_FILE = "mapping.jsonl"
RESOLVED_FILE = "resolved_assignments.jsonl"
COMPARISON_FILE = "comparison.json"
SOURCES_FILE = "sources.json"
WORKBOOK_FILE = "comparison.xlsx"
SUMMARY_FILE = "summary.md"

ASSIGNMENTS_FILE = "case_assignments.jsonl"
TAXONOMY_FILE = "final_taxonomy.json"
CASE_WORKBOOKS = ("cases_with_clusters.xlsx", "cases_with_summaries.xlsx", "cases.xlsx")

COMPARISON_TYPES = ("clusters",)
FLAG_BASES = ("count", "per_day")
DEFAULT_BATCH_SIZE = 15
DEFAULT_MIN_CHANGE = 10
DEFAULT_MIN_CASES = 10
FLAG_ORDER = {"NEW": 0, "GONE": 1, "UPTICK": 2, "DECLINE": 3, "": 4}


class CompareError(Exception):
    """Raised for unusable inputs, requests or responses."""


# -- helpers ---------------------------------------------------------------------------------


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _norm(label: str) -> str:
    return re.sub(r"\s+", " ", label).strip().casefold()


def _write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def _read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise CompareError(f"file not found: {path}") from None
    except json.JSONDecodeError as exc:
        raise CompareError(f"{path} is not valid JSON: {exc}") from None


def _write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        raise CompareError(f"file not found: {path}")
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _to_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)[:10]
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def parse_date(value: str | None, name: str) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        raise CompareError(f"{name} must be YYYY-MM-DD, got '{value}'") from None


# -- run discovery ---------------------------------------------------------------------------


def list_cluster_dirs(run_dir: Path) -> list[dict[str, Any]]:
    """Cluster directories of a run, best candidate first (finalized, then most cases)."""
    root = run_dir / "clusters"
    export = _optional_json(run_dir / "export.json") or {}
    out: list[dict[str, Any]] = []
    if root.is_dir():
        for child in sorted(root.iterdir()):
            manifest = _optional_json(child / "manifest.json")
            if not child.is_dir() or manifest is None:
                continue
            finalized = (child / ASSIGNMENTS_FILE).is_file() and (child / TAXONOMY_FILE).is_file()
            out.append(
                {
                    "name": child.name,
                    "finalized": finalized,
                    "case_count": manifest.get("case_count"),
                    "batch_size": manifest.get("batch_size"),
                    "run_row_count": export.get("row_count"),
                    "modified": datetime.fromtimestamp(child.stat().st_mtime, UTC).isoformat(timespec="seconds"),
                }
            )
    out.sort(key=lambda item: (not item["finalized"], -(item["case_count"] or 0), item["modified"]))
    return out


def _optional_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return None


def case_filters(run_dir: Path) -> list[str] | None:
    """Case-level filters of a run's query with the date range removed, for consistency checks."""
    path = run_dir / "query.soql"
    if not path.is_file():
        return None
    match = re.search(r"\bFROM\s+Case\s+WHERE\b(.*)$", path.read_text(encoding="utf-8"), re.DOTALL | re.IGNORECASE)
    if not match:
        return None
    conditions: list[str] = []
    for line in match.group(1).splitlines():
        line = line.strip()
        if not line or line.upper().startswith("ORDER BY") or "CreatedDate" in line:
            continue
        conditions.append(re.sub(r"^AND\s+", "", line, flags=re.IGNORECASE))
    return conditions


# -- loading a period ------------------------------------------------------------------------


class PeriodSpec:
    def __init__(
        self,
        key: str,
        run_dir: Path,
        cluster_name: str | None,
        start: date | None,
        end: date | None,
        label: str | None,
    ) -> None:
        if start and end and start > end:
            raise CompareError(f"period {key}: start {start} is after end {end}")
        self.key = key
        self.run_dir = run_dir
        self.cluster_name = cluster_name
        self.start = start
        self.end = end
        self._label = label

    @property
    def label(self) -> str:
        if self._label:
            return self._label
        if self.start and self.end:
            return f"{self.start}..{self.end}"
        return self.run_dir.name


def _resolve_cluster_dir(spec: PeriodSpec) -> Path:
    if not spec.run_dir.is_dir():
        raise CompareError(f"run folder not found: {spec.run_dir}")
    name = spec.cluster_name
    if not name:
        candidates = [c for c in list_cluster_dirs(spec.run_dir) if c["finalized"]]
        if not candidates:
            raise CompareError(f"no finalized cluster directory under {spec.run_dir}\\clusters")
        name = candidates[0]["name"]
        spec.cluster_name = name
    directory = spec.run_dir / "clusters" / name
    for required in (ASSIGNMENTS_FILE, TAXONOMY_FILE):
        if not (directory / required).is_file():
            raise CompareError(f"{directory} is not finalized (missing {required})")
    return directory


def _read_case_dates(run_dir: Path, id_field: str, date_field: str) -> tuple[dict[str, date | None], str | None]:
    for name in CASE_WORKBOOKS:
        path = run_dir / name
        if not path.is_file():
            continue
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = book[book.sheetnames[0]]
            rows = ws.iter_rows(values_only=True)
            header = [_text(cell) for cell in next(rows, ())]
            if id_field not in header or date_field not in header:
                continue
            id_i, date_i = header.index(id_field), header.index(date_field)
            return {_text(r[id_i]): _to_date(r[date_i]) for r in rows if _text(r[id_i])}, name
        finally:
            book.close()
    return {}, None


def _taxonomy(directory: Path, prefix: str, counts: dict[tuple[str, str], int]) -> dict[str, Any]:
    """The run's final taxonomy with ids (``<prefix><n>`` / ``<prefix><n>.S<m>``) and in-period counts."""
    raw = _read_json(directory / TAXONOMY_FILE)
    clusters = []
    for i, cluster in enumerate(raw["clusters"], 1):
        subs = []
        for j, sub in enumerate(cluster["subclusters"], 1):
            subs.append(
                {
                    "subcluster_id": f"{prefix}{i}.S{j}",
                    "label": sub["label"],
                    "description": sub.get("description", ""),
                    "rationale": sub.get("rationale", ""),
                    "case_count": counts.get((cluster["label"], sub["label"]), 0),
                }
            )
        clusters.append(
            {
                "cluster_id": f"{prefix}{i}",
                "label": cluster["label"],
                "description": cluster.get("description", ""),
                "rationale": cluster.get("rationale", ""),
                "subclusters": subs,
            }
        )
    return {"clusters": clusters}


def load_period(spec: PeriodSpec, *, id_field: str, date_field: str) -> dict[str, Any]:
    directory = _resolve_cluster_dir(spec)
    assignments = _read_jsonl(directory / ASSIGNMENTS_FILE)
    taxonomy_pairs = {
        (c["label"], s["label"])
        for c in _read_json(directory / TAXONOMY_FILE)["clusters"]
        for s in c["subclusters"]
    }
    dates, date_source = _read_case_dates(spec.run_dir, id_field, date_field)
    warnings: list[str] = []
    filtering = bool(spec.start or spec.end)
    if filtering and date_source is None:
        raise CompareError(f"period {spec.key}: no case workbook with {id_field} and {date_field} in {spec.run_dir}")

    cases: list[dict[str, Any]] = []
    excluded = undated = 0
    seen: set[str] = set()
    for row in assignments:
        number = _text(row.get("case_number"))
        pair = (row["final_cluster"], row["final_subcluster"])
        if pair not in taxonomy_pairs:
            raise CompareError(f"case {number} is assigned to a label missing from {directory / TAXONOMY_FILE}")
        if number in seen:
            raise CompareError(f"duplicate case {number} in {directory / ASSIGNMENTS_FILE}")
        seen.add(number)
        created = dates.get(number)
        if filtering:
            if created is None:
                undated += 1
                continue
            if (spec.start and created < spec.start) or (spec.end and created > spec.end):
                excluded += 1
                continue
        cases.append(
            {
                "case_number": number,
                "created": created.isoformat() if created else None,
                "cluster": pair[0],
                "subcluster": pair[1],
            }
        )
    if not cases:
        raise CompareError(f"period {spec.key} ({spec.label}) has no cases after filtering")
    if undated:
        warnings.append(f"period {spec.key}: {undated} case(s) without a {date_field} were excluded")

    created_dates = [date.fromisoformat(c["created"]) for c in cases if c["created"]]
    first = spec.start or (min(created_dates) if created_dates else None)
    last = spec.end or (max(created_dates) if created_dates else None)
    days = (last - first).days + 1 if first and last else None

    export = _optional_json(spec.run_dir / "export.json") or {}
    row_count = export.get("row_count")
    if not filtering and isinstance(row_count, int) and len(assignments) < row_count:
        warnings.append(
            f"period {spec.key}: cluster directory '{spec.cluster_name}' covers {len(assignments)} of "
            f"{row_count} exported cases"
        )
    counts: dict[tuple[str, str], int] = {}
    for case in cases:
        key = (case["cluster"], case["subcluster"])
        counts[key] = counts.get(key, 0) + 1
    return {
        "spec": spec,
        "cluster_dir": directory,
        "cases": cases,
        "excluded_by_date": excluded,
        "undated": undated,
        "taxonomy": _taxonomy(directory, spec.key, counts),
        "first_date": first.isoformat() if first else None,
        "last_date": last.isoformat() if last else None,
        "days": days,
        "date_source": date_source,
        "assignments_rows": len(assignments),
        "warnings": warnings,
    }


# -- prepare ---------------------------------------------------------------------------------


def _render_baseline(period: dict[str, Any]) -> str:
    parts = [f"=== BASELINE (period A: {period['spec'].label}, {len(period['cases'])} cases) ==="]
    for cluster in period["taxonomy"]["clusters"]:
        subs = [s for s in cluster["subclusters"] if s["case_count"]]
        if not subs:
            continue
        total = sum(s["case_count"] for s in subs)
        parts.append(
            f"Cluster {cluster['cluster_id']}: {cluster['label']} ({total} cases)\n"
            f"  description: {cluster['description']}"
        )
        for sub in subs:
            parts.append(
                f"  Sub-cluster {sub['subcluster_id']}: {sub['label']} ({sub['case_count']} cases)\n"
                f"    description: {sub['description']}\n    rationale: {sub['rationale']}"
            )
    return "\n".join(parts)


def _period_b_subclusters(period: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for cluster in period["taxonomy"]["clusters"]:
        for sub in cluster["subclusters"]:
            if sub["case_count"]:
                out.append({**sub, "cluster_id": cluster["cluster_id"], "cluster_label": cluster["label"]})
    return out


def _render_b(subs: Sequence[dict[str, Any]], label: str) -> str:
    parts = [f"=== PERIOD B SUB-CLUSTERS TO RESOLVE ({label}) ==="]
    for sub in subs:
        parts.append(
            f"Sub-cluster {sub['subcluster_id']}: {sub['label']} ({sub['case_count']} cases)\n"
            f"  period B cluster: {sub['cluster_label']}\n"
            f"  description: {sub['description']}\n  rationale: {sub['rationale']}"
        )
    return "\n".join(parts)


def default_out_dir(a: PeriodSpec, b: PeriodSpec, base: Path = Path("data") / "runs") -> Path:
    def slug(text: str) -> str:
        return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")

    return base / f"compare_{slug(a.label)}_vs_{slug(b.label)}_{datetime.now().date().isoformat()}"


def prepare(
    out_dir: Path,
    a: PeriodSpec,
    b: PeriodSpec,
    *,
    date_field: str = "CreatedDate",
    id_field: str = "CaseNumber",
    comparison_type: str = "clusters",
    batch_size: int = DEFAULT_BATCH_SIZE,
    min_change: int = DEFAULT_MIN_CHANGE,
    min_cases: int = DEFAULT_MIN_CASES,
    flag_basis: str = "count",
    prompt_version: str | None = None,
    request_text: str = "",
) -> dict[str, Any]:
    if comparison_type not in COMPARISON_TYPES:
        raise CompareError(
            f"comparison type '{comparison_type}' is not supported yet (supported: {', '.join(COMPARISON_TYPES)})"
        )
    if flag_basis not in FLAG_BASES:
        raise CompareError(f"--flag-basis must be one of {', '.join(FLAG_BASES)}")
    if batch_size < 1:
        raise CompareError("--batch-size must be at least 1")
    if min_change < 0 or min_cases < 0:
        raise CompareError("thresholds must not be negative")
    if (out_dir / MANIFEST).exists():
        raise CompareError(f"{out_dir} is already prepared; resume it or choose a new directory")

    period_a = load_period(a, id_field=id_field, date_field=date_field)
    period_b = load_period(b, id_field=id_field, date_field=date_field)
    same_taxonomy = period_a["cluster_dir"].resolve() == period_b["cluster_dir"].resolve()

    warnings = [*period_a["warnings"], *period_b["warnings"]]
    overlap = {c["case_number"] for c in period_a["cases"]} & {c["case_number"] for c in period_b["cases"]}
    if overlap:
        warnings.append(f"{len(overlap)} case(s) appear in both periods (periods overlap)")
    days_a, days_b = period_a["days"], period_b["days"]
    if days_a and days_b and days_a != days_b:
        warnings.append(
            f"period lengths differ: A covers {days_a} days, B covers {days_b} days; compare shares and "
            "cases per day, not raw counts"
        )
    elif not (days_a and days_b):
        warnings.append("period length could not be determined for one period; per-day rates are blank")
    filters_a, filters_b = case_filters(a.run_dir), case_filters(b.run_dir)
    if filters_a is None or filters_b is None:
        warnings.append("case filters could not be read from query.soql for one run; verify they are equivalent")
    elif sorted(map(_norm, filters_a)) != sorted(map(_norm, filters_b)):
        warnings.append("the two runs were fetched with different case filters; the comparison may not be like for like")

    analysis = get_analysis("cluster_resolve")
    version = prompt_version or sorted(analysis.prompts, key=lambda v: int(v[1:]))[-1]
    b_subs = _period_b_subclusters(period_b)
    requests: list[dict[str, Any]] = []
    if not same_taxonomy:
        baseline_text = _render_baseline(period_a)
        for start in range(0, len(b_subs), batch_size):
            chunk = b_subs[start : start + batch_size]
            request_id = f"resolve_{len(requests) + 1:02d}"
            requests.append(
                {
                    "request_id": request_id,
                    "kind": "resolve",
                    "prompt": "cluster_resolve",
                    "prompt_version": version,
                    "prompt_hash": analysis.prompt_hash(version),
                    "system": analysis.system_prompt + "\n\n" + analysis.prompt(version),
                    "user": "Resolve these period B sub-clusters onto the baseline.\n\n"
                    + baseline_text
                    + "\n\n"
                    + _render_b(chunk, b.label),
                    "subclusters": [s["subcluster_id"] for s in chunk],
                    "baseline_subclusters": [
                        s["subcluster_id"]
                        for c in period_a["taxonomy"]["clusters"]
                        for s in c["subclusters"]
                        if s["case_count"]
                    ],
                }
            )
        for request in requests:
            _write_json(out_dir / REQUESTS_DIR / f"{request['request_id']}.json", request)

    out_dir.mkdir(parents=True, exist_ok=True)
    _write_jsonl(
        out_dir / PERIOD_CASES_FILE,
        [{"period": "A", **c} for c in period_a["cases"]] + [{"period": "B", **c} for c in period_b["cases"]],
    )
    _write_json(out_dir / BASELINE_FILE, period_a["taxonomy"])
    _write_json(out_dir / PERIOD_B_FILE, period_b["taxonomy"])
    manifest = {
        "created_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "request": request_text,
        "comparison_type": comparison_type,
        "date_field": date_field,
        "id_field": id_field,
        "batch_size": batch_size,
        "thresholds": {"min_change": min_change, "min_cases": min_cases, "flag_basis": flag_basis},
        "prompt_version": version,
        "same_taxonomy": same_taxonomy,
        "resolution_required": not same_taxonomy,
        "periods": {key: _period_manifest(period, filters) for key, period, filters in (
            ("A", period_a, filters_a), ("B", period_b, filters_b)
        )},
        "b_subcluster_count": len(b_subs),
        "request_ids": [r["request_id"] for r in requests],
        "warnings": warnings,
    }
    _write_json(out_dir / MANIFEST, manifest)
    return {
        "out_dir": str(out_dir),
        "resolution_required": not same_taxonomy,
        "request_count": len(requests),
        "b_subcluster_count": len(b_subs),
        "case_counts": {"A": len(period_a["cases"]), "B": len(period_b["cases"])},
        "days": {"A": days_a, "B": days_b},
        "cluster_dirs": {"A": a.cluster_name, "B": b.cluster_name},
        "warnings": warnings,
    }


def _period_manifest(period: dict[str, Any], filters: list[str] | None) -> dict[str, Any]:
    spec: PeriodSpec = period["spec"]
    directory: Path = period["cluster_dir"]
    request = ""
    request_path = spec.run_dir / "request.txt"
    if request_path.is_file():
        request = request_path.read_text(encoding="utf-8").strip()
    export = _optional_json(spec.run_dir / "export.json") or {}
    return {
        "label": spec.label,
        "run_dir": str(spec.run_dir),
        "cluster_name": spec.cluster_name,
        "cluster_dir": str(directory),
        "assignments_file": str(directory / ASSIGNMENTS_FILE),
        "assignments_sha256": _sha256(directory / ASSIGNMENTS_FILE),
        "assignments_rows": period["assignments_rows"],
        "case_workbook": period["date_source"],
        "start": spec.start.isoformat() if spec.start else None,
        "end": spec.end.isoformat() if spec.end else None,
        "first_case_date": period["first_date"],
        "last_case_date": period["last_date"],
        "days": period["days"],
        "case_count": len(period["cases"]),
        "excluded_by_date": period["excluded_by_date"],
        "run_request": request,
        "run_export_rows": export.get("row_count"),
        "run_fields_based_on": export.get("fields_based_on"),
        "case_filters": filters,
    }


# -- validation / publish / check ------------------------------------------------------------


def validate_resolve(request: dict[str, Any], data: Any) -> tuple[ResolveResult | None, list[str]]:
    try:
        result = ResolveResult.model_validate(data)
    except ValidationError as exc:
        return None, [f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors()[:20]]
    errors: list[str] = []
    expected = list(request["subclusters"])
    baseline = set(request["baseline_subclusters"])
    seen: dict[str, int] = {}
    for mapping in result.mappings:
        seen[mapping.subcluster_id] = seen.get(mapping.subcluster_id, 0) + 1
        if mapping.decision == "match" and mapping.baseline_subcluster_id not in baseline:
            errors.append(
                f"{mapping.subcluster_id}: baseline_subcluster_id '{mapping.baseline_subcluster_id}' "
                "is not a baseline sub-cluster"
            )
    unknown = sorted(set(seen) - set(expected))
    missing = [s for s in expected if s not in seen]
    repeated = sorted(s for s, count in seen.items() if count > 1)
    if unknown:
        errors.append(f"unknown sub-cluster ids: {', '.join(unknown[:20])}")
    if missing:
        errors.append(f"unmapped sub-clusters ({len(missing)}): {', '.join(missing[:20])}")
    if repeated:
        errors.append(f"sub-clusters mapped more than once: {', '.join(repeated[:20])}")
    return (None if errors else result), errors


def _load_candidate(path: Path) -> Any:
    text = path.read_text(encoding="utf-8").strip()
    fenced = re.match(r"^```(?:json)?\s*(.*?)\s*```$", text, re.DOTALL)
    if fenced:
        text = fenced.group(1)
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise CompareError(f"candidate is not valid JSON: {exc}") from None


def publish_response(directory: Path, request_id: str, candidate: Path) -> list[str]:
    """Validate a candidate and, when valid, store it as the request's response. Returns errors."""
    if not re.fullmatch(r"[A-Za-z0-9_]+", request_id):
        raise CompareError(f"unsafe request id: {request_id}")
    request_path = directory / REQUESTS_DIR / f"{request_id}.json"
    if not request_path.is_file():
        raise CompareError(f"unknown request id: {request_id}")
    result, errors = validate_resolve(_read_json(request_path), _load_candidate(candidate))
    if result is None:
        return errors
    _write_json(directory / RESPONSES_DIR / f"{request_id}.json", result.model_dump())
    return []


def check(directory: Path) -> dict[str, Any]:
    manifest = _read_json(directory / MANIFEST)
    incomplete: list[str] = []
    errors: dict[str, list[str]] = {}
    for request_id in manifest["request_ids"]:
        response = directory / RESPONSES_DIR / f"{request_id}.json"
        if not response.is_file():
            incomplete.append(request_id)
            continue
        result, problems = validate_resolve(_read_json(directory / REQUESTS_DIR / f"{request_id}.json"), _read_json(response))
        if result is None:
            incomplete.append(request_id)
            errors[request_id] = problems
    return {
        "ready": not incomplete,
        "resolution_required": manifest["resolution_required"],
        "incomplete_request_ids": incomplete,
        "errors": errors,
    }


# -- finalize --------------------------------------------------------------------------------


def build_mapping(directory: Path, manifest: dict[str, Any]) -> list[dict[str, Any]]:
    """One row per period B sub-cluster with the baseline label it resolves to."""
    baseline = _read_json(directory / BASELINE_FILE)
    period_b = _read_json(directory / PERIOD_B_FILE)
    by_id: dict[str, tuple[str, str]] = {}
    pair_label: dict[tuple[str, str], tuple[str, str]] = {}
    cluster_label: dict[str, str] = {}
    for cluster in baseline["clusters"]:
        cluster_label[_norm(cluster["label"])] = cluster["label"]
        for sub in cluster["subclusters"]:
            by_id[sub["subcluster_id"]] = (cluster["label"], sub["label"])
            pair_label[(_norm(cluster["label"]), _norm(sub["label"]))] = (cluster["label"], sub["label"])

    answers: dict[str, dict[str, Any]] = {}
    if manifest["resolution_required"]:
        status = check(directory)
        if not status["ready"]:
            raise CompareError(f"responses missing or invalid: {', '.join(status['incomplete_request_ids'])}")
        for request_id in manifest["request_ids"]:
            for mapping in _read_json(directory / RESPONSES_DIR / f"{request_id}.json")["mappings"]:
                answers[mapping["subcluster_id"]] = mapping

    rows: list[dict[str, Any]] = []
    for cluster in period_b["clusters"]:
        for sub in cluster["subclusters"]:
            if not sub["case_count"]:
                continue
            row: dict[str, Any] = {
                "subcluster_id": sub["subcluster_id"],
                "b_cluster": cluster["label"],
                "b_subcluster": sub["label"],
                "case_count": sub["case_count"],
            }
            answer = answers.get(sub["subcluster_id"])
            if answer is None:
                row.update(
                    decision="match",
                    resolved_cluster=cluster["label"],
                    resolved_subcluster=sub["label"],
                    baseline_subcluster_id=None,
                    confidence="high",
                    rationale="Both periods use the same taxonomy.",
                    status="same_taxonomy",
                )
            elif answer["decision"] == "match":
                target = by_id[answer["baseline_subcluster_id"]]
                row.update(
                    decision="match",
                    resolved_cluster=target[0],
                    resolved_subcluster=target[1],
                    baseline_subcluster_id=answer["baseline_subcluster_id"],
                    confidence=answer["confidence"],
                    rationale=answer["rationale"],
                    status="unchanged" if (_norm(target[0]), _norm(target[1])) == (_norm(cluster["label"]), _norm(sub["label"])) else "renamed",
                )
            else:
                new_cluster, new_sub = answer["new_cluster"], answer["new_subcluster"]
                existing = pair_label.get((_norm(new_cluster), _norm(new_sub)))
                if existing:
                    resolved, status = existing, "renamed"
                else:
                    resolved, status = (cluster_label.get(_norm(new_cluster), new_cluster), new_sub), "new"
                row.update(
                    decision=answer["decision"],
                    resolved_cluster=resolved[0],
                    resolved_subcluster=resolved[1],
                    baseline_subcluster_id=None,
                    confidence=answer["confidence"],
                    rationale=answer["rationale"],
                    status=status,
                )
            rows.append(row)
    # New labels written differently by different batches coalesce on their normalised form.
    canonical: dict[tuple[str, str], tuple[str, str]] = {}
    for row in rows:
        key = (_norm(row["resolved_cluster"]), _norm(row["resolved_subcluster"]))
        resolved = canonical.setdefault(key, (row["resolved_cluster"], row["resolved_subcluster"]))
        row["resolved_cluster"], row["resolved_subcluster"] = resolved
    return rows


def flag_for(
    count_a: int, count_b: int, days_a: int | None, days_b: int | None, thresholds: dict[str, Any]
) -> str:
    if count_a == 0 and count_b > 0:
        return "NEW"
    if count_b == 0 and count_a > 0:
        return "GONE"
    expected_a: float = count_a
    if thresholds["flag_basis"] == "per_day" and days_a and days_b:
        expected_a = count_a * days_b / days_a
    delta = count_b - expected_a
    if abs(delta) > thresholds["min_change"] and max(count_a, count_b) >= thresholds["min_cases"]:
        return "UPTICK" if delta > 0 else "DECLINE"
    return ""


def _metrics(
    count_a: int, count_b: int, total_a: int, total_b: int, days_a: int | None, days_b: int | None,
    thresholds: dict[str, Any],
) -> dict[str, Any]:
    share_a = round(100 * count_a / total_a, 2) if total_a else 0.0
    share_b = round(100 * count_b / total_b, 2) if total_b else 0.0
    return {
        "count_a": count_a,
        "share_a_pct": share_a,
        "per_day_a": round(count_a / days_a, 2) if days_a else None,
        "count_b": count_b,
        "share_b_pct": share_b,
        "per_day_b": round(count_b / days_b, 2) if days_b else None,
        "change_count": count_b - count_a,
        "change_share_pp": round(share_b - share_a, 2),
        "change_pct": round(100 * (count_b - count_a) / count_a, 1) if count_a else None,
        "flag": flag_for(count_a, count_b, days_a, days_b, thresholds),
    }


def compute(
    cases: Sequence[dict[str, Any]], manifest: dict[str, Any], b_sources: dict[tuple[str, str], list[str]]
) -> dict[str, Any]:
    """Aggregate resolved case labels into cluster and sub-cluster comparisons."""
    days_a, days_b = manifest["periods"]["A"]["days"], manifest["periods"]["B"]["days"]
    thresholds = manifest["thresholds"]
    sub_counts: dict[tuple[str, str], dict[str, int]] = {}
    cluster_counts: dict[str, dict[str, int]] = {}
    for case in cases:
        side = case["period"]
        sub = sub_counts.setdefault((case["resolved_cluster"], case["resolved_subcluster"]), {"A": 0, "B": 0})
        sub[side] += 1
        cluster_counts.setdefault(case["resolved_cluster"], {"A": 0, "B": 0})[side] += 1
    total_a = sum(1 for c in cases if c["period"] == "A")
    total_b = len(cases) - total_a

    subclusters = []
    for (cluster, sub), counts in sub_counts.items():
        subclusters.append(
            {
                "cluster": cluster,
                "subcluster": sub,
                **_metrics(counts["A"], counts["B"], total_a, total_b, days_a, days_b, thresholds),
                "b_source_labels": b_sources.get((cluster, sub), []),
            }
        )
    clusters = [
        {"cluster": cluster, **_metrics(c["A"], c["B"], total_a, total_b, days_a, days_b, thresholds)}
        for cluster, c in cluster_counts.items()
    ]
    key = lambda item: (-abs(item["change_count"]), -(item["count_a"] + item["count_b"]), item["cluster"])
    clusters.sort(key=key)
    subclusters.sort(key=lambda item: (*key(item), item["subcluster"]))
    return {
        "totals": {"count_a": total_a, "count_b": total_b, "days_a": days_a, "days_b": days_b},
        "clusters": clusters,
        "subclusters": subclusters,
    }


def finalize(directory: Path) -> dict[str, Any]:
    manifest = _read_json(directory / MANIFEST)
    mapping = build_mapping(directory, manifest)
    by_b = {(m["b_cluster"], m["b_subcluster"]): m for m in mapping}
    resolved: list[dict[str, Any]] = []
    for case in _read_jsonl(directory / PERIOD_CASES_FILE):
        if case["period"] == "A":
            row = {**case, "resolved_cluster": case["cluster"], "resolved_subcluster": case["subcluster"],
                   "mapping_status": "baseline", "mapping_confidence": ""}
        else:
            m = by_b[(case["cluster"], case["subcluster"])]
            row = {**case, "resolved_cluster": m["resolved_cluster"], "resolved_subcluster": m["resolved_subcluster"],
                   "mapping_status": m["status"], "mapping_confidence": m["confidence"]}
        resolved.append(row)
    b_sources: dict[tuple[str, str], list[str]] = {}
    for m in mapping:
        labels = b_sources.setdefault((m["resolved_cluster"], m["resolved_subcluster"]), [])
        original = f"{m['b_cluster']} / {m['b_subcluster']}"
        if original not in labels:
            labels.append(original)

    comparison = compute(resolved, manifest, b_sources)
    comparison["periods"] = {key: _public_period(p) for key, p in manifest["periods"].items()}
    comparison["thresholds"] = manifest["thresholds"]
    comparison["comparison_type"] = manifest["comparison_type"]
    comparison["warnings"] = manifest["warnings"]
    comparison["mapping_summary"] = _mapping_summary(mapping)

    _write_jsonl(
        directory / RESOLVED_FILE,
        [
            {
                "period": r["period"],
                "case_number": r["case_number"],
                "created": r["created"],
                "original_cluster": r["cluster"],
                "original_subcluster": r["subcluster"],
                "resolved_cluster": r["resolved_cluster"],
                "resolved_subcluster": r["resolved_subcluster"],
                "mapping_status": r["mapping_status"],
                "mapping_confidence": r["mapping_confidence"],
            }
            for r in resolved
        ],
    )
    _write_jsonl(directory / MAPPING_FILE, mapping)
    _write_json(directory / COMPARISON_FILE, comparison)
    _write_json(directory / SOURCES_FILE, _sources(manifest, comparison))
    _write_workbook(directory / WORKBOOK_FILE, manifest, comparison, resolved, mapping)
    (directory / SUMMARY_FILE).write_text(render_summary(manifest, comparison), encoding="utf-8")
    flags = {}
    for item in comparison["clusters"]:
        if item["flag"]:
            flags[item["flag"]] = flags.get(item["flag"], 0) + 1
    return {
        "out_dir": str(directory),
        "case_counts": {"A": comparison["totals"]["count_a"], "B": comparison["totals"]["count_b"]},
        "cluster_count": len(comparison["clusters"]),
        "subcluster_count": len(comparison["subclusters"]),
        "cluster_flags": flags,
        "mapping": comparison["mapping_summary"],
        "files": [WORKBOOK_FILE, SUMMARY_FILE, COMPARISON_FILE, RESOLVED_FILE, MAPPING_FILE, SOURCES_FILE],
        "warnings": manifest["warnings"],
    }


def _public_period(period: dict[str, Any]) -> dict[str, Any]:
    keys = ("label", "run_dir", "cluster_name", "start", "end", "first_case_date", "last_case_date", "days", "case_count")
    return {k: period[k] for k in keys}


def _mapping_summary(mapping: Sequence[dict[str, Any]]) -> dict[str, Any]:
    status: dict[str, int] = {}
    for m in mapping:
        status[m["status"]] = status.get(m["status"], 0) + 1
    return {
        "b_subclusters": len(mapping),
        "by_status": status,
        "low_confidence": sum(1 for m in mapping if m["confidence"] == "low"),
    }


def _sources(manifest: dict[str, Any], comparison: dict[str, Any]) -> dict[str, Any]:
    return {
        "request": manifest["request"],
        "created_at": manifest["created_at"],
        "finalized_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "comparison_type": manifest["comparison_type"],
        "date_field": manifest["date_field"],
        "prompt": {"name": "cluster_resolve", "version": manifest["prompt_version"]},
        "resolution_required": manifest["resolution_required"],
        "thresholds": manifest["thresholds"],
        "periods": manifest["periods"],
        "totals": comparison["totals"],
        "warnings": manifest["warnings"],
    }


# -- outputs ---------------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FLAG_FILLS = {
    "NEW": "C6E0B4",
    "GONE": "D9D9D9",
    "UPTICK": "F8CBAD",
    "DECLINE": "BDD7EE",
}
METRIC_COLUMNS = (
    ("Count A", "count_a"),
    ("Share A %", "share_a_pct"),
    ("Cases/day A", "per_day_a"),
    ("Count B", "count_b"),
    ("Share B %", "share_b_pct"),
    ("Cases/day B", "per_day_b"),
    ("Change (count)", "change_count"),
    ("Change (share, pp)", "change_share_pp"),
    ("Change %", "change_pct"),
    ("Flag", "flag"),
)


def _sheet(book: Workbook, title: str, header: Sequence[str], rows: Iterable[Sequence[Any]], *, first: bool = False):
    ws = book.active if first else book.create_sheet()
    ws.title = title
    ws.append(list(header))
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A2"
    return ws


def _auto_width(ws: Any, limit: int = 60) -> None:
    for column in ws.columns:
        width = max((len(str(cell.value)) for cell in column if cell.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 8), limit)


def _color_flags(ws: Any, flag_col: int) -> None:
    for row in ws.iter_rows(min_row=2):
        fill = FLAG_FILLS.get(str(row[flag_col].value or ""))
        if fill:
            row[flag_col].fill = PatternFill("solid", fgColor=fill)


def _write_workbook(
    path: Path, manifest: dict[str, Any], comparison: dict[str, Any], resolved: Sequence[dict[str, Any]],
    mapping: Sequence[dict[str, Any]],
) -> None:
    book = Workbook()
    a, b = comparison["periods"]["A"], comparison["periods"]["B"]
    t = manifest["thresholds"]
    summary_rows = [
        ("Comparison type", manifest["comparison_type"]),
        ("Request", manifest["request"]),
        ("Time point A (baseline)", a["label"]),
        ("A run", a["run_dir"]),
        ("A cluster directory", a["cluster_name"]),
        ("A date range", f"{a['first_case_date']} to {a['last_case_date']}"),
        ("A days", a["days"]),
        ("A cases", a["case_count"]),
        ("Time point B", b["label"]),
        ("B run", b["run_dir"]),
        ("B cluster directory", b["cluster_name"]),
        ("B date range", f"{b['first_case_date']} to {b['last_case_date']}"),
        ("B days", b["days"]),
        ("B cases", b["case_count"]),
        ("Date field", manifest["date_field"]),
        ("Case filters (A)", "; ".join(manifest["periods"]["A"]["case_filters"] or ["unknown"])),
        ("Case filters (B)", "; ".join(manifest["periods"]["B"]["case_filters"] or ["unknown"])),
        ("Taxonomy resolution", "mapped B onto A" if manifest["resolution_required"] else "same taxonomy (no mapping)"),
        ("Resolve prompt version", manifest["prompt_version"]),
        ("Flag: min absolute change", t["min_change"]),
        ("Flag: min cases in either period", t["min_cases"]),
        ("Flag basis", t["flag_basis"]),
        ("Warnings", "\n".join(manifest["warnings"]) or "none"),
    ]
    ws = _sheet(book, "Summary", ("Item", "Value"), summary_rows, first=True)
    _auto_width(ws, 100)
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    def totals_row(label_cols: int) -> list[Any]:
        tot = comparison["totals"]
        row: list[Any] = ["Total"] + [""] * (label_cols - 1)
        row += [tot["count_a"], 100.0, round(tot["count_a"] / tot["days_a"], 2) if tot["days_a"] else None,
                tot["count_b"], 100.0, round(tot["count_b"] / tot["days_b"], 2) if tot["days_b"] else None,
                tot["count_b"] - tot["count_a"], "", "", ""]
        return row

    ws = _sheet(
        book, "Clusters", ("Cluster", *(h for h, _ in METRIC_COLUMNS)),
        [[c["cluster"], *(c[k] for _, k in METRIC_COLUMNS)] for c in comparison["clusters"]],
    )
    ws.append(totals_row(1))
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    _color_flags(ws, len(METRIC_COLUMNS))
    _auto_width(ws)
    ws = _sheet(
        book, "Subclusters", ("Cluster", "Subcluster", *(h for h, _ in METRIC_COLUMNS), "B source sub-clusters"),
        [
            [s["cluster"], s["subcluster"], *(s[k] for _, k in METRIC_COLUMNS), " | ".join(s["b_source_labels"])]
            for s in comparison["subclusters"]
        ],
    )
    ws.append(totals_row(2))
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    _color_flags(ws, len(METRIC_COLUMNS) + 1)
    _auto_width(ws)
    ws = _sheet(
        book, "Case Assignments",
        ("Period", "CaseNumber", "Created", "Original cluster", "Original subcluster", "Resolved cluster",
         "Resolved subcluster", "Mapping status", "Mapping confidence"),
        [
            [r["period"], r["case_number"], r["created"], r["cluster"], r["subcluster"], r["resolved_cluster"],
             r["resolved_subcluster"], r["mapping_status"], r["mapping_confidence"]]
            for r in resolved
        ],
    )
    ws.auto_filter.ref = ws.dimensions
    _auto_width(ws, 45)
    ws = _sheet(
        book, "Mapping",
        ("B sub-cluster id", "B cluster", "B subcluster", "B cases", "Decision", "Status", "Resolved cluster",
         "Resolved subcluster", "Baseline sub-cluster id", "Confidence", "Rationale"),
        [
            [m["subcluster_id"], m["b_cluster"], m["b_subcluster"], m["case_count"], m["decision"], m["status"],
             m["resolved_cluster"], m["resolved_subcluster"], m["baseline_subcluster_id"], m["confidence"],
             m["rationale"]]
            for m in mapping
        ],
    )
    ws.auto_filter.ref = ws.dimensions
    _auto_width(ws, 50)
    sources = []
    for key in ("A", "B"):
        p = manifest["periods"][key]
        sources.append([key, p["label"], p["run_dir"], p["cluster_dir"], p["assignments_file"],
                        p["assignments_sha256"], p["assignments_rows"], p["case_workbook"], p["start"], p["end"],
                        p["excluded_by_date"], p["run_export_rows"], p["run_request"]])
    ws = _sheet(
        book, "Sources",
        ("Period", "Label", "Run folder", "Cluster directory", "Assignments file", "Assignments SHA-256",
         "Assignment rows", "Date source workbook", "Start", "End", "Excluded by date", "Run export rows",
         "Original run request"),
        sources,
    )
    _auto_width(ws, 60)
    legend = [
        ("Period A / B", "A is the baseline time point; B is compared against it."),
        ("Share %", "Cases in the cluster as a percentage of all cases in that period."),
        ("Cases/day", "Cases in the cluster divided by the number of days in the period."),
        ("Change (count)", "Count B minus Count A."),
        ("Change (share, pp)", "Share B minus Share A, in percentage points."),
        ("Change %", "Relative change in count versus A; blank when A is 0."),
        ("NEW", "No cases in A, at least one in B."),
        ("GONE", "Cases in A, none in B."),
        ("UPTICK / DECLINE",
         f"B minus A (basis: {t['flag_basis']}) exceeds {t['min_change']} cases and at least one period has "
         f"{t['min_cases']} or more cases."),
        ("Mapping status", "baseline = period A; unchanged / renamed = B sub-cluster matched a baseline sub-cluster "
                           "(renamed when the label differs); new = no baseline match; same_taxonomy = no mapping needed."),
    ]
    ws = _sheet(book, "Legend", ("Term", "Meaning"), legend)
    _auto_width(ws, 110)
    for key, fill in FLAG_FILLS.items():
        for row in ws.iter_rows(min_row=2):
            if str(row[0].value).startswith(key):
                row[0].fill = PatternFill("solid", fgColor=fill)
                break
    book.save(path)


def _fmt(value: Any, suffix: str = "") -> str:
    return "-" if value is None else f"{value}{suffix}"


def render_summary(manifest: dict[str, Any], comparison: dict[str, Any]) -> str:
    a, b = comparison["periods"]["A"], comparison["periods"]["B"]
    t = manifest["thresholds"]
    lines = [
        f"# Period comparison: {a['label']} (baseline) vs {b['label']}",
        "",
        f"Comparison type: {manifest['comparison_type']}. Created {manifest['created_at']}.",
        "",
        "| | A (baseline) | B |",
        "|---|---|---|",
        f"| Label | {a['label']} | {b['label']} |",
        f"| Run | {Path(a['run_dir']).name} | {Path(b['run_dir']).name} |",
        f"| Cluster directory | {a['cluster_name']} | {b['cluster_name']} |",
        f"| Dates | {a['first_case_date']} to {a['last_case_date']} | {b['first_case_date']} to {b['last_case_date']} |",
        f"| Days | {_fmt(a['days'])} | {_fmt(b['days'])} |",
        f"| Cases | {a['case_count']} | {b['case_count']} |",
        "",
    ]
    if manifest["warnings"]:
        lines += ["## Warnings", ""] + [f"- {w}" for w in manifest["warnings"]] + [""]
    mapping = comparison["mapping_summary"]
    if manifest["resolution_required"]:
        status = ", ".join(f"{n} {k}" for k, n in sorted(mapping["by_status"].items()))
        lines += [
            "## Taxonomy resolution",
            "",
            f"{mapping['b_subclusters']} period B sub-clusters mapped onto the baseline ({status}); "
            f"{mapping['low_confidence']} low-confidence.",
            "",
        ]
    else:
        lines += ["## Taxonomy resolution", "", "Both periods come from one clustering; no mapping was needed.", ""]

    def table(items: list[dict[str, Any]], *, sub: bool) -> list[str]:
        head = "| Cluster | " + ("Sub-cluster | " if sub else "")
        out = [head + "A | A % | B | B % | Change | Flag |", "|---|" + ("---|" if sub else "") + "---|---|---|---|---|---|"]
        for item in items:
            name = f"{item['cluster']} | {item['subcluster']} | " if sub else f"{item['cluster']} | "
            out.append(
                f"| {name}{item['count_a']} | {item['share_a_pct']} | {item['count_b']} | {item['share_b_pct']} | "
                f"{item['change_count']:+d} | {item['flag']} |"
            )
        return out

    flagged_clusters = [c for c in comparison["clusters"] if c["flag"]]
    lines += [
        f"## Clusters flagged (NEW, GONE, or change > {t['min_change']} with >= {t['min_cases']} cases, "
        f"basis {t['flag_basis']})",
        "",
    ]
    lines += table(flagged_clusters, sub=False) if flagged_clusters else ["None."]
    days_a, days_b = comparison["totals"]["days_a"], comparison["totals"]["days_b"]
    if t["flag_basis"] == "count" and days_a and days_b and days_a != days_b:
        lines += [
            "",
            f"Note: flags use raw counts but the periods are {days_a} and {days_b} days long, so UPTICK / DECLINE "
            "are skewed toward the longer period. Use the share columns, or rerun with `--flag-basis per_day`.",
        ]
    lines += ["", "## All clusters", ""] + table(comparison["clusters"], sub=False) + [""]
    for flag, title in (("NEW", "New in B"), ("GONE", "Gone in B"), ("UPTICK", "Largest upticks"), ("DECLINE", "Largest declines")):
        items = [s for s in comparison["subclusters"] if s["flag"] == flag][:10]
        lines += [f"## Sub-clusters: {title}", ""]
        lines += table(items, sub=True) if items else ["None."]
        lines.append("")
    lines += [
        "Full detail, including the case-level original and resolved clusters, is in `comparison.xlsx`, "
        "`resolved_assignments.jsonl` and `mapping.jsonl`; sources in `sources.json`.",
        "",
    ]
    return "\n".join(lines)
