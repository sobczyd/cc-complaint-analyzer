"""Batch clustering of summarized cases, then a final merge of the per-batch taxonomies.

The Python side is deterministic: it builds request files, validates responses and applies the merge
mapping. Reasoning (the clustering and the merge) is done by Copilot worker agents that write
candidate JSON, which ``publish_response`` validates before accepting.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, ValidationError

from ...prompts import get_analysis
from ...prompts.base import AnalysisError
from ...prompts.case_cluster.schema import BatchClusterResult, MergeResult

SUMMARY_COLUMNS = (
    "case_summary.summary",
    "case_summary.reported_issue",
    "case_summary.product",
    "case_summary.actions_taken",
    "case_summary.resolution",
    "case_summary.status",
    "case_summary.key_points",
)
CASE_SEPARATOR = "=== Case ==="
MERGE_REQUEST_ID = "merge"
SAMPLE_CASES = 5

MANIFEST = "manifest.json"
REQUESTS_DIR = "requests"
RESPONSES_DIR = "responses"
CASES_FILE = "cases.jsonl"
TAXONOMY_FILE = "final_taxonomy.json"
ASSIGNMENTS_FILE = "case_assignments.jsonl"
WORKBOOK_FILE = "clusters.xlsx"


class ClusterError(Exception):
    """Raised for unusable inputs, requests or responses."""


# -- helpers ---------------------------------------------------------------------------------


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def _read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise ClusterError(f"file not found: {path}") from None
    except json.JSONDecodeError as exc:
        raise ClusterError(f"{path} is not valid JSON: {exc}") from None


def _norm(label: str) -> str:
    return re.sub(r"\s+", " ", label).strip().casefold()


def _load_manifest(directory: Path) -> dict[str, Any]:
    return _read_json(directory / MANIFEST)


def render_case(case_number: str, values: dict[str, str]) -> str:
    """One case as ``{Column name}: {Column contents}`` lines; empty columns are skipped."""
    lines = [CASE_SEPARATOR, f"CaseNumber: {case_number}"]
    for column in SUMMARY_COLUMNS:
        value = values.get(column, "")
        if value:
            lines.append(f"{column}: {value}")
    return "\n".join(lines)


def read_cases(
    workbook: Path, *, sheet: str | None, id_field: str, offset: int, limit: int | None
) -> list[dict[str, Any]]:
    """Return summarized cases (analysis ok and a non-empty summary) in file order."""
    if not workbook.is_file():
        raise ClusterError(f"workbook not found: {workbook}")
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        ws = book[sheet] if sheet else book[book.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [_text(cell) for cell in next(rows, ())]
        index = {name: i for i, name in enumerate(header) if name}
        missing = [name for name in (id_field, *SUMMARY_COLUMNS) if name not in index]
        if missing:
            raise ClusterError(f"workbook is missing column(s): {', '.join(missing)}")
        status_index = index.get("case_summary.analysis_status")
        cases: list[dict[str, Any]] = []
        seen: set[str] = set()
        for row in rows:
            case_number = _text(row[index[id_field]])
            if not case_number:
                continue
            if status_index is not None and _text(row[status_index]).lower() != "ok":
                continue
            values = {column: _text(row[index[column]]) for column in SUMMARY_COLUMNS}
            if not values["case_summary.summary"]:
                continue
            if case_number in seen:
                raise ClusterError(f"duplicate {id_field} {case_number}")
            seen.add(case_number)
            cases.append({"case_number": case_number, "values": values})
    finally:
        book.close()
    selected = cases[offset:] if limit is None else cases[offset : offset + limit]
    if not selected:
        raise ClusterError("no summarized cases selected (check --offset / --limit)")
    return selected


# -- prepare ---------------------------------------------------------------------------------


def prepare(
    workbook: Path,
    out_dir: Path,
    *,
    sheet: str | None = None,
    id_field: str = "CaseNumber",
    offset: int = 0,
    limit: int | None = 100,
    batch_size: int = 100,
    prompt_version: str | None = None,
    merge_version: str | None = None,
) -> dict[str, Any]:
    if batch_size < 1:
        raise ClusterError("--batch-size must be at least 1")
    if (out_dir / MANIFEST).exists():
        raise ClusterError(f"{out_dir} is already prepared; resume it or choose a new directory")
    analysis = get_analysis("case_cluster")
    version = prompt_version or sorted(analysis.prompts, key=lambda v: int(v[1:]))[-1]
    merge = get_analysis("case_cluster_merge")
    merge_ver = merge_version or sorted(merge.prompts, key=lambda v: int(v[1:]))[-1]

    cases = read_cases(workbook, sheet=sheet, id_field=id_field, offset=offset, limit=limit)
    batches: list[dict[str, Any]] = []
    for start in range(0, len(cases), batch_size):
        chunk = cases[start : start + batch_size]
        batch_id = f"batch_{len(batches) + 1:02d}"
        case_numbers = [case["case_number"] for case in chunk]
        user = (
            f"Cluster these {len(chunk)} cases.\n\n"
            + "\n\n".join(render_case(c["case_number"], c["values"]) for c in chunk)
        )
        _write_json(
            out_dir / REQUESTS_DIR / f"{batch_id}.json",
            {
                "request_id": batch_id,
                "kind": "batch",
                "prompt": "case_cluster",
                "prompt_version": version,
                "prompt_hash": analysis.prompt_hash(version),
                "system": analysis.system_prompt + "\n\n" + analysis.prompt(version),
                "user": user,
                "case_numbers": case_numbers,
            },
        )
        batches.append({"batch_id": batch_id, "case_count": len(case_numbers)})

    with (out_dir / CASES_FILE).open("w", encoding="utf-8") as handle:
        for case in cases:
            handle.write(json.dumps(case, ensure_ascii=False) + "\n")
    manifest = {
        "workbook": str(workbook),
        "sheet": sheet,
        "id_field": id_field,
        "offset": offset,
        "limit": limit,
        "batch_size": batch_size,
        "case_count": len(cases),
        "prompt_version": version,
        "merge_prompt_version": merge_ver,
        "batches": batches,
        "merge_required": len(batches) > 1,
    }
    _write_json(out_dir / MANIFEST, manifest)
    return manifest


# -- validation ------------------------------------------------------------------------------


def _parse(model: type[BaseModel], data: Any) -> tuple[BaseModel | None, list[str]]:
    try:
        return model.model_validate(data), []
    except ValidationError as exc:
        return None, [f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors()[:20]]


def validate_batch(request: dict[str, Any], data: Any) -> tuple[BatchClusterResult | None, list[str]]:
    result, errors = _parse(BatchClusterResult, data)
    if result is None:
        return None, errors
    assert isinstance(result, BatchClusterResult)
    expected = list(request["case_numbers"])
    seen: dict[str, int] = {}
    cluster_ids: set[str] = set()
    sub_ids: set[str] = set()
    for cluster in result.clusters:
        if cluster.cluster_id in cluster_ids:
            errors.append(f"duplicate cluster_id {cluster.cluster_id}")
        cluster_ids.add(cluster.cluster_id)
        for sub in cluster.subclusters:
            if sub.subcluster_id in sub_ids:
                errors.append(f"duplicate subcluster_id {sub.subcluster_id}")
            sub_ids.add(sub.subcluster_id)
            if not sub.subcluster_id.startswith(f"{cluster.cluster_id}."):
                errors.append(f"subcluster_id {sub.subcluster_id} must start with {cluster.cluster_id}.")
            for case in sub.case_numbers:
                seen[case] = seen.get(case, 0) + 1
    unknown = sorted(set(seen) - set(expected))
    missing = [case for case in expected if case not in seen]
    repeated = sorted(case for case, count in seen.items() if count > 1)
    if unknown:
        errors.append(f"unknown case numbers: {', '.join(unknown[:20])}")
    if missing:
        errors.append(f"missing case numbers ({len(missing)}): {', '.join(missing[:20])}")
    if repeated:
        errors.append(f"case numbers in more than one sub-cluster: {', '.join(repeated[:20])}")
    return (None if errors else result), errors


def validate_merge(request: dict[str, Any], data: Any) -> tuple[MergeResult | None, list[str]]:
    result, errors = _parse(MergeResult, data)
    if result is None:
        return None, errors
    assert isinstance(result, MergeResult)
    final: dict[str, dict[str, str]] = {}
    for cluster in result.final_clusters:
        key = _norm(cluster.label)
        if key in final:
            errors.append(f"duplicate final cluster label '{cluster.label}'")
        subs = final.setdefault(key, {})
        for sub in cluster.subclusters:
            if _norm(sub.label) in subs:
                errors.append(f"duplicate sub-cluster label '{sub.label}' in '{cluster.label}'")
            subs[_norm(sub.label)] = sub.label
    expected = {(item["batch_id"], item["subcluster_id"]) for item in request["subclusters"]}
    mapped: dict[tuple[str, str], int] = {}
    targets: set[tuple[str, str]] = set()
    for mapping in result.mappings:
        key = (mapping.batch_id, mapping.subcluster_id)
        mapped[key] = mapped.get(key, 0) + 1
        cluster_key, sub_key = _norm(mapping.final_cluster), _norm(mapping.final_subcluster)
        if cluster_key not in final or sub_key not in final[cluster_key]:
            errors.append(
                f"mapping {key[0]}/{key[1]} targets undefined '{mapping.final_cluster}' / "
                f"'{mapping.final_subcluster}'"
            )
        else:
            targets.add((cluster_key, sub_key))
    unknown = sorted(set(mapped) - expected)
    missing = sorted(expected - set(mapped))
    repeated = sorted(key for key, count in mapped.items() if count > 1)
    if unknown:
        errors.append(f"mappings for unknown sub-clusters: {unknown[:20]}")
    if missing:
        errors.append(f"unmapped batch sub-clusters ({len(missing)}): {missing[:20]}")
    if repeated:
        errors.append(f"sub-clusters mapped more than once: {repeated[:20]}")
    unused = [
        f"{cluster.label} / {sub.label}"
        for cluster in result.final_clusters
        for sub in cluster.subclusters
        if (_norm(cluster.label), _norm(sub.label)) not in targets
    ]
    if unused:
        errors.append(f"final sub-clusters with no mapping: {unused[:20]}")
    return (None if errors else result), errors


def _load_candidate(path: Path) -> Any:
    text = path.read_text(encoding="utf-8").strip()
    fenced = re.match(r"^```(?:json)?\s*(.*?)\s*```$", text, re.DOTALL)
    if fenced:
        text = fenced.group(1)
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise ClusterError(f"candidate is not valid JSON: {exc}") from None


def _validate_request_response(directory: Path, request_id: str, data: Any) -> tuple[BaseModel | None, list[str]]:
    request = _read_json(directory / REQUESTS_DIR / f"{request_id}.json")
    if request["kind"] == "merge":
        return validate_merge(request, data)
    return validate_batch(request, data)


def publish_response(directory: Path, request_id: str, candidate: Path) -> list[str]:
    """Validate a candidate and, when valid, store it as the request's response. Returns errors."""
    if not re.fullmatch(r"[A-Za-z0-9_]+", request_id):
        raise ClusterError(f"unsafe request id: {request_id}")
    if not (directory / REQUESTS_DIR / f"{request_id}.json").is_file():
        raise ClusterError(f"unknown request id: {request_id}")
    result, errors = _validate_request_response(directory, request_id, _load_candidate(candidate))
    if result is None:
        return errors
    _write_json(directory / RESPONSES_DIR / f"{request_id}.json", result.model_dump())
    return []


# -- check / merge request -------------------------------------------------------------------


def _request_ids(directory: Path) -> list[str]:
    manifest = _load_manifest(directory)
    ids = [batch["batch_id"] for batch in manifest["batches"]]
    if (directory / REQUESTS_DIR / f"{MERGE_REQUEST_ID}.json").is_file():
        ids.append(MERGE_REQUEST_ID)
    return ids


def check(directory: Path) -> dict[str, Any]:
    incomplete: list[str] = []
    errors: dict[str, list[str]] = {}
    for request_id in _request_ids(directory):
        response = directory / RESPONSES_DIR / f"{request_id}.json"
        if not response.is_file():
            incomplete.append(request_id)
            continue
        result, problems = _validate_request_response(directory, request_id, _read_json(response))
        if result is None:
            incomplete.append(request_id)
            errors[request_id] = problems
    return {"ready": not incomplete, "incomplete_request_ids": incomplete, "errors": errors}


def prepare_merge(directory: Path) -> dict[str, Any]:
    manifest = _load_manifest(directory)
    if not manifest["merge_required"]:
        return {"merge_required": False}
    status = check(directory)
    pending = [i for i in status["incomplete_request_ids"] if i != MERGE_REQUEST_ID]
    if pending:
        raise ClusterError(f"batch responses missing or invalid: {', '.join(pending)}")
    analysis = get_analysis("case_cluster_merge")
    version = manifest["merge_prompt_version"]
    parts: list[str] = []
    subclusters: list[dict[str, str]] = []
    for batch in manifest["batches"]:
        batch_id = batch["batch_id"]
        result = BatchClusterResult.model_validate(_read_json(directory / RESPONSES_DIR / f"{batch_id}.json"))
        parts.append(f"=== Batch {batch_id} ({batch['case_count']} cases) ===")
        for cluster in result.clusters:
            parts.append(
                f"Cluster {cluster.cluster_id}: {cluster.label} ({len(cluster.case_numbers)} cases)\n"
                f"  description: {cluster.description}\n  rationale: {cluster.rationale}"
            )
            for sub in cluster.subclusters:
                subclusters.append({"batch_id": batch_id, "subcluster_id": sub.subcluster_id})
                parts.append(
                    f"  Sub-cluster {sub.subcluster_id}: {sub.label} ({len(sub.case_numbers)} cases; "
                    f"samples: {', '.join(sub.case_numbers[:SAMPLE_CASES])})\n"
                    f"    description: {sub.description}\n    rationale: {sub.rationale}"
                )
    request = {
        "request_id": MERGE_REQUEST_ID,
        "kind": "merge",
        "prompt": "case_cluster_merge",
        "prompt_version": version,
        "prompt_hash": analysis.prompt_hash(version),
        "system": analysis.system_prompt + "\n\n" + analysis.prompt(version),
        "user": "Merge these batch clusterings.\n\n" + "\n\n".join(parts),
        "subclusters": subclusters,
    }
    _write_json(directory / REQUESTS_DIR / f"{MERGE_REQUEST_ID}.json", request)
    return {"merge_required": True, "request_id": MERGE_REQUEST_ID, "subclusters": len(subclusters)}


# -- finalize --------------------------------------------------------------------------------


@dataclass
class _Sub:
    batch_id: str
    cluster: Any
    sub: Any


def finalize(directory: Path) -> dict[str, Any]:
    manifest = _load_manifest(directory)
    status = check(directory)
    if not status["ready"]:
        raise ClusterError(f"responses incomplete: {status}")
    merge_path = directory / RESPONSES_DIR / f"{MERGE_REQUEST_ID}.json"
    if manifest["merge_required"] and not merge_path.is_file():
        raise ClusterError("merge response missing; run prepare-merge and publish the merge response")

    subs: dict[tuple[str, str], _Sub] = {}
    for batch in manifest["batches"]:
        result = BatchClusterResult.model_validate(
            _read_json(directory / RESPONSES_DIR / f"{batch['batch_id']}.json")
        )
        for cluster in result.clusters:
            for sub in cluster.subclusters:
                subs[(batch["batch_id"], sub.subcluster_id)] = _Sub(batch["batch_id"], cluster, sub)

    if manifest["merge_required"]:
        merged = MergeResult.model_validate(_read_json(merge_path))
        canonical = {
            (_norm(c.label), _norm(s.label)): (c.label, s.label)
            for c in merged.final_clusters
            for s in c.subclusters
        }
        mapping = {
            (m.batch_id, m.subcluster_id): (
                canonical[(_norm(m.final_cluster), _norm(m.final_subcluster))],
                m.rationale,
            )
            for m in merged.mappings
        }
        final_defs = merged.final_clusters
    else:
        mapping = {
            key: ((item.cluster.label, item.sub.label), "") for key, item in subs.items()
        }
        final_defs = None

    cases = [json.loads(line) for line in (directory / CASES_FILE).read_text(encoding="utf-8").splitlines() if line]
    case_values = {case["case_number"]: case["values"] for case in cases}

    assignments: list[dict[str, Any]] = []
    groups: dict[tuple[str, str], list[str]] = {}
    sources: dict[tuple[str, str], list[str]] = {}
    for key, item in subs.items():
        (final_cluster, final_sub), _ = mapping[key]
        groups.setdefault((final_cluster, final_sub), []).extend(item.sub.case_numbers)
        sources.setdefault((final_cluster, final_sub), []).append(f"{key[0]}/{key[1]}")
        for case in item.sub.case_numbers:
            assignments.append(
                {
                    "case_number": case,
                    "batch_id": item.batch_id,
                    "batch_cluster_id": item.cluster.cluster_id,
                    "batch_cluster": item.cluster.label,
                    "batch_subcluster_id": item.sub.subcluster_id,
                    "batch_subcluster": item.sub.label,
                    "final_cluster": final_cluster,
                    "final_subcluster": final_sub,
                }
            )
    order = {c["case_number"]: i for i, c in enumerate(cases)}
    assignments.sort(key=lambda row: order[row["case_number"]])
    if sorted(order) != sorted(row["case_number"] for row in assignments):
        raise ClusterError("final assignments do not cover each case exactly once")

    taxonomy: list[dict[str, Any]] = []
    if final_defs is not None:
        for cluster in final_defs:
            subs_out = []
            for sub in cluster.subclusters:
                key = (cluster.label, sub.label)
                subs_out.append(
                    {
                        "label": sub.label,
                        "description": sub.description,
                        "rationale": sub.rationale,
                        "case_numbers": groups.get(key, []),
                        "source_subclusters": sources.get(key, []),
                    }
                )
            taxonomy.append(
                {
                    "label": cluster.label,
                    "description": cluster.description,
                    "rationale": cluster.rationale,
                    "subclusters": subs_out,
                }
            )
    else:
        for item in subs.values():
            cluster = next((c for c in taxonomy if c["label"] == item.cluster.label), None)
            if cluster is None:
                cluster = {
                    "label": item.cluster.label,
                    "description": item.cluster.description,
                    "rationale": item.cluster.rationale,
                    "subclusters": [],
                }
                taxonomy.append(cluster)
            cluster["subclusters"].append(
                {
                    "label": item.sub.label,
                    "description": item.sub.description,
                    "rationale": item.sub.rationale,
                    "case_numbers": list(item.sub.case_numbers),
                    "source_subclusters": [f"{item.batch_id}/{item.sub.subcluster_id}"],
                }
            )
    for cluster in taxonomy:
        cluster["case_numbers"] = [c for sub in cluster["subclusters"] for c in sub["case_numbers"]]
        cluster["case_count"] = len(cluster["case_numbers"])
        for sub in cluster["subclusters"]:
            sub["case_count"] = len(sub["case_numbers"])
    taxonomy.sort(key=lambda c: -c["case_count"])

    _write_json(directory / TAXONOMY_FILE, {"case_count": len(assignments), "clusters": taxonomy})
    with (directory / ASSIGNMENTS_FILE).open("w", encoding="utf-8") as handle:
        for row in assignments:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    _write_workbook(directory / WORKBOOK_FILE, taxonomy, assignments, mapping, case_values)
    return {
        "case_count": len(assignments),
        "final_clusters": len(taxonomy),
        "final_subclusters": sum(len(c["subclusters"]) for c in taxonomy),
        "workbook": str(directory / WORKBOOK_FILE),
    }


def _write_workbook(
    path: Path,
    taxonomy: Sequence[dict[str, Any]],
    assignments: Sequence[dict[str, Any]],
    mapping: dict[tuple[str, str], tuple[tuple[str, str], str]],
    case_values: dict[str, dict[str, str]],
) -> None:
    book = Workbook()
    book.remove(book.active)

    def sheet(name: str, header: Sequence[str], rows: Sequence[Sequence[Any]], widths: Sequence[int]) -> None:
        ws = book.create_sheet(name)
        ws.append(list(header))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in rows:
            ws.append(list(row))
        for letter, width in zip("ABCDEFGHIJ", widths):
            ws.column_dimensions[letter].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    sheet(
        "Clusters",
        ["Cluster", "Cases", "Description", "Rationale", "Sub-clusters", "Case numbers"],
        [
            [
                c["label"],
                c["case_count"],
                c["description"],
                c["rationale"],
                "\n".join(f"{s['label']} ({s['case_count']})" for s in c["subclusters"]),
                ", ".join(c["case_numbers"]),
            ]
            for c in taxonomy
        ],
        [32, 8, 60, 60, 40, 60],
    )
    sheet(
        "Subclusters",
        ["Cluster", "Sub-cluster", "Cases", "Description", "Rationale", "Source batch sub-clusters", "Case numbers"],
        [
            [c["label"], s["label"], s["case_count"], s["description"], s["rationale"],
             ", ".join(s["source_subclusters"]), ", ".join(s["case_numbers"])]
            for c in taxonomy
            for s in c["subclusters"]
        ],
        [32, 32, 8, 60, 60, 28, 60],
    )
    sheet(
        "Case Assignments",
        ["CaseNumber", "Final cluster", "Final sub-cluster", "Batch", "Batch cluster", "Batch sub-cluster",
         "case_summary.reported_issue", "case_summary.summary"],
        [
            [a["case_number"], a["final_cluster"], a["final_subcluster"], a["batch_id"], a["batch_cluster"],
             a["batch_subcluster"], case_values[a["case_number"]].get("case_summary.reported_issue", ""),
             case_values[a["case_number"]].get("case_summary.summary", "")]
            for a in assignments
        ],
        [14, 32, 32, 10, 32, 32, 50, 70],
    )
    sheet(
        "Merge Map",
        ["Batch", "Batch sub-cluster", "Final cluster", "Final sub-cluster", "Rationale"],
        [[k[0], k[1], v[0][0], v[0][1], v[1]] for k, v in sorted(mapping.items())],
        [10, 16, 32, 32, 60],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
