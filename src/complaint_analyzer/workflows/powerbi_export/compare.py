"""Parity check: compare a generated ``powerbi.xlsx`` with the Power BI dashboard workbook.

Rows are joined on ``Source Identifier Raw`` and (by default) limited to cases that are Closed in
both files, because open cases keep changing between exports. Values are normalized before
comparing (line endings, ``_x000D_``, whitespace, pandas NA placeholders, int vs. text), and a
dashboard value that is a cut-off prefix of ours (the Salesforce report truncates long text) counts
as a match.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .derived import ASSESSMENT_QUESTIONS
from .export import PowerBIExportError
from .mapping import NA_STRINGS

KEY = "Source Identifier Raw"
STATE = "Source State"

# Trackwise / customer columns with no Salesforce source yet (handled in separate work).
UNSOURCED = frozenset(
    {
        "Project",
        "Complaint Receipt Date",
        "Date Opened",
        "Date Closed",
        "PR State",
        "Source Status",
        "Source Created Date",
        "Source Closed Date",
        "Event Description",
        "Evaluation Rationale",
        "kept_file",
        "Customer",
        "Customer / Technical Note",
    }
)
TIMING = frozenset({"export_date"})
# Columns this export derives itself; the dashboard leaves them blank.
DERIVED = frozenset({"serial_numbers"})
# The dashboard cuts long text (usually at 255 characters); shorter prefixes are real differences.
TRUNCATION_MIN = 250

_FILLS = {
    "Matched": "C6EFCE",
    "Not Matched": "FFC7CE",
    "Unsourced": "D9D9D9",
    "Timing": "FFEB9C",
    "Derived": "E2EFDA",
    "Not in dashboard": "DDEBF7",
}


def is_unsourced(column: str) -> bool:
    return column in UNSOURCED or column.endswith("^") or column.endswith("_Merged")


def normalize(value: Any) -> str:
    """Comparable text for a cell value."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value)
    if text in NA_STRINGS:
        return ""
    text = text.replace("_x000D_", "").replace("\r\n", "\n").replace("\r", "\n")
    text = " ".join(text.split())
    if re.fullmatch(r"-?\d+\.0+", text):
        text = text.split(".")[0]
    return text


def _text_match(ours: Any, dashboard: Any) -> bool:
    left, right = normalize(ours), normalize(dashboard)
    if left == right:
        return True
    if not right or not left.startswith(right):
        return False
    raw = "" if dashboard is None else str(dashboard).replace("_x000D_", "\r")
    return len(raw) >= TRUNCATION_MIN


def split_answers(text: Any, questions: tuple[str, ...]) -> list[str] | None:
    """Answers of a ``question:answer`` assessment block, or ``None`` when it does not follow ``questions``."""
    body = "" if text is None else str(text)
    answers: list[str] = []
    position = 0
    for index, question in enumerate(questions):
        marker = f"{question}:"
        start = body.find(marker, position)
        if start < 0:
            return None
        start += len(marker)
        if index + 1 < len(questions):
            end = body.find(f"{questions[index + 1]}:", start)
            if end < 0:
                return None
        else:
            end = len(body)
        answers.append(body[start:end])
        position = end
    return answers


def values_match(ours: Any, dashboard: Any, column: str | None = None) -> bool:
    if _text_match(ours, dashboard):
        return True
    questions = ASSESSMENT_QUESTIONS.get(column or "")
    if not questions:
        return False
    # The dashboard builds the blocks from report fields that are already cut, so compare answer by answer.
    mine, theirs = split_answers(ours, questions), split_answers(dashboard, questions)
    return mine is not None and theirs is not None and all(map(_text_match, mine, theirs))


def _key(value: Any) -> str:
    text = normalize(value)
    return text.lstrip("0") or text


def _read(path: Path, keys: set[str] | None = None) -> tuple[list[str], dict[str, dict[str, Any]]]:
    if not path.is_file():
        raise PowerBIExportError(f"workbook not found: {path}")
    book = load_workbook(path, read_only=True)
    try:
        rows = book.worksheets[0].iter_rows(values_only=True)
        headers = [str(name) for name in next(rows, ())]
        if KEY not in headers:
            raise PowerBIExportError(f"{path} has no '{KEY}' column")
        index = headers.index(KEY)
        records: dict[str, dict[str, Any]] = {}
        for row in rows:
            key = _key(row[index])
            if key and (keys is None or key in keys):
                records.setdefault(key, dict(zip(headers, row)))
        return headers, records
    finally:
        book.close()


def compare(ours_path: Path, dashboard_path: Path, *, state: str | None = "Closed") -> dict[str, Any]:
    """Per-column comparison for the cases present in both workbooks."""
    our_headers, ours = _read(ours_path)
    dash_headers, dashboard = _read(dashboard_path, set(ours))
    keys = [
        key
        for key in ours
        if key in dashboard
        and (state is None or (ours[key].get(STATE) == state and dashboard[key].get(STATE) == state))
    ]
    columns: list[dict[str, Any]] = []
    mismatches: list[dict[str, Any]] = []
    for column in dash_headers + [name for name in our_headers if name not in dash_headers]:
        in_ours, in_dash = column in our_headers, column in dash_headers
        if not in_dash:
            status = "Not in dashboard"
        elif is_unsourced(column):
            status = "Unsourced"
        elif column in TIMING:
            status = "Timing"
        elif column in DERIVED:
            status = "Derived"
        else:
            status = None
        matched = 0
        example: dict[str, Any] | None = None
        for key in keys:
            mine = ours[key].get(column) if in_ours else None
            theirs = dashboard[key].get(column) if in_dash else None
            if values_match(mine, theirs, column):
                matched += 1
                continue
            if status is None:
                mismatches.append({"case": key, "column": column, "ours": mine, "dashboard": theirs})
            example = example or {"case": key, "ours": mine, "dashboard": theirs}
        if status is None:
            status = "Matched" if matched == len(keys) else "Not Matched"
        columns.append(
            {
                "column": column,
                "status": status,
                "matched": matched,
                "compared": len(keys),
                "example_case": example["case"] if example else None,
                "example_ours": example["ours"] if example else None,
                "example_dashboard": example["dashboard"] if example else None,
            }
        )
    return {
        "cases": len(keys),
        "ours_cases": len(ours),
        "overlap": sum(key in dashboard for key in ours),
        "columns": columns,
        "mismatches": mismatches,
    }


def _cell(value: Any) -> Any:
    if isinstance(value, str) and len(value) > 32000:
        return value[:32000]
    return value


def write_report(result: dict[str, Any], out: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Columns"
    names = ["column", "status", "matched", "compared", "example_case", "example_ours", "example_dashboard"]
    sheet.append(names)
    for row in result["columns"]:
        sheet.append([_cell(row[name]) for name in names])
        fill = _FILLS.get(row["status"])
        if fill:
            sheet.cell(sheet.max_row, 2).fill = PatternFill("solid", start_color=fill)
    mismatch_sheet = book.create_sheet("Mismatches")
    mismatch_sheet.append(["case", "column", "ours", "dashboard"])
    for row in result["mismatches"]:
        mismatch_sheet.append([_cell(row[name]) for name in ("case", "column", "ours", "dashboard")])
    for target in (sheet, mismatch_sheet):
        for cell in target[1]:
            cell.font = Font(bold=True)
        target.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)
