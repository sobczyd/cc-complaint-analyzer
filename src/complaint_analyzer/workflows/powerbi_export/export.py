"""Write the Power BI workbook from an exported cases workbook."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from ..defect_mapping.kb import parse_link
from .mapping import build_row
from .serials import SERIAL_COLUMNS

DEFAULT_TEMPLATE = Path("data") / "input" / "salesforce.xlsx"


class PowerBIExportError(Exception):
    pass


def read_headers(template: Path) -> list[str]:
    if not template.is_file():
        raise PowerBIExportError(f"template workbook not found: {template}")
    book = load_workbook(template, read_only=True)
    try:
        header = next(book.active.iter_rows(max_row=1, values_only=True), None)
    finally:
        book.close()
    if not header:
        raise PowerBIExportError(f"template workbook has no header row: {template}")
    return [str(name) for name in header]


def read_sheet(workbook: Path, sheet: str) -> list[dict[str, Any]]:
    book = load_workbook(workbook, read_only=True)
    try:
        if sheet not in book.sheetnames:
            return []
        rows = book[sheet].iter_rows(values_only=True)
        names = [str(name) for name in next(rows, ())]
        return [dict(zip(names, row)) for row in rows if any(value is not None for value in row)]
    finally:
        book.close()


def load_kb_titles(kb_dir: Path | None) -> dict[str, str]:
    """KB article key -> title from a ``defect_mapping kb-normalize`` corpus (``articles.jsonl``)."""
    if kb_dir is None:
        return {}
    path = kb_dir / "articles.jsonl"
    if not path.is_file():
        raise PowerBIExportError(f"KB corpus has no articles.jsonl: {kb_dir}")
    titles: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            article = json.loads(line)
            if article.get("key") and article.get("title"):
                titles[article["key"]] = article["title"]
    return titles


def export_powerbi(
    cases: Path, out: Path, *, template: Path = DEFAULT_TEMPLATE, kb_dir: Path | None = None
) -> dict[str, Any]:
    template_headers = read_headers(template)
    headers = template_headers + [name for name in SERIAL_COLUMNS if name not in template_headers]
    case_rows = read_sheet(cases, "Cases")
    if not case_rows:
        raise PowerBIExportError(f"no rows on the Cases sheet of {cases}")
    problem_codes = _by_case(read_sheet(cases, "Problem Codes"))
    qualifications = _by_case(read_sheet(cases, "Complaint Qualifications"))
    titles = load_kb_titles(kb_dir)

    def kb_title(url: str) -> str:
        link = parse_link(url) if url else None
        return titles.get(link["key"], "") if link else ""

    exported_at = datetime.now().replace(microsecond=0)
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(headers)
    stats = {
        "cases": 0,
        "extra_problem_codes": 0,
        "extra_qualifications": 0,
        "kb_titles": 0,
        "serial_numbers": 0,
        "pre_post_serials": 0,
    }
    reference_index = headers.index("Reference Name")
    serials_index = headers.index("serial_numbers") if "serial_numbers" in headers else None
    pre_index, post_index = headers.index("pre_sn"), headers.index("post_sn")
    for case in case_rows:
        number = str(case.get("CaseNumber"))
        codes, cqs = problem_codes.get(number, []), qualifications.get(number, [])
        stats["extra_problem_codes"] += max(len(codes) - 1, 0)
        stats["extra_qualifications"] += max(len(cqs) - 1, 0)
        row = build_row(
            headers,
            case,
            codes[0] if codes else None,
            cqs[-1] if cqs else None,
            kb_title=kb_title,
            export_date=exported_at,
        )
        stats["kb_titles"] += bool(row[reference_index])
        stats["serial_numbers"] += serials_index is not None and bool(row[serials_index])
        stats["pre_post_serials"] += bool(row[pre_index] or row[post_index])
        sheet.append(row)
        stats["cases"] += 1
    sheet.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)
    return {**stats, "columns": len(headers), "template_columns": len(template_headers), "out": str(out)}


def _by_case(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(str(row.get("CaseNumber")), []).append(row)
    return grouped
