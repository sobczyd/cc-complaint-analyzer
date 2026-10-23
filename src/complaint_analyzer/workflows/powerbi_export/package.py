"""Power BI data-model package: one Excel file per table, written to ``<run>\\powerbi_export_package``."""

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from .export import DEFAULT_TEMPLATE, PowerBIExportError, export_powerbi, read_sheet
from .headlines import UNMAPPED, issue_group
from .serials import serial_attributes
from .trackwise import (
    CODES_FILE, FACT_COMPLAINT_COLUMNS, PRODUCT_FILE, QUERY_FILE, TrackwiseData, TrackwiseError, complaint_rows,
    load_trackwise,
)

PACKAGE_DIR = "powerbi_export_package"
KEY = "Source Identifier Raw"
DIM_ID_COLUMNS = [KEY, "Source Identifier", "PR ID", "Trackwise_Status"]
DIM_CLUSTER_COLUMNS = [
    "Case Number", "Period", "Resolved cluster", "Resolved subcluster", "Mapping status", "Mapping confidence",
        "Issue Group",
    ]
DIM_DEFECT_COLUMNS = [
    "Case Number", "issue_key", "issue_type", "confidence", "rationale", "version_justification", "kb_support",
]
SUMMARY_PREFIX = "case_summary."
DIM_SERIAL_COLUMNS = [
    KEY, "sn_type", "serial_number", "pre_sn", "post_sn", "pre_sn_date_wwyy", "post_sn_date_wwyy",
    "sn_date_wwyy", "sn_date_year", "sn_age_in_years", "sn_age_in_months", "sn_age_category",
]
DIM_SERIAL_DATE_COLUMNS = ("pre_sn_date_wwyy", "post_sn_date_wwyy", "sn_date_wwyy")
SUMMARIES_FILE = "cases_with_summaries.xlsx"
TRACKWISE_FILES = {"trackwise_query": QUERY_FILE, "trackwise_product": PRODUCT_FILE, "trackwise_codes": CODES_FILE}

README = """# Power BI export package

Generated {generated} from `{run}`. Each Excel file is one table; point Power BI at this folder and refresh.

| File | Table | Key |
|---|---|---|
| FactSource.xlsx (Sheet1) | FactSource | `Source Identifier Raw` |
| DimID.xlsx (DimID) | DimID | `Source Identifier Raw` |
| DimCluster.xlsx (DimCluster) | DimCluster | `Case Number` |
| DimDefect.xlsx (DimDefect) | DimDefect | `Case Number` (many rows per case) |
| DimSummary.xlsx (DimSummary) | DimSummary | `Source Identifier Raw` (the `case_summary.*` columns) |
| DimSerial.xlsx (DimSerial) | DimSerial | `Source Identifier Raw` (one row per old/new serial) |
| FactComplaint.xlsx (FactComplaint) | FactComplaint | `PR ID` (only when Trackwise data is supplied) |

## Relationships (single direction, filtering from the dimension to FactSource)
- FactSource[Source Identifier Raw] 1 : 1 DimID[Source Identifier Raw]
- FactSource[Source Identifier Raw] 1 : 1 DimCluster[Case Number]
- FactSource[Source Identifier Raw] 1 : * DimDefect[Case Number]
- FactSource[Source Identifier Raw] 1 : 1 DimSummary[Source Identifier Raw]
- FactSource[Source Identifier Raw] 1 : * DimSerial[Source Identifier Raw] (`sn_type` = old / new)
- FactSource[Source Identifier Raw] 1 : * FactComplaint[Source Identifier Raw] (a case can have several PRs)

All keys are text (leading zeros kept). A case can match several defects, so use a distinct count of cases
when slicing by defect.

## Notes
- FactComplaint is one row per Trackwise PR ID (Project = Complaint). Its `Source Identifier Raw` is the
  matching case number in FactSource; PRs with no case in this run keep the Trackwise Source Identifier.
  FactSource also carries the Trackwise columns of the PR matched to each case.
- DimID.`Trackwise_Status` is the matched PR's `PR State`; it is blank when no Trackwise data was supplied.
- DimCluster.`Issue Group` is a short headline derived from `Resolved cluster` by regex rules in
  `headlines.py`; new clusters show as `Unmapped` (and in the manifest warnings) until a rule is added.
- DimDefect holds real defect matches only; cases with no defect have no rows.
- `package_manifest.json` lists the sources, row counts and coverage gaps.
"""


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write(
    path: Path, sheet_name: str, headers: list[str], rows: list[list[Any]], date_columns: tuple[str, ...] = ()
) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for name in date_columns:
        for cell in sheet[get_column_letter(headers.index(name) + 1)][1:]:
            cell.number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"
    book.save(path)


def _as_date(value: Any) -> Any:
    return value.date() if isinstance(value, datetime) else value


def dim_serial_rows(fact_rows: list[dict[str, Any]], today: date | None = None) -> list[dict[str, Any]]:
    """Explode FactSource's old/new serial lists to one de-duplicated row per case, type and serial."""
    today = today or date.today()
    out: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for row in fact_rows:
        key = _text(row.get(KEY))
        case_level = {
            "pre_sn": row.get("pre_sn"),
            "post_sn": row.get("post_sn"),
            "pre_sn_date_wwyy": _as_date(row.get("pre_sn_date_wwyy")),
            "post_sn_date_wwyy": _as_date(row.get("post_sn_date_wwyy")),
        }
        for sn_type in ("old", "new"):
            for serial in _text(row.get(f"{sn_type}_serial_numbers")).split(","):
                serial = serial.strip()
                if not serial or (key, sn_type, serial) in seen:
                    continue
                seen.add((key, sn_type, serial))
                out.append(
                    {KEY: key, "sn_type": sn_type, "serial_number": serial, **case_level,
                     **serial_attributes(serial, today=today)}
                )
    return out


def _latest(directory: Path, filename: str, name: str | None, what: str) -> Path | None:
    """Pick ``directory\\<name>\\filename``; without a name use the newest one."""
    if name:
        path = directory / name / filename
        if not path.is_file():
            raise PowerBIExportError(f"{what} file not found: {path}")
        return path
    found = sorted(directory.glob(f"*/{filename}"), key=lambda item: item.stat().st_mtime) if directory.is_dir() else []
    return found[-1] if found else None


def _comparison_path(value: Path) -> Path:
    path = value / "comparison.xlsx" if value.is_dir() else value
    if not path.is_file():
        raise PowerBIExportError(f"comparison workbook not found: {path}")
    return path


def read_cluster_rows(comparison: Path | None, assignments: Path | None) -> tuple[list[dict[str, Any]], str]:
    if comparison is not None:
        book = load_workbook(comparison, read_only=True, data_only=True)
        try:
            if "Case Assignments" not in book.sheetnames:
                raise PowerBIExportError(f"no 'Case Assignments' sheet in {comparison}")
            rows = list(book["Case Assignments"].iter_rows(values_only=True))
        finally:
            book.close()
        header = [_text(cell) for cell in rows[0]]
        out = []
        for values in rows[1:]:
            record = dict(zip(header, values))
            if not _text(record.get("CaseNumber")):
                continue
            out.append(
                {
                    "Case Number": _text(record["CaseNumber"]),
                    "Period": _text(record.get("Period")),
                    "Resolved cluster": _text(record.get("Resolved cluster")),
                    "Resolved subcluster": _text(record.get("Resolved subcluster")),
                    "Mapping status": _text(record.get("Mapping status")),
                    "Mapping confidence": record.get("Mapping confidence"),
                }
            )
        return out, "comparison"
    if assignments is not None:
        out = []
        for line in assignments.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            record = json.loads(line)
            out.append(
                {
                    "Case Number": _text(record["case_number"]),
                    "Period": "",
                    "Resolved cluster": _text(record.get("final_cluster")),
                    "Resolved subcluster": _text(record.get("final_subcluster")),
                    "Mapping status": "baseline",
                    "Mapping confidence": None,
                }
            )
        return out, "case_assignments"
    return [], "none"


def read_defect_rows(mappings: Path | None) -> list[dict[str, Any]]:
    if mappings is None:
        return []
    book = load_workbook(mappings, read_only=True, data_only=True)
    try:
        if "Mappings" not in book.sheetnames:
            raise PowerBIExportError(f"no 'Mappings' sheet in {mappings}")
        rows = list(book["Mappings"].iter_rows(values_only=True))
    finally:
        book.close()
    header = [_text(cell) for cell in rows[0]]
    out = []
    for values in rows[1:]:
        record = dict(zip(header, values))
        if not _text(record.get("case_number")) or not _text(record.get("issue_key")):
            continue
        row = {column: record.get(column) for column in DIM_DEFECT_COLUMNS[1:]}
        row["Case Number"] = _text(record["case_number"])
        out.append(row)
    return out


def read_summary_rows(summaries: Path | None) -> tuple[list[str], list[dict[str, Any]]]:
    """Return the ``case_summary.*`` column names and one record per case (key in ``CaseNumber``)."""
    if summaries is None:
        return [], []
    book = load_workbook(summaries, read_only=True, data_only=True)
    try:
        if "Cases" not in book.sheetnames:
            raise PowerBIExportError(f"no 'Cases' sheet in {summaries}")
        rows = list(book["Cases"].iter_rows(values_only=True))
    finally:
        book.close()
    header = [_text(cell) for cell in rows[0]]
    columns = [name for name in header if name.startswith(SUMMARY_PREFIX)]
    out = []
    for values in rows[1:]:
        record = dict(zip(header, values))
        if _text(record.get("CaseNumber")):
            out.append(record)
    return columns, out


def build_package(
    run_dir: Path,
    *,
    comparison: Path | None = None,
    clusters: str | None = None,
    defects: str | None = None,
    out_dir: Path | None = None,
    template: Path = DEFAULT_TEMPLATE,
    kb_dir: Path | None = None,
    trackwise_dir: Path | None = None,
) -> dict[str, Any]:
    cases = run_dir / "cases.xlsx"
    if not cases.is_file():
        raise PowerBIExportError(f"cases.xlsx not found in {run_dir}")
    out_dir = out_dir or run_dir / PACKAGE_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    warnings: list[str] = []

    comparison_file = _comparison_path(comparison) if comparison else None
    assignments = None if comparison_file else _latest(run_dir / "clusters", "case_assignments.jsonl", clusters, "clusters")
    mappings = _latest(run_dir / "defects", "mappings.xlsx", defects, "defects")
    if comparison_file is None and assignments is None:
        warnings.append("no cluster source (comparison.xlsx or clusters\\<name>\\case_assignments.jsonl): DimCluster is empty")
    if mappings is None:
        warnings.append("no defects\\<name>\\mappings.xlsx: DimDefect is empty")

    fact_path = out_dir / "FactSource.xlsx"
    tw: TrackwiseData | None = None
    if trackwise_dir is not None:
        try:
            tw = load_trackwise(trackwise_dir)
        except TrackwiseError as error:
            raise PowerBIExportError(str(error)) from error
    stats = export_powerbi(cases, fact_path, template=template, kb_dir=kb_dir, trackwise=tw)
    fact_rows = read_sheet(fact_path, "Sheet1")
    keys = [_text(row.get(KEY)) for row in fact_rows]
    key_set = set(keys)
    if len(key_set) != len(keys):
        warnings.append(f"FactSource has {len(keys) - len(key_set)} duplicate {KEY} value(s)")

    dim_id = [
        [key, _text(row.get("Source Identifier")), row.get("PR ID"), _text(row.get("PR State"))]
        for key, row in zip(keys, fact_rows)
    ]
    _write(out_dir / "DimID.xlsx", "DimID", DIM_ID_COLUMNS, dim_id)

    complaints: list[dict[str, Any]] = []
    if tw is None:
        (out_dir / "FactComplaint.xlsx").unlink(missing_ok=True)
    else:
        complaints = complaint_rows(tw, fact_rows)
        _write(
            out_dir / "FactComplaint.xlsx", "FactComplaint", FACT_COMPLAINT_COLUMNS,
            [[row[column] for column in FACT_COMPLAINT_COLUMNS] for row in complaints],
        )
        if stats["trackwise_ambiguous"]:
            warnings.append(
                f"{stats['trackwise_ambiguous']} case(s) match several Trackwise PRs (the last Closed, else last, is used)"
            )
        if stats["trackwise_matched"] < len(fact_rows):
            warnings.append(f"{len(fact_rows) - stats['trackwise_matched']} FactSource case(s) have no Trackwise PR")

    cluster_rows, cluster_source = read_cluster_rows(comparison_file, assignments)
    chosen: dict[str, dict[str, Any]] = {}
    for row in cluster_rows:
        if row["Case Number"] in key_set:
            chosen[row["Case Number"]] = row
    orphan_clusters = len({row["Case Number"] for row in cluster_rows}) - len(chosen)
    unmapped_clusters: set[str] = set()
    for row in chosen.values():
        row["Issue Group"] = issue_group(row["Resolved cluster"])
        if row["Issue Group"] == UNMAPPED:
            unmapped_clusters.add(row["Resolved cluster"])
    if unmapped_clusters:
        warnings.append(
            "Issue Group is 'Unmapped' for new Resolved cluster(s); add a rule in headlines.py: "
            + "; ".join(sorted(unmapped_clusters))
        )
    _write(
        out_dir / "DimCluster.xlsx", "DimCluster", DIM_CLUSTER_COLUMNS,
        [[row[column] for column in DIM_CLUSTER_COLUMNS] for row in chosen.values()],
    )

    all_defects = read_defect_rows(mappings)
    defect_rows = [row for row in all_defects if row["Case Number"] in key_set]
    _write(
        out_dir / "DimDefect.xlsx", "DimDefect", DIM_DEFECT_COLUMNS,
        [[row[column] for column in DIM_DEFECT_COLUMNS] for row in defect_rows],
    )

    serial_rows = dim_serial_rows(fact_rows)
    _write(
        out_dir / "DimSerial.xlsx", "DimSerial", DIM_SERIAL_COLUMNS,
        [[row[column] for column in DIM_SERIAL_COLUMNS] for row in serial_rows],
        date_columns=DIM_SERIAL_DATE_COLUMNS,
    )

    summaries = run_dir / SUMMARIES_FILE
    summaries = summaries if summaries.is_file() else None
    if summaries is None:
        warnings.append(f"no {SUMMARIES_FILE}: DimSummary is empty")
    summary_columns, summary_records = read_summary_rows(summaries)
    if summaries is not None and not summary_columns:
        warnings.append(f"no {SUMMARY_PREFIX}* columns in {SUMMARIES_FILE}: DimSummary has no summary columns")
    summary_rows = [
        [_text(record["CaseNumber"])] + [record.get(column) for column in summary_columns]
        for record in summary_records
        if _text(record["CaseNumber"]) in key_set
    ]
    _write(out_dir / "DimSummary.xlsx", "DimSummary", [KEY] + summary_columns, summary_rows)
    missing_summary = key_set - {row[0] for row in summary_rows}
    if summaries is not None and missing_summary:
        warnings.append(f"{len(missing_summary)} FactSource case(s) have no summary")

    missing_cluster = [key for key in key_set if key not in chosen] if cluster_source != "none" else []
    if missing_cluster:
        warnings.append(f"{len(missing_cluster)} FactSource case(s) have no cluster")
    (out_dir / "README.md").write_text(
        README.format(generated=datetime.now().strftime("%Y-%m-%d %H:%M"), run=run_dir.name), encoding="utf-8"
    )
    sources = {
        "cases": cases, "comparison": comparison_file, "case_assignments": assignments, "mappings": mappings,
        "summaries": summaries,
    }
    if tw is not None:
        sources.update({name: tw.directory / file for name, file in TRACKWISE_FILES.items()})
    manifest = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "run_dir": str(run_dir),
        "cluster_source": cluster_source,
        "sources": {
            name: {"path": str(path), "sha256": _sha256(path)} for name, path in sources.items() if path is not None
        },
        "rows": {
            "FactSource": len(fact_rows),
            "DimID": len(dim_id),
            "DimCluster": len(chosen),
            "DimDefect": len(defect_rows),
            "DimSummary": len(summary_rows),
            "DimSerial": len(serial_rows),
            **({"FactComplaint": len(complaints)} if tw is not None else {}),
        },
        "coverage": {
            "cases_without_summary": len(missing_summary) if summaries is not None else len(key_set),
            "cases_without_cluster": len(missing_cluster),
            "cases_with_defect": len({row["Case Number"] for row in defect_rows}),
            "cases_with_serials": len({row[KEY] for row in serial_rows}),
            "cases_without_defect": len(key_set - {row["Case Number"] for row in defect_rows}),
            "orphan_cluster_cases_dropped": orphan_clusters,
            "orphan_defect_rows_dropped": len(all_defects) - len(defect_rows),
            **(
                {
                    "cases_with_trackwise": stats["trackwise_matched"],
                    "cases_without_trackwise": len(fact_rows) - stats["trackwise_matched"],
                    "cases_with_several_trackwise_prs": stats["trackwise_ambiguous"],
                }
                if tw is not None
                else {}
            ),
        },
        "warnings": warnings,
    }
    (out_dir / "package_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return {**manifest, "out_dir": str(out_dir), "export": stats}
