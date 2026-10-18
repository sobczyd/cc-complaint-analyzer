"""Run a versioned SOQL query with the Salesforce CLI and write the result to Excel.

The export produces one row per top-level record on a ``Cases`` sheet. Lookup fields
(``Owner.Name``) become dotted column names. Each child relationship subquery
(``Problem_Codes__r``) becomes:

* a joined-text column on the main sheet (``key: value`` lines, records separated by a
  blank line) plus a ``<name> (count)`` column, and
* its own detail sheet with one row per child record, keyed by the parent id field.
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ...workbook.writer import _sanitize
from ...queries import QueryError, load_queries, version_sort_key

QUERIES_DIR = Path(__file__).resolve().parents[2] / "queries"
MAIN_SHEET = "Cases"
DEFAULT_ID_FIELD = "CaseNumber"
DEFAULT_TIMEOUT_SECONDS = 900
COUNT_SUFFIX = " (count)"

QUERY_FILENAME = "query.soql"
RAW_FILENAME = "raw.json"
EXPORT_META_FILENAME = "export.json"
CASES_WORKBOOK_FILENAME = "cases.xlsx"

_SELECT = re.compile(r"\s*SELECT\s", re.IGNORECASE)
_FROM = re.compile(r"\sFROM\s+(\w+)", re.IGNORECASE)


class ExportError(Exception):
    """Raised when the query cannot be run or its output cannot be understood."""


# --------------------------------------------------------------------------- query


def resolve_query(query_set: str, version: str | None = None) -> tuple[str, str]:
    """Return ``(version, soql)`` for a query set; the latest version when omitted."""
    try:
        versions = load_queries(QUERIES_DIR / query_set)
    except QueryError as exc:
        raise ExportError(str(exc)) from exc
    if version is None:
        version = sorted(versions, key=version_sort_key)[-1]
    if version not in versions:
        raise ExportError(
            f"query set '{query_set}' has no version '{version}'. "
            f"Available: {', '.join(versions)}"
        )
    return version, versions[version]


def _split_top_level(text: str) -> list[str]:
    parts: list[str] = []
    depth = 0
    current: list[str] = []
    for char in text:
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
        if char == "," and depth == 0:
            parts.append("".join(current).strip())
            current = []
        else:
            current.append(char)
    tail = "".join(current).strip()
    if tail:
        parts.append(tail)
    return parts


def _select_and_from(soql: str) -> tuple[str, str] | None:
    """``(select list, FROM object)`` for the outermost SELECT, or None."""
    match = _SELECT.match(soql)
    if not match:
        return None
    depth = 0
    for index in range(match.end(), len(soql)):
        char = soql[index]
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
        elif depth == 0 and char.isspace():
            found = _FROM.match(soql, index)
            if found:
                return soql[match.end() : index], found.group(1)
    return None


@dataclass(frozen=True)
class SelectShape:
    """Fields named in a SOQL select list, used to fix column order."""

    fields: tuple[str, ...]
    children: Mapping[str, tuple[str, ...]]


def parse_select(soql: str) -> SelectShape:
    """Parse top-level fields and child subqueries from a SOQL statement.

    Best effort: anything not understood is ignored and discovered from the data.
    """
    parsed = _select_and_from(soql)
    if parsed is None:
        return SelectShape(fields=(), children={})
    fields: list[str] = []
    children: dict[str, tuple[str, ...]] = {}
    for part in _split_top_level(parsed[0]):
        if part.startswith("(") and part.endswith(")"):
            inner = _select_and_from(part[1:-1])
            if inner is not None:
                children[inner[1]] = tuple(
                    name for name in _split_top_level(inner[0]) if name
                )
        elif part and not any(char.isspace() for char in part):
            fields.append(part)
    return SelectShape(fields=tuple(fields), children=children)


# --------------------------------------------------------------------------- sf CLI


def find_sf() -> str:
    for name in ("sf", "sf.cmd"):
        found = shutil.which(name)
        if found:
            return found
    raise ExportError(
        "Salesforce CLI 'sf' was not found on PATH. Install it and run "
        "'sf org login web' first."
    )


def run_query(
    soql_path: Path,
    *,
    target_org: str | None = None,
    timeout: int = DEFAULT_TIMEOUT_SECONDS,
) -> dict[str, Any]:
    """Run ``sf data query --file`` and return the parsed JSON response.

    ``--file`` is required: a multi-line ``--query`` value is cut at the first newline.
    """
    command = [find_sf(), "data", "query", "--file", str(soql_path), "--json"]
    if target_org:
        command += ["--target-org", target_org]
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise ExportError(f"sf data query timed out after {timeout} seconds") from exc
    return parse_response(completed.stdout, stderr=completed.stderr)


def parse_response(stdout: str, *, stderr: str = "") -> dict[str, Any]:
    """Validate the ``sf ... --json`` envelope and return it."""
    try:
        payload = json.loads(stdout)
    except json.JSONDecodeError as exc:
        detail = (stderr or stdout).strip()[:2000]
        raise ExportError(f"sf did not return JSON: {detail}") from exc
    if not isinstance(payload, dict):
        raise ExportError("sf returned JSON that is not an object")
    if payload.get("status") != 0:
        message = payload.get("message") or payload.get("name") or "unknown error"
        raise ExportError(f"sf data query failed: {message}")
    result = payload.get("result")
    if not isinstance(result, dict) or not isinstance(result.get("records"), list):
        raise ExportError("sf response has no result.records list")
    return payload


# --------------------------------------------------------------------------- flatten


def _is_child_relationship(value: Any) -> bool:
    return isinstance(value, dict) and isinstance(value.get("records"), list)


def _scalar(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _flatten(record: Mapping[str, Any], prefix: str = "") -> dict[str, Any]:
    """Scalar and lookup values keyed by dotted path. Child relationships are skipped."""
    flat: dict[str, Any] = {}
    for key, value in record.items():
        if key == "attributes" or _is_child_relationship(value):
            continue
        name = f"{prefix}{key}"
        if isinstance(value, dict):
            flat.update(_flatten(value, prefix=f"{name}."))
        else:
            flat[name] = _scalar(value)
    return flat


def _ordered_columns(
    preferred: Sequence[str], discovered: Iterable[str]
) -> tuple[list[str], dict[str, str]]:
    """Column order: the SELECT list first, then anything else seen in the data.

    Returns the columns and a map from lower-cased data key to column name, because
    Salesforce echoes API names in their canonical case, not the query's case.
    """
    columns: list[str] = []
    by_lower: dict[str, str] = {}
    for name in preferred:
        if name.lower() not in by_lower:
            by_lower[name.lower()] = name
            columns.append(name)
    for name in discovered:
        lowered = name.lower()
        if lowered in by_lower:
            continue
        # A null lookup (``Asset: null``) arrives as a bare key; skip it when the
        # SELECT list already names columns underneath it.
        if any(column.lower().startswith(lowered + ".") for column in columns):
            continue
        by_lower[lowered] = name
        columns.append(name)
    return columns, by_lower


def friendly_name(relationship: str) -> str:
    """``Problem_Codes__r`` -> ``Problem Codes``."""
    base = re.sub(r"__r$", "", relationship)
    return base.replace("_", " ").strip() or relationship


@dataclass
class Table:
    columns: list[str]
    rows: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class FlatExport:
    """The main table plus one detail table per child relationship."""

    main: Table
    children: dict[str, Table]  # keyed by friendly name
    relationships: dict[str, str]  # friendly name -> relationship API name


def _child_text(child_rows: Sequence[Mapping[str, Any]], columns: Sequence[str]) -> str:
    blocks = []
    for row in child_rows:
        lines = [
            f"{column}: {row[column]}"
            for column in columns
            if column != "Id" and row.get(column) not in ("", None)
        ]
        if lines:
            blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def _lookup(record: Mapping[str, Any], name: str) -> Any:
    lowered = name.lower()
    for key, value in record.items():
        if key.lower() == lowered:
            return value
    return None


def _normalize(flat: Mapping[str, Any], by_lower: Mapping[str, str]) -> dict[str, Any]:
    row: dict[str, Any] = {}
    for key, value in flat.items():
        column = by_lower.get(key.lower())
        if column is not None:
            row[column] = value
    return row


def flatten_records(
    records: Sequence[Mapping[str, Any]],
    *,
    shape: SelectShape | None = None,
    id_field: str = DEFAULT_ID_FIELD,
) -> FlatExport:
    shape = shape or SelectShape(fields=(), children={})

    relationship_names: list[str] = list(shape.children)
    for record in records:
        for key, value in record.items():
            if _is_child_relationship(value) and not any(
                key.lower() == known.lower() for known in relationship_names
            ):
                relationship_names.append(key)
    relationship_lower = {name.lower() for name in relationship_names}

    flat_main = [_flatten(record) for record in records]
    discovered: dict[str, None] = {}
    for flat in flat_main:
        discovered.update(dict.fromkeys(flat))
    # A child relationship with no rows comes back as ``null``; it is not a column.
    discovered = {key: None for key in discovered if key.lower() not in relationship_lower}
    main_columns, main_map = _ordered_columns(shape.fields, discovered)
    main_rows = [_normalize(flat, main_map) for flat in flat_main]
    id_column = main_map.get(id_field.lower(), id_field)

    children: dict[str, Table] = {}
    relationships: dict[str, str] = {}
    for relationship in relationship_names:
        friendly = friendly_name(relationship)
        relationships[friendly] = relationship
        per_record: list[list[dict[str, Any]]] = []
        child_discovered: dict[str, None] = {}
        for record in records:
            value = _lookup(record, relationship)
            items = value["records"] if _is_child_relationship(value) else []
            flats = [_flatten(item) for item in items]
            for flat in flats:
                child_discovered.update(dict.fromkeys(flat))
            per_record.append(flats)
        child_columns, child_map = _ordered_columns(
            shape.children.get(relationship, ()), child_discovered
        )
        table = Table(columns=[id_column, *child_columns])
        for row, flats in zip(main_rows, per_record):
            normalized = [_normalize(flat, child_map) for flat in flats]
            for child in normalized:
                table.rows.append({id_column: row.get(id_column, ""), **child})
            row[friendly] = _child_text(normalized, child_columns)
            row[friendly + COUNT_SUFFIX] = len(normalized)
        children[friendly] = table
        main_columns += [friendly, friendly + COUNT_SUFFIX]

    return FlatExport(
        main=Table(columns=main_columns, rows=main_rows),
        children=children,
        relationships=relationships,
    )


# --------------------------------------------------------------------------- Excel


def make_cell(sheet, value: Any, *, header: bool = False, fill=None) -> WriteOnlyCell:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        cell = WriteOnlyCell(sheet, value=value)
    else:
        cell = WriteOnlyCell(sheet, value=_sanitize("" if value is None else str(value)))
        cell.data_type = "s"
    if header:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    if fill is not None:
        cell.fill = fill
    return cell


def sheet_title(name: str) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "-", name).strip() or "Sheet"
    return cleaned[:31]


def write_table(book, title: str, table: Table, *, fills: Mapping[str, Any] | None = None) -> None:
    """Append ``table`` as a new write-only sheet; ``fills`` maps column -> fill."""
    fills = fills or {}
    sheet = book.create_sheet(title=sheet_title(title))
    sheet.freeze_panes = "A2"
    for index, name in enumerate(table.columns, start=1):
        width = min(max(len(name) + 2, 14), 60)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.append(
        [make_cell(sheet, name, header=True, fill=fills.get(name)) for name in table.columns]
    )
    for row in table.rows:
        sheet.append(
            [make_cell(sheet, row.get(name, ""), fill=fills.get(name)) for name in table.columns]
        )
    if table.columns:
        last = get_column_letter(len(table.columns))
        sheet.auto_filter.ref = f"A1:{last}{len(table.rows) + 1}"


def write_cases_workbook(path: Path, export: FlatExport) -> None:
    """Write the main sheet and one sheet per child relationship.

    A header row is always written, even with zero records, so the workbook can be
    loaded as a dataset.
    """
    book = OpenpyxlWorkbook(write_only=True)
    write_table(book, MAIN_SHEET, export.main)
    for friendly, table in export.children.items():
        write_table(book, friendly, table)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


# --------------------------------------------------------------------------- orchestration


@dataclass(frozen=True)
class ExportResult:
    out_dir: Path
    query_set: str
    version: str
    target_org: str | None
    row_count: int
    done: bool
    columns: tuple[str, ...]
    child_sheets: tuple[str, ...]
    workbook: Path

    def as_json(self) -> dict[str, Any]:
        return {
            "query_set": self.query_set,
            "version": self.version,
            "target_org": self.target_org,
            "row_count": self.row_count,
            "done": self.done,
            "exported_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "main_sheet": MAIN_SHEET,
            "columns": list(self.columns),
            "child_sheets": list(self.child_sheets),
            "workbook": self.workbook.name,
        }


def export_query_set(
    query_set: str,
    out_dir: Path,
    *,
    version: str | None = None,
    target_org: str | None = None,
    from_json: Path | None = None,
    id_field: str = DEFAULT_ID_FIELD,
    timeout: int = DEFAULT_TIMEOUT_SECONDS,
) -> ExportResult:
    """Run (or replay) a query set and write ``query.soql``, ``raw.json``, ``cases.xlsx``.

    ``from_json`` replays a saved ``sf data query --json`` response instead of calling
    Salesforce, which is useful for tests and for rebuilding a workbook.
    """
    version, soql = resolve_query(query_set, version)
    out_dir.mkdir(parents=True, exist_ok=True)
    query_path = out_dir / QUERY_FILENAME
    query_path.write_text(soql, encoding="utf-8")

    if from_json is not None:
        payload = parse_response(Path(from_json).read_text(encoding="utf-8"))
    else:
        payload = run_query(query_path, target_org=target_org, timeout=timeout)
    raw_path = out_dir / RAW_FILENAME
    if from_json is None or Path(from_json).resolve() != raw_path.resolve():
        raw_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    result = payload["result"]
    export = flatten_records(result["records"], shape=parse_select(soql), id_field=id_field)
    workbook = out_dir / CASES_WORKBOOK_FILENAME
    write_cases_workbook(workbook, export)

    outcome = ExportResult(
        out_dir=out_dir,
        query_set=query_set,
        version=version,
        target_org=target_org,
        row_count=len(export.main.rows),
        done=bool(result.get("done", True)),
        columns=tuple(export.main.columns),
        child_sheets=tuple(sheet_title(name) for name in export.children),
        workbook=workbook,
    )
    (out_dir / EXPORT_META_FILENAME).write_text(
        json.dumps(outcome.as_json(), indent=2), encoding="utf-8"
    )
    return outcome
