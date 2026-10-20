"""Map complaints to existing Jira defects and RCAs, one sub-cluster (or batch) at a time.

Python is deterministic here: it short-circuits cases that are already linked to an issue by
CRM case number, narrows the candidate issues by software version and lexical similarity,
builds the request files, validates responses and writes the final mapping. A Copilot worker
agent does the reasoning.
"""

from __future__ import annotations

import json
import math
import re
from collections import Counter
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from pydantic import BaseModel, ValidationError

from ...prompts import get_analysis
from ...prompts.defect_mapping.schema import MappingResult
from . import jira, kb

MANIFEST = "manifest.json"
REQUESTS_DIR = "requests"
RESPONSES_DIR = "responses"
DIRECT_LINKS_FILE = "direct_links.jsonl"
MAPPINGS_FILE = "mappings.jsonl"
WORKBOOK_FILE = "mappings.xlsx"

SUMMARY_COLUMNS = (
    "case_summary.summary",
    "case_summary.reported_issue",
    "case_summary.product",
    "case_summary.actions_taken",
    "case_summary.resolution",
    "case_summary.status",
    "case_summary.key_points",
)
VERSION_COLUMN = "HRCSVC_Software_Version__r.Name"
ID_COLUMN = "CaseNumber"
CASE_SEPARATOR = "=== Case ==="
ISSUE_SEPARATOR = "=== Candidate ==="
ARTICLE_SEPARATOR = "=== KB Article ==="

DEFAULT_MIN_BATCH = 10
DEFAULT_MAX_BATCH = 20
DEFAULT_CANDIDATES = 40
DEFAULT_PER_CASE = 10

ISSUE_TEXT_FIELDS = (
    "summary",
    "description",
    "comment_summary",
    "impact_description",
    "steps_taken",
    "actual_results",
    "expected_results",
    "work_around",
    "root_cause_summary",
    "investigation_details",
    "troubleshooting_steps",
    "troubleshooting_results",
    "verification_summary",
    "rca_review_comments",
)


class MappingError(Exception):
    """Raised for unusable inputs, requests or responses."""


# -- helpers ---------------------------------------------------------------------------------


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def _read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise MappingError(f"file not found: {path}") from None
    except json.JSONDecodeError as exc:
        raise MappingError(f"{path} is not valid JSON: {exc}") from None


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except FileNotFoundError:
        raise MappingError(f"file not found: {path}") from None
    return [json.loads(line) for line in lines if line.strip()]


def latest_version(name: str) -> str:
    analysis = get_analysis(name)
    return sorted(analysis.prompts, key=lambda version: int(version[1:]))[-1]


# -- batching ---------------------------------------------------------------------------------


def split_sizes(count: int, *, min_batch: int, max_batch: int) -> list[int]:
    """Batch sizes for one sub-cluster, kept inside ``[min_batch, max_batch]`` where possible.

    A sub-cluster no larger than ``max_batch`` stays whole (even when below ``min_batch``);
    otherwise it is split into as few balanced batches as the bounds allow, so 23 cases with
    bounds 10-20 become 12 + 11.
    """
    if min_batch < 1 or max_batch < min_batch:
        raise MappingError("--min-batch must be >= 1 and <= --max-batch")
    if count <= max_batch:
        return [count]
    batches = math.ceil(count / max_batch)
    while batches > 1 and batches * min_batch > count:
        batches -= 1
    base, remainder = divmod(count, batches)
    return [base + (1 if index < remainder else 0) for index in range(batches)]


# -- similarity -------------------------------------------------------------------------------

_WORD = re.compile(r"[a-z0-9][a-z0-9_.-]{2,}")
_STOP = frozenset(
    """the and for with that this from have has was were are not but you your they them
    their there when what which while would could should about into over under then than
    issue case customer site system device unit room staff call user users report reported
    reports""".split()
)


def tokenize(text: str) -> list[str]:
    return [word for word in _WORD.findall(text.lower()) if word not in _STOP]


class Similarity:
    """Small TF-IDF cosine index over the candidate issues (no extra dependencies)."""

    def __init__(self, documents: Sequence[tuple[str, str]]) -> None:
        self.ids = [doc_id for doc_id, _ in documents]
        counts = [Counter(tokenize(text)) for _, text in documents]
        frequency: Counter[str] = Counter()
        for count in counts:
            frequency.update(count.keys())
        total = max(len(counts), 1)
        self.idf = {term: math.log(1 + total / (1 + df)) for term, df in frequency.items()}
        self.vectors = [self._vector(count) for count in counts]

    def _vector(self, count: Counter[str]) -> dict[str, float]:
        vector = {term: (1 + math.log(n)) * self.idf.get(term, 0.0) for term, n in count.items()}
        norm = math.sqrt(sum(value * value for value in vector.values())) or 1.0
        return {term: value / norm for term, value in vector.items()}

    def rank(self, query: str, *, allowed: set[str] | None = None) -> list[tuple[str, float]]:
        vector = self._vector(Counter(tokenize(query)))
        scored = []
        for doc_id, document in zip(self.ids, self.vectors):
            if allowed is not None and doc_id not in allowed:
                continue
            score = sum(weight * document.get(term, 0.0) for term, weight in vector.items())
            if score > 0:
                scored.append((doc_id, round(score, 4)))
        scored.sort(key=lambda item: (-item[1], item[0]))
        return scored


# -- inputs -----------------------------------------------------------------------------------


def read_cases(workbook: Path, *, case_numbers: Sequence[str]) -> dict[str, dict[str, Any]]:
    """Summary columns and software version for the selected cases."""
    if not workbook.is_file():
        raise MappingError(f"workbook not found: {workbook}")
    wanted = set(case_numbers)
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        sheet = book[book.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = [_text(cell) for cell in next(rows, ())]
        index = {name: position for position, name in enumerate(header) if name}
        missing = [name for name in (ID_COLUMN, *SUMMARY_COLUMNS) if name not in index]
        if missing:
            raise MappingError(f"workbook is missing column(s): {', '.join(missing)}")
        cases: dict[str, dict[str, Any]] = {}
        for row in rows:
            case_number = _text(row[index[ID_COLUMN]])
            if case_number not in wanted or case_number in cases:
                continue
            values = {column: _text(row[index[column]]) for column in SUMMARY_COLUMNS}
            version_text = _text(row[index[VERSION_COLUMN]]) if VERSION_COLUMN in index else ""
            cases[case_number] = {
                "case_number": case_number,
                "values": values,
                "software_version": version_text,
                "version": list(jira.parse_version(version_text) or ()),
            }
    finally:
        book.close()
    unknown = [number for number in case_numbers if number not in cases]
    if unknown:
        raise MappingError(f"{len(unknown)} case(s) not found in {workbook}: {', '.join(unknown[:10])}")
    return cases


def load_corpus(directories: Sequence[Path]) -> list[dict[str, Any]]:
    issues: dict[str, dict[str, Any]] = {}
    for directory in directories:
        for issue in jira.read_issues(directory):
            issues[issue["issue_key"]] = issue
    if not issues:
        raise MappingError("the Jira corpus is empty")
    return list(issues.values())


# -- version rule -------------------------------------------------------------------------------


def version_match(case_version: Sequence[int], issue: dict[str, Any]) -> tuple[bool, str]:
    """First-pass filter: same reported version, or a fix in a later version.

    An issue with no usable version stays a candidate, flagged so the model knows the version
    evidence is missing (most RCAs have no version fields).
    """
    reported = [tuple(version) for version in issue.get("versions_reported") or []]
    fixed = [tuple(version) for version in issue.get("versions_fix") or []]
    if not case_version:
        return True, "case_version_unknown"
    case = tuple(case_version)
    if not reported and not fixed:
        return True, "issue_version_unknown"
    if case in reported:
        return True, "reported_in_same_version"
    if any(version > case for version in fixed):
        return True, "fixed_in_later_version"
    return False, "version_excluded"


def select_candidates(
    cases: Sequence[dict[str, Any]],
    issues: Sequence[dict[str, Any]],
    index: Similarity,
    *,
    per_case: int,
    limit: int,
) -> list[dict[str, Any]]:
    """Version-filtered, similarity-ranked candidates for one batch of cases."""
    by_key = {issue["issue_key"]: issue for issue in issues}
    reasons: dict[str, str] = {}
    best: dict[str, float] = {}
    for case in cases:
        allowed = {}
        for issue in issues:
            keep, reason = version_match(case["version"], issue)
            if keep:
                allowed[issue["issue_key"]] = reason
        query = " ".join(case["values"][column] for column in SUMMARY_COLUMNS)
        for key, score in index.rank(query, allowed=set(allowed))[:per_case]:
            if score > best.get(key, 0.0):
                best[key] = score
                reasons[key] = allowed[key]
    ordered = sorted(best, key=lambda key: (-best[key], key))[:limit]
    return [{"issue": by_key[key], "score": best[key], "version_reason": reasons[key]} for key in ordered]


# -- rendering ---------------------------------------------------------------------------------


def render_case(case: dict[str, Any], kb_keys: Sequence[str]) -> str:
    lines = [CASE_SEPARATOR, f"CaseNumber: {case['case_number']}"]
    if case["software_version"]:
        lines.append(f"Reported software version: {case['software_version']}")
    for column in SUMMARY_COLUMNS:
        value = case["values"].get(column, "")
        if value:
            lines.append(f"{column}: {value}")
    lines.append(f"KB article ids: {', '.join(kb_keys) if kb_keys else '(none)'}")
    return "\n".join(lines)


def render_issue(candidate: dict[str, Any], *, max_chars: int = 2500) -> str:
    issue = candidate["issue"]
    lines = [
        ISSUE_SEPARATOR,
        f"issue_key: {issue['issue_key']}",
        f"issue_type: {'rca' if issue['project'] == 'RCA' else 'defect'}",
        f"status: {issue['status']}",
        f"version_reported: {issue.get('version_reported') or '(none)'}",
        f"version_fix: {issue.get('version_fix') or '(none)'}",
        f"version_rule: {candidate['version_reason']}",
    ]
    for field in ISSUE_TEXT_FIELDS:
        value = _text(issue.get(field))
        if value and value.upper() != "TBD":
            lines.append(f"{field}: {value[:max_chars]}")
    return "\n".join(lines)


def render_article(article: dict[str, Any], *, max_chars: int = 6000) -> str:
    return "\n".join(
        [
            ARTICLE_SEPARATOR,
            f"kb_article_id: {article['key']}",
            f"title: {article.get('title') or '(untitled)'}",
            f"text: {_text(article.get('text'))[:max_chars]}",
        ]
    )


# -- prepare -----------------------------------------------------------------------------------


def prepare(
    assignments_path: Path,
    out_dir: Path,
    *,
    workbook: Path,
    corpora: Sequence[Path],
    kb_dir: Path | None = None,
    min_batch: int = DEFAULT_MIN_BATCH,
    max_batch: int = DEFAULT_MAX_BATCH,
    candidates: int = DEFAULT_CANDIDATES,
    per_case: int = DEFAULT_PER_CASE,
    prompt_version: str | None = None,
) -> dict[str, Any]:
    """Build one request per sub-cluster batch, after short-circuiting direct links."""
    if (out_dir / MANIFEST).exists():
        raise MappingError(f"{out_dir} is already prepared; resume it or choose a new directory")
    assignments = _read_jsonl(assignments_path)
    if not assignments:
        raise MappingError(f"no case assignments in {assignments_path}")
    case_numbers = [row["case_number"] for row in assignments]
    cases = read_cases(workbook, case_numbers=case_numbers)
    issues = load_corpus(corpora)
    analysis = get_analysis("defect_mapping")
    version = prompt_version or latest_version("defect_mapping")

    linked: dict[str, list[str]] = {}
    for issue in issues:
        for number in issue["case_numbers"]:
            linked.setdefault(number, []).append(issue["issue_key"])
    direct = [
        {
            "case_number": number,
            "match_type": "direct_link",
            "issue_keys": sorted(linked[number]),
            "rationale": "The Jira CRM Case Number field already names this case.",
        }
        for number in case_numbers
        if number in linked
    ]
    _write_jsonl(out_dir / DIRECT_LINKS_FILE, direct)
    direct_numbers = {row["case_number"] for row in direct}

    kb_index = kb.read_index(kb_dir) if kb_dir else {}
    articles = kb.read_articles(kb_dir) if kb_dir else {}
    index = Similarity([(issue["issue_key"], _issue_text(issue)) for issue in issues])

    groups: dict[tuple[str, str], list[str]] = {}
    for row in assignments:
        if row["case_number"] in direct_numbers:
            continue
        key = (row.get("final_cluster", ""), row.get("final_subcluster", ""))
        groups.setdefault(key, []).append(row["case_number"])

    requests: list[dict[str, Any]] = []
    for (cluster, subcluster), numbers in sorted(groups.items()):
        sizes = split_sizes(len(numbers), min_batch=min_batch, max_batch=max_batch)
        start = 0
        for position, size in enumerate(sizes, start=1):
            chunk = numbers[start : start + size]
            start += size
            batch_cases = [cases[number] for number in chunk]
            case_articles: dict[str, list[str]] = {}
            for case in batch_cases:
                keys = [link["key"] for link in kb_index.get(case["case_number"], []) if link.get("key")]
                case_articles[case["case_number"]] = [key for key in keys if key in articles]
            used = sorted({key for keys in case_articles.values() for key in keys})
            selected = select_candidates(
                batch_cases, issues, index, per_case=per_case, limit=candidates
            )
            request_id = f"{_slug(subcluster) or 'unlabeled'}_b{position:02d}"
            user = "\n\n".join(
                [
                    f"Cluster: {cluster}\nSub-cluster: {subcluster}",
                    f"--- Complaints ({len(batch_cases)}) ---",
                    *(render_case(case, case_articles[case["case_number"]]) for case in batch_cases),
                    f"--- KB articles ({len(used)}) ---",
                    *(render_article(articles[key]) for key in used),
                    f"--- Candidate defects and RCAs ({len(selected)}) ---",
                    *(render_issue(candidate) for candidate in selected),
                ]
            )
            _write_json(
                out_dir / REQUESTS_DIR / f"{request_id}.json",
                {
                    "request_id": request_id,
                    "cluster": cluster,
                    "subcluster": subcluster,
                    "prompt": "defect_mapping",
                    "prompt_version": version,
                    "prompt_hash": analysis.prompt_hash(version),
                    "system": analysis.system_prompt + "\n\n" + analysis.prompt(version),
                    "user": user,
                    "case_numbers": chunk,
                    "candidate_keys": [item["issue"]["issue_key"] for item in selected],
                    "kb_article_ids": used,
                },
            )
            requests.append(
                {
                    "request_id": request_id,
                    "cluster": cluster,
                    "subcluster": subcluster,
                    "case_count": len(chunk),
                    "candidate_count": len(selected),
                    "kb_article_count": len(used),
                }
            )

    manifest = {
        "assignments": str(assignments_path),
        "workbook": str(workbook),
        "corpora": [str(path) for path in corpora],
        "kb_dir": str(kb_dir) if kb_dir else None,
        "prompt_version": version,
        "min_batch": min_batch,
        "max_batch": max_batch,
        "candidates": candidates,
        "per_case": per_case,
        "case_count": len(case_numbers),
        "direct_link_count": len(direct),
        "mapped_case_count": len(case_numbers) - len(direct),
        "subcluster_count": len(groups),
        "issue_count": len(issues),
        "requests": requests,
    }
    _write_json(out_dir / MANIFEST, manifest)
    return manifest


def _issue_text(issue: dict[str, Any]) -> str:
    return " ".join(_text(issue.get(field)) for field in ISSUE_TEXT_FIELDS)


def _slug(value: str) -> str:
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", value.lower())).strip("_")[:40]


def _write_jsonl(path: Path, rows: Sequence[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


# -- validate ------------------------------------------------------------------------------------


def validate(request: dict[str, Any], data: Any) -> tuple[MappingResult | None, list[str]]:
    try:
        result = MappingResult.model_validate(data)
    except ValidationError as exc:
        return None, [f"{'.'.join(str(p) for p in e['loc'])}: {e['msg']}" for e in exc.errors()[:20]]
    errors: list[str] = []
    expected = list(request["case_numbers"])
    allowed_issues = set(request["candidate_keys"])
    allowed_articles = set(request["kb_article_ids"])
    seen: Counter[str] = Counter(item.case_number for item in result.cases)
    unknown = sorted(set(seen) - set(expected))
    missing = [number for number in expected if number not in seen]
    repeated = sorted(number for number, count in seen.items() if count > 1)
    if unknown:
        errors.append(f"unknown case numbers: {', '.join(unknown[:20])}")
    if missing:
        errors.append(f"missing case numbers ({len(missing)}): {', '.join(missing[:20])}")
    if repeated:
        errors.append(f"case numbers returned more than once: {', '.join(repeated[:20])}")
    for item in result.cases:
        if item.match_type == "none" and item.matches:
            errors.append(f"{item.case_number}: match_type 'none' must have no matches")
        if item.match_type != "none" and not item.matches:
            errors.append(f"{item.case_number}: match_type '{item.match_type}' needs at least one match")
        for match in item.matches:
            if match.issue_key not in allowed_issues:
                errors.append(f"{item.case_number}: {match.issue_key} was not a candidate")
            expected_type = "rca" if match.issue_key.startswith("RCA-") else "defect"
            if match.issue_type != expected_type:
                errors.append(
                    f"{item.case_number}: {match.issue_key} is a {expected_type}, not {match.issue_type}"
                )
            for article in match.kb_support:
                if article not in allowed_articles:
                    errors.append(f"{item.case_number}: KB article {article} was not supplied")
    return (None if errors else result), errors


def _load_candidate(path: Path) -> Any:
    text = path.read_text(encoding="utf-8").strip()
    fenced = re.match(r"^```(?:json)?\s*(.*?)\s*```$", text, re.DOTALL)
    if fenced:
        text = fenced.group(1)
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise MappingError(f"candidate is not valid JSON: {exc}") from None


def publish_response(directory: Path, request_id: str, candidate: Path) -> list[str]:
    """Validate a candidate response and store it when valid. Returns the errors found."""
    if not re.fullmatch(r"[A-Za-z0-9_]+", request_id):
        raise MappingError(f"invalid request id '{request_id}'")
    request = _read_json(directory / REQUESTS_DIR / f"{request_id}.json")
    result, errors = validate(request, _load_candidate(candidate))
    if result is not None:
        _write_json(directory / RESPONSES_DIR / f"{request_id}.json", result.model_dump())
    return errors


def check(directory: Path) -> dict[str, Any]:
    manifest = _read_json(directory / MANIFEST)
    incomplete: list[str] = []
    errors: dict[str, list[str]] = {}
    for entry in manifest["requests"]:
        request_id = entry["request_id"]
        response = directory / RESPONSES_DIR / f"{request_id}.json"
        if not response.is_file():
            incomplete.append(request_id)
            continue
        request = _read_json(directory / REQUESTS_DIR / f"{request_id}.json")
        _, found = validate(request, _read_json(response))
        if found:
            incomplete.append(request_id)
            errors[request_id] = found
    return {"ready": not incomplete, "incomplete_request_ids": incomplete, "errors": errors}


def candidate_keys(directory: Path) -> dict[str, Any]:
    """Every issue key offered as a candidate in a prepared run.

    Used to fetch and summarize comments for just the issues that can still win a match,
    instead of the whole corpus.
    """
    manifest = _read_json(directory / MANIFEST)
    keys: set[str] = set()
    for entry in manifest["requests"]:
        request = _read_json(directory / REQUESTS_DIR / f"{entry['request_id']}.json")
        keys.update(request["candidate_keys"])
    ordered = sorted(keys)
    return {"count": len(ordered), "issue_keys": ordered}


# -- finalize ------------------------------------------------------------------------------------

MAPPING_COLUMNS = (
    "case_number",
    "cluster",
    "subcluster",
    "match_type",
    "issue_key",
    "issue_type",
    "confidence",
    "rationale",
    "version_justification",
    "kb_support",
)


def finalize(directory: Path) -> dict[str, Any]:
    """Combine direct links and worker responses into ``mappings.jsonl`` / ``mappings.xlsx``."""
    manifest = _read_json(directory / MANIFEST)
    status = check(directory)
    if not status["ready"]:
        raise MappingError(
            f"{len(status['incomplete_request_ids'])} request(s) still lack a valid response"
        )
    rows: list[dict[str, Any]] = []
    for link in _read_jsonl(directory / DIRECT_LINKS_FILE):
        rows.append(
            {
                "case_number": link["case_number"],
                "cluster": "",
                "subcluster": "",
                "match_type": "direct_link",
                "issue_key": ", ".join(link["issue_keys"]),
                "issue_type": "direct",
                "confidence": "",
                "rationale": link["rationale"],
                "version_justification": "",
                "kb_support": "",
            }
        )
    for entry in manifest["requests"]:
        response = _read_json(directory / RESPONSES_DIR / f"{entry['request_id']}.json")
        for case in response["cases"]:
            if not case["matches"]:
                rows.append(
                    {
                        "case_number": case["case_number"],
                        "cluster": entry["cluster"],
                        "subcluster": entry["subcluster"],
                        "match_type": "none",
                        "issue_key": "",
                        "issue_type": "",
                        "confidence": "",
                        "rationale": case.get("rationale", ""),
                        "version_justification": "",
                        "kb_support": "",
                    }
                )
            for match in case["matches"]:
                rows.append(
                    {
                        "case_number": case["case_number"],
                        "cluster": entry["cluster"],
                        "subcluster": entry["subcluster"],
                        "match_type": case["match_type"],
                        "issue_key": match["issue_key"],
                        "issue_type": match["issue_type"],
                        "confidence": match["confidence"],
                        "rationale": match["rationale"],
                        "version_justification": match.get("version_justification", ""),
                        "kb_support": ", ".join(match.get("kb_support", [])),
                    }
                )
    rows.sort(key=lambda row: (row["case_number"], row["issue_key"]))
    _write_jsonl(directory / MAPPINGS_FILE, rows)
    _write_workbook(directory / WORKBOOK_FILE, rows, manifest)

    mapped = {row["case_number"] for row in rows if row["issue_key"]}
    counts = Counter(row["match_type"] for row in rows)
    summary = {
        "case_count": manifest["case_count"],
        "direct_links": manifest["direct_link_count"],
        "cases_with_a_match": len(mapped),
        "cases_without_a_match": manifest["case_count"] - len(mapped),
        "match_type_counts": dict(sorted(counts.items())),
        "top_issues": [
            {"issue_key": key, "cases": count}
            for key, count in Counter(
                row["issue_key"] for row in rows if row["issue_key"] and row["match_type"] != "direct_link"
            ).most_common(10)
        ],
        "mappings": str(directory / MAPPINGS_FILE),
        "workbook": str(directory / WORKBOOK_FILE),
    }
    _write_json(directory / "summary.json", summary)
    return summary


def _write_workbook(path: Path, rows: Sequence[dict[str, Any]], manifest: dict[str, Any]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Mappings"
    sheet.append(list(MAPPING_COLUMNS))
    for row in rows:
        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", str(row.get(name, "")))[:32000] for name in MAPPING_COLUMNS])
    batches = book.create_sheet("Batches")
    batches.append(["request_id", "cluster", "subcluster", "case_count", "candidate_count", "kb_article_count"])
    for entry in manifest["requests"]:
        batches.append(
            [
                entry["request_id"],
                entry["cluster"],
                entry["subcluster"],
                entry["case_count"],
                entry["candidate_count"],
                entry["kb_article_count"],
            ]
        )
    book.save(path)
