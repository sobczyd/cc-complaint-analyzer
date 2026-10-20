"""Reusable corpus of Confluence KB articles referenced by complaint qualifications.

The link lives in ``HRCSVC21_URL__c`` on the cases workbook's *Complaint Qualifications*
sheet; most are Confluence tiny links (``/wiki/x/<key>``). ``plan`` lists the unique articles
to fetch and the case index, an agent reads each page once with the Confluence MCP tools and
saves it under ``pages\\``, and ``normalize`` writes ``articles.jsonl``.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Sequence

from openpyxl import load_workbook

MANIFEST = "manifest.json"
PAGES_DIR = "pages"
INDEX_FILE = "index.jsonl"
ARTICLES_FILE = "articles.jsonl"
CQ_SHEET = "Complaint Qualifications"
URL_COLUMN = "HRCSVC21_URL__c"
ID_COLUMN = "CaseNumber"

ACCESSIBLE_HOSTS = ("baxter-international.atlassian.net",)
_TINY = re.compile(r"https?://([^/]+)/wiki/x/([A-Za-z0-9_-]+)")
_PAGE_ID = re.compile(r"https?://([^/]+)/wiki/spaces/[^/]+/pages/(\d+)")


class KbCorpusError(Exception):
    """Raised for an unusable workbook, corpus directory or page file."""


def parse_link(value: str) -> dict[str, str] | None:
    """A Confluence tiny link or page URL as ``{key, host, url, kind}``; ``None`` otherwise."""
    text = (value or "").strip()
    if not text.lower().startswith("http"):
        return None
    for pattern, kind in ((_TINY, "tiny"), (_PAGE_ID, "page_id")):
        match = pattern.search(text)
        if match:
            host, key = match.group(1), match.group(2)
            return {"key": key, "host": host, "url": match.group(0), "kind": kind}
    return None


def read_case_links(workbook: Path, *, case_numbers: Sequence[str] | None = None) -> list[dict[str, Any]]:
    """One row per (case, KB link) from the Complaint Qualifications sheet.

    Values that are not links (for example ``No KB Required``) are kept with ``key`` empty and
    a reason, so the mapping step can say why a case has no article.
    """
    if not workbook.is_file():
        raise KbCorpusError(f"workbook not found: {workbook}")
    wanted = set(case_numbers) if case_numbers else None
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        if CQ_SHEET not in book.sheetnames:
            raise KbCorpusError(f"workbook has no '{CQ_SHEET}' sheet: {workbook}")
        sheet = book[CQ_SHEET]
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell) if cell is not None else "" for cell in next(rows, ())]
        for column in (ID_COLUMN, URL_COLUMN):
            if column not in header:
                raise KbCorpusError(f"'{CQ_SHEET}' is missing column {column}")
        id_index, url_index = header.index(ID_COLUMN), header.index(URL_COLUMN)
        links: list[dict[str, Any]] = []
        for row in rows:
            case_number = str(row[id_index] or "").strip()
            raw = str(row[url_index] or "").strip()
            if not case_number or not raw:
                continue
            if wanted is not None and case_number not in wanted:
                continue
            parsed = parse_link(raw)
            if parsed is None:
                links.append({"case_number": case_number, "key": "", "raw": raw, "reason": "not_a_link"})
            elif parsed["host"] not in ACCESSIBLE_HOSTS:
                links.append(
                    {
                        "case_number": case_number,
                        "key": "",
                        "raw": raw,
                        "host": parsed["host"],
                        "reason": "site_not_accessible",
                    }
                )
            else:
                links.append({"case_number": case_number, **parsed, "raw": raw})
        return links
    finally:
        book.close()


def plan(
    out_dir: Path, *, workbook: Path, case_numbers: Sequence[str] | None = None
) -> dict[str, Any]:
    """Record which articles to fetch and which case cites which article."""
    links = read_case_links(workbook, case_numbers=case_numbers)
    articles: dict[str, dict[str, str]] = {}
    for link in links:
        if link["key"] and link["key"] not in articles:
            articles[link["key"]] = {"key": link["key"], "host": link["host"], "url": link["url"]}
    (out_dir / PAGES_DIR).mkdir(parents=True, exist_ok=True)
    with (out_dir / INDEX_FILE).open("w", encoding="utf-8") as handle:
        for link in links:
            handle.write(json.dumps(link, ensure_ascii=False) + "\n")
    manifest = {
        "workbook": str(workbook),
        "case_count": len({link["case_number"] for link in links}),
        "link_count": len(links),
        "article_count": len(articles),
        "unavailable": sorted({link["reason"] for link in links if not link["key"]}),
        "articles": sorted(articles.values(), key=lambda item: item["key"]),
        "pages_dir": PAGES_DIR,
        "status": "planned",
    }
    (out_dir / MANIFEST).write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return {key: manifest[key] for key in ("case_count", "link_count", "article_count", "unavailable")}


def read_manifest(directory: Path) -> dict[str, Any]:
    path = directory / MANIFEST
    if not path.is_file():
        raise KbCorpusError(f"not a KB corpus directory (no {MANIFEST}): {directory}")
    return json.loads(path.read_text(encoding="utf-8"))


def _page_text(payload: Any) -> tuple[str, str]:
    """``(title, body)`` from a saved Confluence MCP result, whichever shape it has."""
    if isinstance(payload, str):
        return "", payload
    if not isinstance(payload, dict):
        return "", ""
    if "data" in payload:
        return _page_text(payload["data"])
    title = str(payload.get("title") or "")
    body = payload.get("body")
    if isinstance(body, dict):
        for key in ("value", "markdown", "html", "view", "storage"):
            candidate = body.get(key)
            if isinstance(candidate, str) and candidate.strip():
                return title, candidate
            if isinstance(candidate, dict) and isinstance(candidate.get("value"), str):
                return title, candidate["value"]
    for key in ("markdown", "content", "text", "excerpt"):
        candidate = payload.get(key)
        if isinstance(candidate, str) and candidate.strip():
            return title, candidate
    return title, ""


def normalize(directory: Path) -> dict[str, Any]:
    """Turn saved pages into ``articles.jsonl`` keyed by tiny-link key."""
    manifest = read_manifest(directory)
    pages_dir = directory / PAGES_DIR
    articles: list[dict[str, Any]] = []
    missing: list[str] = []
    for entry in manifest["articles"]:
        key = entry["key"]
        path = next(iter(sorted(pages_dir.glob(f"{key}.*"))), None)
        if path is None:
            missing.append(key)
            continue
        raw = path.read_text(encoding="utf-8")
        if path.suffix == ".json":
            try:
                title, body = _page_text(json.loads(raw))
            except json.JSONDecodeError as exc:
                raise KbCorpusError(f"{path} is not valid JSON: {exc}") from None
        else:
            title, body = "", raw
        articles.append({**entry, "title": title, "text": body.strip()})
    with (directory / ARTICLES_FILE).open("w", encoding="utf-8") as handle:
        for article in articles:
            handle.write(json.dumps(article, ensure_ascii=False) + "\n")
    manifest.update(status="normalized", fetched=len(articles), missing=missing)
    (directory / MANIFEST).write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return {"fetched": len(articles), "missing": missing, "article_count": manifest["article_count"]}


def read_articles(directory: Path) -> dict[str, dict[str, Any]]:
    path = directory / ARTICLES_FILE
    if not path.is_file():
        raise KbCorpusError(f"KB corpus is not normalized (no {ARTICLES_FILE}): {directory}")
    articles = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            article = json.loads(line)
            articles[article["key"]] = article
    return articles


def read_index(directory: Path) -> dict[str, list[dict[str, Any]]]:
    """Case number -> its KB links (including unavailable ones)."""
    path = directory / INDEX_FILE
    if not path.is_file():
        raise KbCorpusError(f"KB corpus has no {INDEX_FILE}: {directory}")
    index: dict[str, list[dict[str, Any]]] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            link = json.loads(line)
            index.setdefault(link["case_number"], []).append(link)
    return index
