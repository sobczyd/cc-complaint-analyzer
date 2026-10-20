"""Build a per-case "Case Context" text column from chosen workbook columns.

Each field becomes a ``ColumnName: value`` line (blank values are skipped). Columns are
looked up on the ``Cases`` sheet first, then on the child detail sheets joined by
``CaseNumber``; a child field with several rows yields one line per row.
"""

from __future__ import annotations

from pathlib import Path
from typing import Sequence

from openpyxl import Workbook, load_workbook

from ..analysis_run import AnalysisRunError
from ..salesforce_export.export import MAIN_SHEET

CONTEXT_WORKBOOK_FILENAME = "case_context.xlsx"
CONTEXT_COLUMN = "Case Context"


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _read_sheet(book, name: str) -> tuple[list[str], list[tuple]]:
    rows = book[name].iter_rows(values_only=True)
    header = [_text(cell) for cell in next(rows, ())]
    return header, [row for row in rows if any(cell is not None for cell in row)]


def build_context_workbook(
    run_dir: Path, fields: Sequence[str], *, id_field: str = "CaseNumber"
) -> Path:
    """Write ``case_context.xlsx`` (sheet ``Cases``: id column + ``Case Context``)."""
    source = run_dir / "cases.xlsx"
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        header, rows = _read_sheet(book, MAIN_SHEET)
        main = {name.lower(): index for index, name in enumerate(header)}
        if id_field.lower() not in main:
            raise AnalysisRunError(f"id field '{id_field}' is not a column of {source.name}")

        children: dict[str, tuple[dict[str, int], dict[str, list[tuple]]]] = {}
        for sheet in book.sheetnames:
            if sheet == MAIN_SHEET:
                continue
            child_header, child_rows = _read_sheet(book, sheet)
            columns = {name.lower(): index for index, name in enumerate(child_header)}
            if "casenumber" not in columns:
                continue
            by_case: dict[str, list[tuple]] = {}
            for row in child_rows:
                by_case.setdefault(_text(row[columns["casenumber"]]), []).append(row)
            children[sheet] = (columns, by_case)

        plan: list[tuple[str, str, str | None]] = []  # (label, source, child sheet)
        for field in fields:
            key = field.lower()
            if key in main:
                plan.append((header[main[key]], "main", None))
                continue
            sheet = next((name for name, (cols, _) in children.items() if key in cols), None)
            if sheet is None:
                raise AnalysisRunError(f"context field '{field}' is not in any sheet of {source.name}")
            plan.append((field, "child", sheet))

        id_index = main[id_field.lower()]
        output = Workbook()
        out_sheet = output.active
        out_sheet.title = MAIN_SHEET
        out_sheet.append([id_field, CONTEXT_COLUMN])
        for row in rows:
            case_id = _text(row[id_index])
            lines: list[str] = []
            for label, kind, sheet in plan:
                if kind == "main":
                    values = [_text(row[main[label.lower()]])]
                else:
                    columns, by_case = children[sheet]
                    index = columns[label.lower()]
                    values = [_text(child[index]) for child in by_case.get(case_id, [])]
                lines += [f"{label}: {value}" for value in values if value]
            out_sheet.append([case_id, "\n".join(lines)])
    finally:
        book.close()

    path = run_dir / CONTEXT_WORKBOOK_FILENAME
    output.save(path)
    return path
