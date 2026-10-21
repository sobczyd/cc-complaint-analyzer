"""Period-over-period comparison: period loading, mapping validation, metrics, flags and outputs."""

from __future__ import annotations

import json
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from complaint_analyzer.workflows.compare_periods import compare
from complaint_analyzer.workflows.compare_periods.__main__ import main as cli

QUERY = """SELECT Id FROM Case
WHERE HRCSVC_Complaint_Code__r.Name = 'CareComm - Nurse Call'
  AND CreatedDate >= {start}T00:00:00Z
  AND CreatedDate < {end}T00:00:00Z
  AND Id IN (SELECT HRCSVC_Case__c FROM X WHERE Name LIKE 'CQ-%')
ORDER BY CreatedDate ASC
"""


def make_run(
    root: Path,
    name: str,
    taxonomy: dict[str, dict[str, list[str]]],
    created: dict[str, date],
    *,
    query: str | None = None,
) -> Path:
    """A finalized run: taxonomy {cluster: {subcluster: [case numbers]}} plus case dates."""
    run = root / name
    cdir = run / "clusters" / "all"
    cdir.mkdir(parents=True)
    clusters = []
    rows = []
    for cluster, subs in taxonomy.items():
        clusters.append(
            {
                "label": cluster,
                "description": f"{cluster} d",
                "rationale": "r",
                "subclusters": [{"label": s, "description": f"{s} d", "rationale": "r"} for s in subs],
            }
        )
        for sub, cases in subs.items():
            rows += [
                {"case_number": c, "final_cluster": cluster, "final_subcluster": sub} for c in cases
            ]
    (cdir / "final_taxonomy.json").write_text(json.dumps({"case_count": len(rows), "clusters": clusters}))
    (cdir / "case_assignments.jsonl").write_text("\n".join(json.dumps(r) for r in rows) + "\n")
    (cdir / "manifest.json").write_text(json.dumps({"case_count": len(rows), "batch_size": 10}))
    (run / "export.json").write_text(json.dumps({"row_count": len(rows)}))
    (run / "request.txt").write_text("test run")
    (run / "query.soql").write_text(
        query or QUERY.format(start="2026-09-01", end="2026-10-01"), encoding="utf-8"
    )
    book = Workbook()
    ws = book.active
    ws.title = "Cases"
    ws.append(["CaseNumber", "CreatedDate"])
    for case, day in created.items():
        ws.append([case, f"{day.isoformat()}T10:00:00.000+0000"])
    book.save(run / "cases_with_clusters.xlsx")
    return run


def numbers(prefix: str, n: int) -> list[str]:
    return [f"{prefix}{i:03d}" for i in range(n)]


def dates(cases: list[str], start: date) -> dict[str, date]:
    return {c: start + timedelta(days=i % 10) for i, c in enumerate(cases)}


def spec(key: str, run: Path, **kwargs: Any) -> compare.PeriodSpec:
    return compare.PeriodSpec(
        key, run, kwargs.get("cluster_name"), kwargs.get("start"), kwargs.get("end"), kwargs.get("label")
    )


THRESHOLDS = {"min_change": 10, "min_cases": 10, "flag_basis": "count"}


def test_flag_rules() -> None:
    assert compare.flag_for(0, 3, 30, 30, THRESHOLDS) == "NEW"
    assert compare.flag_for(4, 0, 30, 30, THRESHOLDS) == "GONE"
    assert compare.flag_for(20, 31, 30, 30, THRESHOLDS) == "UPTICK"
    assert compare.flag_for(31, 20, 30, 30, THRESHOLDS) == "DECLINE"
    assert compare.flag_for(20, 30, 30, 30, THRESHOLDS) == ""  # change must exceed 10
    assert compare.flag_for(2, 14, 30, 30, THRESHOLDS) == "UPTICK"  # one period has >= 10
    assert compare.flag_for(2, 9, 30, 30, THRESHOLDS) == ""  # change under threshold
    assert compare.flag_for(5, 14, 30, 30, {**THRESHOLDS, "min_change": 8}) == "UPTICK"
    assert compare.flag_for(5, 14, 30, 30, {**THRESHOLDS, "min_cases": 20}) == ""


def test_flag_basis_per_day_scales_baseline() -> None:
    per_day = {**THRESHOLDS, "flag_basis": "per_day"}
    # 30 days vs 15 days: 40 cases in A is about 20 expected in B, so 22 is steady, not a decline of 18.
    assert compare.flag_for(40, 22, 30, 15, THRESHOLDS) == "DECLINE"
    assert compare.flag_for(40, 22, 30, 15, per_day) == ""
    assert compare.flag_for(40, 40, 30, 15, per_day) == "UPTICK"


def test_single_run_split_by_date_needs_no_mapping(tmp_path: Path) -> None:
    taxonomy = {
        "Network": {"PoE fault": numbers("P", 30), "Switch down": numbers("S", 5)},
        "Audio": {"No audio": numbers("N", 12)},
    }
    created: dict[str, date] = {}
    for cases in (numbers("P", 30), numbers("S", 5), numbers("N", 12)):
        created.update(dates(cases, date(2026, 9, 1)))
    for case in numbers("P", 30)[20:]:  # 10 PoE cases land in the second period
        created[case] = date(2026, 9, 20)
    for case in numbers("N", 12):
        created[case] = date(2026, 9, 2)
    run = make_run(tmp_path, "run", taxonomy, created)
    out = tmp_path / "cmp"
    result = compare.prepare(
        out,
        spec("A", run, start=date(2026, 9, 1), end=date(2026, 9, 14), label="First half"),
        spec("B", run, start=date(2026, 9, 15), end=date(2026, 9, 30), label="Second half"),
    )
    assert result["resolution_required"] is False
    assert result["request_count"] == 0
    assert result["days"] == {"A": 14, "B": 16}
    assert any("period lengths differ" in w for w in result["warnings"])
    assert compare.check(out)["ready"] is True

    summary = compare.finalize(out)
    comparison = json.loads((out / compare.COMPARISON_FILE).read_text())
    by_cluster = {c["cluster"]: c for c in comparison["clusters"]}
    assert by_cluster["Network"]["count_a"] + by_cluster["Network"]["count_b"] == 35
    assert by_cluster["Audio"]["count_b"] == 0 and by_cluster["Audio"]["flag"] == "GONE"
    assert sum(c["count_a"] for c in comparison["clusters"]) == summary["case_counts"]["A"]
    assert sum(c["count_b"] for c in comparison["clusters"]) == summary["case_counts"]["B"]
    for name in (
        compare.WORKBOOK_FILE, compare.SUMMARY_FILE, compare.RESOLVED_FILE, compare.MAPPING_FILE,
        compare.SOURCES_FILE, compare.COMPARISON_FILE,
    ):
        assert (out / name).is_file()
    resolved = [json.loads(line) for line in (out / compare.RESOLVED_FILE).read_text().splitlines()]
    assert all(r["original_cluster"] == r["resolved_cluster"] for r in resolved)
    assert {r["period"] for r in resolved} == {"A", "B"}


def two_runs(tmp_path: Path) -> tuple[Path, Path]:
    a_tax = {
        "Network faults": {"PoE or switch fault": numbers("a", 40), "Cabling": numbers("b", 20)},
        "Audio": {"No audio": numbers("c", 15)},
    }
    b_tax = {
        "Infrastructure problems": {"Power over Ethernet failure": numbers("x", 30)},
        "Audio": {"Silent calls": numbers("y", 12)},
        "Software": {"Upgrade failure": numbers("z", 14)},
    }
    a_cases = [c for subs in a_tax.values() for cs in subs.values() for c in cs]
    b_cases = [c for subs in b_tax.values() for cs in subs.values() for c in cs]
    a = make_run(tmp_path, "sep", a_tax, dates(a_cases, date(2026, 9, 1)))
    b = make_run(
        tmp_path, "oct", b_tax, dates(b_cases, date(2026, 10, 1)),
        query=QUERY.format(start="2026-10-01", end="2026-10-20"),
    )
    return a, b


def candidate(tmp_path: Path, mappings: list[dict[str, Any]]) -> Path:
    path = tmp_path / "candidate.json"
    path.write_text(json.dumps({"mappings": mappings}))
    return path


def test_two_runs_are_resolved_onto_the_baseline(tmp_path: Path) -> None:
    sep, octo = two_runs(tmp_path)
    out = tmp_path / "cmp"
    result = compare.prepare(
        out, spec("A", sep, label="Sep 2026"), spec("B", octo, label="Oct 2026"), batch_size=2
    )
    assert result["resolution_required"] is True
    assert result["b_subcluster_count"] == 3
    assert result["request_count"] == 2  # 3 sub-clusters in batches of 2
    assert not any("case filters" in w for w in result["warnings"])
    assert compare.check(out)["incomplete_request_ids"] == ["resolve_01", "resolve_02"]

    request = json.loads((out / "requests" / "resolve_01.json").read_text())
    assert request["system"] and "BASELINE" in request["user"]
    b_ids = request["subclusters"]
    assert len(b_ids) == 2
    baseline_ids = request["baseline_subclusters"]

    # An unmapped sub-cluster, an unknown id and a bad baseline id are all rejected.
    bad = candidate(tmp_path, [
        {"subcluster_id": b_ids[0], "decision": "match", "baseline_subcluster_id": "A9.S9"},
        {"subcluster_id": "B7.S7", "decision": "new", "new_cluster": "X", "new_subcluster": "Y"},
    ])
    errors = compare.publish_response(out, "resolve_01", bad)
    assert any("not a baseline sub-cluster" in e for e in errors)
    assert any("unknown sub-cluster ids" in e for e in errors)
    assert any("unmapped sub-clusters" in e for e in errors)
    assert not (out / "responses" / "resolve_01.json").exists()

    by_label = {
        s["label"]: s["subcluster_id"]
        for c in json.loads((out / compare.BASELINE_FILE).read_text())["clusters"]
        for s in c["subclusters"]
    }
    b_labels = {
        s["subcluster_id"]: s["label"]
        for c in json.loads((out / compare.PERIOD_B_FILE).read_text())["clusters"]
        for s in c["subclusters"]
    }
    answers = {
        "Power over Ethernet failure": {"decision": "match", "baseline_subcluster_id": by_label["PoE or switch fault"], "confidence": "high"},
        "Silent calls": {"decision": "match", "baseline_subcluster_id": by_label["No audio"], "confidence": "low"},
        "Upgrade failure": {"decision": "new", "new_cluster": "Software", "new_subcluster": "Upgrade failure", "confidence": "high"},
    }
    for request_id in ("resolve_01", "resolve_02"):
        ids = json.loads((out / "requests" / f"{request_id}.json").read_text())["subclusters"]
        path = candidate(tmp_path, [{"subcluster_id": i, **answers[b_labels[i]]} for i in ids])
        assert compare.publish_response(out, request_id, path) == []
    assert compare.check(out)["ready"] is True
    assert set(baseline_ids) >= set(by_label.values())

    compare.finalize(out)
    comparison = json.loads((out / compare.COMPARISON_FILE).read_text())
    clusters = {c["cluster"]: c for c in comparison["clusters"]}
    assert clusters["Network faults"]["count_a"] == 60 and clusters["Network faults"]["count_b"] == 30
    assert clusters["Audio"]["count_a"] == 15 and clusters["Audio"]["count_b"] == 12
    assert clusters["Software"]["flag"] == "NEW" and clusters["Software"]["count_a"] == 0
    subs = {(s["cluster"], s["subcluster"]): s for s in comparison["subclusters"]}
    assert subs[("Network faults", "Cabling")]["flag"] == "GONE"
    assert subs[("Network faults", "PoE or switch fault")]["b_source_labels"] == [
        "Infrastructure problems / Power over Ethernet failure"
    ]
    assert comparison["mapping_summary"]["low_confidence"] == 1
    assert comparison["mapping_summary"]["by_status"] == {"renamed": 2, "new": 1}
    assert comparison["totals"] == {"count_a": 75, "count_b": 56, "days_a": 10, "days_b": 10}

    resolved = [json.loads(line) for line in (out / compare.RESOLVED_FILE).read_text().splitlines()]
    poe = next(r for r in resolved if r["case_number"] == "x000")
    assert (poe["original_cluster"], poe["resolved_cluster"], poe["mapping_status"]) == (
        "Infrastructure problems", "Network faults", "renamed"
    )
    sources = json.loads((out / compare.SOURCES_FILE).read_text())
    assert sources["periods"]["A"]["assignments_sha256"] and sources["prompt"]["name"] == "cluster_resolve"

    book = load_workbook(out / compare.WORKBOOK_FILE)
    assert book.sheetnames == [
        "Summary", "Clusters", "Subclusters", "Case Assignments", "Mapping", "Sources", "Legend"
    ]
    assert book["Case Assignments"].max_row == 1 + 75 + 56
    assert "Software" in (out / compare.SUMMARY_FILE).read_text()


def test_filter_mismatch_and_unsupported_type_are_reported(tmp_path: Path) -> None:
    sep, octo = two_runs(tmp_path)
    query = (octo / "query.soql").read_text().replace("Nurse Call", "Other Product")
    (octo / "query.soql").write_text(query)
    out = tmp_path / "cmp"
    result = compare.prepare(out, spec("A", sep), spec("B", octo))
    assert any("different case filters" in w for w in result["warnings"])
    with pytest.raises(compare.CompareError, match="not supported yet"):
        compare.prepare(tmp_path / "other", spec("A", sep), spec("B", octo), comparison_type="by_account")
    with pytest.raises(compare.CompareError, match="already prepared"):
        compare.prepare(out, spec("A", sep), spec("B", octo))


def test_partial_cluster_directory_is_reported(tmp_path: Path) -> None:
    sep, octo = two_runs(tmp_path)
    (octo / "export.json").write_text(json.dumps({"row_count": 500}))
    result = compare.prepare(tmp_path / "cmp", spec("A", sep), spec("B", octo))
    assert any("covers 56 of 500" in w for w in result["warnings"])


def test_cli_list_clusters_and_prepare(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    sep, octo = two_runs(tmp_path)
    assert cli(["list-clusters", "--run", str(sep)]) == 0
    assert json.loads(capsys.readouterr().out)["suggested"] == "all"
    out = tmp_path / "cmp"
    assert cli(["prepare", "--out-dir", str(out), "--a-run", str(sep), "--b-run", str(octo)]) == 0
    assert (out / compare.MANIFEST).is_file()
    assert cli(["check", "--dir", str(out)]) == 1  # resolution requests still unanswered
    assert cli(["finalize", "--dir", str(out)]) == 2  # cannot finalize without responses
