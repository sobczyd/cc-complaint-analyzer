import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from complaint_analyzer.workflows.report_merge.__main__ import main


def _setup(tmp_path: Path) -> tuple[Path, Path]:
    cases = tmp_path / "cases_with_summaries.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Cases"
    sheet.append(["CaseNumber", "case_summary.summary"])
    sheet.append(["001", "a"])
    sheet.append(["002", "b"])
    book.save(cases)
    cdir = tmp_path / "clusters" / "x"
    cdir.mkdir(parents=True)
    (cdir / "case_assignments.jsonl").write_text(
        json.dumps({"case_number": "001", "final_cluster": "Hardware", "final_subcluster": "Station replaced"}) + "\n",
        encoding="utf-8",
    )
    return cases, cdir


def test_merge_clusters_writes_default_output(tmp_path: Path) -> None:
    cases, cdir = _setup(tmp_path)
    assert main(["merge-clusters", "--cases", str(cases), "--cluster-dir", str(cdir)]) == 0
    out = tmp_path / "cases_with_clusters.xlsx"
    rows = list(load_workbook(out)["Cases"].iter_rows(values_only=True))
    header = list(rows[0])
    assert header[:2] == ["CaseNumber", "case_summary.summary"]
    assert "case_cluster.cluster" in header and "case_cluster.subcluster" in header
    first = dict(zip(header, rows[1]))
    assert first["case_cluster.cluster"] == "Hardware"
    assert first["case_cluster.subcluster"] == "Station replaced"
    second = dict(zip(header, rows[2]))
    assert second["case_cluster.analysis_status"] == "not_run"
    assert "Legend" in load_workbook(out).sheetnames


def test_merge_clusters_missing_assignments_is_error(tmp_path: Path) -> None:
    cases, _ = _setup(tmp_path)
    empty = tmp_path / "empty"
    empty.mkdir()
    assert main(["merge-clusters", "--cases", str(cases), "--cluster-dir", str(empty)]) == 2
