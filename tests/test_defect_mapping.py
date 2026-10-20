import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from complaint_analyzer.prompts import get_analysis
from complaint_analyzer.workflows.defect_mapping import jira, kb, mapping
from complaint_analyzer.workflows.defect_mapping.__main__ import main

ROOT = Path(__file__).resolve().parents[1]


# -- jira corpus ------------------------------------------------------------------


def test_as_text_flattens_html_options_and_lists() -> None:
    assert jira.as_text("<p>one</p><p>two</p>") == "one\ntwo"
    assert jira.as_text({"name": "NC 4.1.200"}) == "NC 4.1.200"
    assert jira.as_text([{"name": "a"}, {"value": "b"}]) == "a, b"
    assert jira.as_text(None) == ""


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("Voalte Nurse Call 4.0.405", (4, 0, 405)),
        ("NC 4.1.200", (4, 1, 200)),
        ("PRJ01099 - NNC 4.2.000", (4, 2, 0)),
        ("", None),
        ("no version here", None),
    ],
)
def test_parse_version(value: str, expected: tuple[int, int, int] | None) -> None:
    assert jira.parse_version(value) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("00979158 00980526", ["00979158", "00980526"]),
        ("00543751 & 00543522", ["00543751", "00543522"]),
        ("01074804/01083150", ["01074804", "01083150"]),
        ("00218496 - Service Cloud", ["00218496"]),
        ("1067661", ["01067661"]),
        ("", []),
    ],
)
def test_parse_case_numbers(value: str, expected: list[str]) -> None:
    assert jira.parse_case_numbers(value) == expected


def _raw_issue(key: str, **fields: object) -> dict:
    base = {
        "summary": "Console drops calls",
        "description": "<p>Calls drop</p>",
        "status": {"name": "Done"},
        "issuetype": {"name": "Defect_CC"},
        "created": "2026-01-01T00:00:00.000-0500",
        "updated": "2026-02-01T00:00:00.000-0500",
        "fixVersions": [{"name": "NC 4.2.000"}],
        "versions": [{"name": "NC 4.0.405"}],
        "customfield_14726": [{"name": "NC 4.0.405"}],
        "customfield_10693": "00123456",
        "comment": {"comments": [{"body": "<p>looked at logs</p>"}]},
    }
    return {"key": key, "fields": {**base, **fields}}


def _corpus(tmp_path: Path, issues: list[dict], *, kind: str = "defect") -> Path:
    directory = tmp_path / f"corpus_{kind}"
    jira.plan(directory, kind=kind)
    (directory / "raw" / "page_01.json").write_text(
        json.dumps({"data": {"issues": issues}}), encoding="utf-8"
    )
    return directory


def test_normalize_writes_issues_and_counts(tmp_path: Path) -> None:
    directory = _corpus(tmp_path, [_raw_issue("NC-1"), _raw_issue("NC-2", comment={"comments": []})])
    out = jira.normalize(directory)
    assert out == {
        "kind": "defect",
        "issue_count": 2,
        "with_comments": 1,
        "with_case_numbers": 2,
        "with_comment_summary": 0,
    }
    issues = jira.read_issues(directory)
    assert [issue["issue_key"] for issue in issues] == ["NC-1", "NC-2"]
    assert issues[0]["description"] == "Calls drop"
    assert issues[0]["versions_fix"] == [[4, 2, 0]]
    assert issues[0]["case_numbers"] == ["00123456"]
    assert (directory / jira.WORKBOOK_FILE).is_file()


def test_normalize_reads_concatenated_json_and_evidence_shape(tmp_path: Path) -> None:
    directory = tmp_path / "corpus"
    jira.plan(directory, kind="defect")
    evidence = {
        "key": "NC-9",
        "fields": {
            "summary": "s",
            "issuetype": {"name": "Defect_CC"},
            "customFields": {"CRM Case Number": {"id": "customfield_10693", "value": "00999999"}},
        },
    }
    (directory / "raw" / "page_01.json").write_text(
        json.dumps({"data": {"issues": [_raw_issue("NC-1")]}}) + "\n" + json.dumps({"data": evidence}),
        encoding="utf-8",
    )
    jira.normalize(directory)
    issues = {issue["issue_key"]: issue for issue in jira.read_issues(directory)}
    assert issues["NC-9"]["case_numbers"] == ["00999999"]
    assert set(issues) == {"NC-1", "NC-9"}


def test_normalize_merges_pages_fetched_for_different_fields(tmp_path: Path) -> None:
    directory = tmp_path / "corpus"
    jira.plan(directory, kind="defect")
    (directory / "raw" / "page_01.json").write_text(
        json.dumps({"data": {"issues": [_raw_issue("NC-1", comment={"total": 2})]}}), encoding="utf-8"
    )
    (directory / "raw" / "page_02_comments.json").write_text(
        json.dumps({"data": {"issues": [{"key": "NC-1", "fields": {}}]}})
        + "\n"
        + json.dumps(
            {"issues": {"nodes": [{"key": "NC-1", "fields": {"comment": {"comments": [{"body": "found it"}]}}}]}}
        ),
        encoding="utf-8",
    )
    jira.normalize(directory)
    issue = jira.read_issues(directory)[0]
    assert issue["summary"] == "Console drops calls"  # kept from the field page
    assert issue["comments"] == ["found it"]  # added by the comments page


def test_comment_dataset_and_summary_merge(tmp_path: Path) -> None:
    directory = _corpus(tmp_path, [_raw_issue("NC-1")])
    jira.normalize(directory)
    out = jira.comment_dataset(directory, tmp_path / "comments.xlsx")
    assert out["issue_count"] == 1
    rows = list(load_workbook(tmp_path / "comments.xlsx")["Issues"].iter_rows(values_only=True))
    assert rows[0] == ("issue_key", "Issue Context")
    assert "looked at logs" in rows[1][1]

    results = tmp_path / "results.jsonl"
    results.write_text(
        json.dumps(
            {"record_id": "NC-1", "status": "ok", "values": {"summary": "root cause found", "key_points": ["p1"]}}
        )
        + "\n",
        encoding="utf-8",
    )
    jira.normalize(directory, summaries=results)
    issue = jira.read_issues(directory)[0]
    assert issue["comment_summary"].startswith("root cause found")
    assert "- p1" in issue["comment_summary"]


def test_normalize_requires_raw_pages(tmp_path: Path) -> None:
    directory = tmp_path / "empty"
    jira.plan(directory, kind="rca")
    with pytest.raises(jira.JiraCorpusError, match="no issues found"):
        jira.normalize(directory)


# -- kb corpus ---------------------------------------------------------------------


def _cases_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "cases_with_summaries.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Cases"
    columns = ["CaseNumber", mapping.VERSION_COLUMN, *mapping.SUMMARY_COLUMNS]
    sheet.append(columns)
    for number in ("00000001", "00000002", "00000003"):
        sheet.append([number, "Voalte Nurse Call 4.0.405", *[f"text {number}"] * len(mapping.SUMMARY_COLUMNS)])
    cq = book.create_sheet(kb.CQ_SHEET)
    cq.append(["CaseNumber", "HRCSVC21_URL__c"])
    cq.append(["00000001", "https://baxter-international.atlassian.net/wiki/x/AbCd"])
    cq.append(["00000002", "No KB Required"])
    cq.append(["00000003", "https://hill-rom.atlassian.net/wiki/x/ZzZz"])
    book.save(path)
    return path


def test_kb_plan_separates_links_from_unusable_values(tmp_path: Path) -> None:
    workbook = _cases_workbook(tmp_path)
    directory = tmp_path / "kb"
    out = kb.plan(directory, workbook=workbook)
    assert out["article_count"] == 1
    assert out["unavailable"] == ["not_a_link", "site_not_accessible"]
    index = kb.read_index(directory)
    assert index["00000001"][0]["key"] == "AbCd"
    assert index["00000003"][0]["reason"] == "site_not_accessible"


def test_kb_normalize_reads_saved_pages(tmp_path: Path) -> None:
    workbook = _cases_workbook(tmp_path)
    directory = tmp_path / "kb"
    kb.plan(directory, workbook=workbook)
    (directory / kb.PAGES_DIR / "AbCd.json").write_text(
        json.dumps({"data": {"title": "Lost audio station", "body": {"value": "Replace the SRS."}}}),
        encoding="utf-8",
    )
    out = kb.normalize(directory)
    assert out == {"fetched": 1, "missing": [], "article_count": 1}
    article = kb.read_articles(directory)["AbCd"]
    assert article["title"] == "Lost audio station"
    assert article["text"] == "Replace the SRS."


# -- batching and version rule -------------------------------------------------------


@pytest.mark.parametrize(
    ("count", "expected"),
    [(5, [5]), (20, [20]), (23, [12, 11]), (40, [20, 20]), (41, [14, 14, 13]), (100, [20] * 5)],
)
def test_split_sizes_stays_inside_the_bounds(count: int, expected: list[int]) -> None:
    assert mapping.split_sizes(count, min_batch=10, max_batch=20) == expected


def test_split_sizes_rejects_bad_bounds() -> None:
    with pytest.raises(mapping.MappingError):
        mapping.split_sizes(10, min_batch=30, max_batch=20)


@pytest.mark.parametrize(
    ("reported", "fixed", "keep", "reason"),
    [
        ([[4, 0, 405]], [], True, "reported_in_same_version"),
        ([[4, 1, 200]], [[4, 2, 0]], True, "fixed_in_later_version"),
        ([[4, 1, 200]], [[4, 0, 300]], False, "version_excluded"),
        ([], [], True, "issue_version_unknown"),
    ],
)
def test_version_match(reported: list, fixed: list, keep: bool, reason: str) -> None:
    issue = {"versions_reported": reported, "versions_fix": fixed}
    assert mapping.version_match([4, 0, 405], issue) == (keep, reason)


def test_version_match_keeps_everything_when_the_case_version_is_unknown() -> None:
    issue = {"versions_reported": [[4, 1, 200]], "versions_fix": [[4, 1, 300]]}
    assert mapping.version_match([], issue) == (True, "case_version_unknown")


def test_similarity_ranks_the_closest_issue_first() -> None:
    index = mapping.Similarity(
        [("NC-1", "console speaker audio distorted"), ("NC-2", "database backup schedule failure")]
    )
    assert index.rank("audio distorted on the console speaker")[0][0] == "NC-1"


# -- prepare / publish / finalize ------------------------------------------------------


def _prepared(tmp_path: Path) -> tuple[Path, Path]:
    workbook = _cases_workbook(tmp_path)
    corpus = _corpus(
        tmp_path,
        [
            _raw_issue("NC-1", summary="text 00000001 console", customfield_10693=""),
            _raw_issue("NC-2", summary="unrelated database issue", customfield_10693="00000002"),
        ],
    )
    jira.normalize(corpus)
    kb_dir = tmp_path / "kb"
    kb.plan(kb_dir, workbook=workbook)
    (kb_dir / kb.PAGES_DIR / "AbCd.json").write_text(
        json.dumps({"data": {"title": "KB", "body": {"value": "article text"}}}), encoding="utf-8"
    )
    kb.normalize(kb_dir)
    assignments = tmp_path / "case_assignments.jsonl"
    assignments.write_text(
        "\n".join(
            json.dumps({"case_number": number, "final_cluster": "Audio", "final_subcluster": "Console audio"})
            for number in ("00000001", "00000002", "00000003")
        ),
        encoding="utf-8",
    )
    out_dir = tmp_path / "run"
    manifest = mapping.prepare(
        assignments, out_dir, workbook=workbook, corpora=[corpus], kb_dir=kb_dir, min_batch=1, max_batch=20
    )
    assert manifest["direct_link_count"] == 1  # 00000002 is named by NC-2's CRM field
    assert manifest["mapped_case_count"] == 2
    return out_dir, manifest["requests"][0]["request_id"]


def test_prepare_short_circuits_direct_links_and_builds_one_request(tmp_path: Path) -> None:
    out_dir, request_id = _prepared(tmp_path)
    request = json.loads((out_dir / mapping.REQUESTS_DIR / f"{request_id}.json").read_text(encoding="utf-8"))
    assert request["case_numbers"] == ["00000001", "00000003"]
    assert "NC-1" in request["candidate_keys"]
    assert request["kb_article_ids"] == ["AbCd"]
    assert "article text" in request["user"]
    assert "KB article ids: AbCd" in request["user"]
    direct = [json.loads(line) for line in (out_dir / mapping.DIRECT_LINKS_FILE).read_text().splitlines()]
    assert direct == [
        {
            "case_number": "00000002",
            "match_type": "direct_link",
            "issue_keys": ["NC-2"],
            "rationale": "The Jira CRM Case Number field already names this case.",
        }
    ]


def test_prepare_refuses_to_overwrite(tmp_path: Path) -> None:
    out_dir, _ = _prepared(tmp_path)
    with pytest.raises(mapping.MappingError, match="already prepared"):
        mapping.prepare(
            tmp_path / "case_assignments.jsonl",
            out_dir,
            workbook=tmp_path / "cases_with_summaries.xlsx",
            corpora=[tmp_path / "corpus_defect"],
        )


def _candidate(cases: list[dict]) -> dict:
    return {"cases": cases}


def test_publish_rejects_unknown_cases_issues_and_articles(tmp_path: Path) -> None:
    out_dir, request_id = _prepared(tmp_path)
    path = tmp_path / "candidate.json"

    path.write_text(json.dumps(_candidate([{"case_number": "00000001", "match_type": "none"}])), encoding="utf-8")
    assert any("missing case numbers" in error for error in mapping.publish_response(out_dir, request_id, path))

    bad = _candidate(
        [
            {
                "case_number": "00000001",
                "match_type": "defect",
                "matches": [
                    {
                        "issue_key": "NC-404",
                        "issue_type": "defect",
                        "confidence": "high",
                        "rationale": "r",
                        "kb_support": ["ZZZZ"],
                    }
                ],
            },
            {"case_number": "00000003", "match_type": "none"},
        ]
    )
    path.write_text(json.dumps(bad), encoding="utf-8")
    errors = mapping.publish_response(out_dir, request_id, path)
    assert any("was not a candidate" in error for error in errors)
    assert any("was not supplied" in error for error in errors)


def test_publish_rejects_none_with_matches(tmp_path: Path) -> None:
    out_dir, request_id = _prepared(tmp_path)
    path = tmp_path / "candidate.json"
    path.write_text(
        json.dumps(
            _candidate(
                [
                    {
                        "case_number": "00000001",
                        "match_type": "none",
                        "matches": [
                            {"issue_key": "NC-1", "issue_type": "defect", "confidence": "low", "rationale": "r"}
                        ],
                    },
                    {"case_number": "00000003", "match_type": "none"},
                ]
            )
        ),
        encoding="utf-8",
    )
    assert any("must have no matches" in error for error in mapping.publish_response(out_dir, request_id, path))


def test_finalize_combines_direct_links_and_responses(tmp_path: Path) -> None:
    out_dir, request_id = _prepared(tmp_path)
    path = tmp_path / "candidate.json"
    path.write_text(
        "```json\n"
        + json.dumps(
            _candidate(
                [
                    {
                        "case_number": "00000001",
                        "match_type": "defect",
                        "matches": [
                            {
                                "issue_key": "NC-1",
                                "issue_type": "defect",
                                "confidence": "high",
                                "rationale": "same console audio fault",
                                "version_justification": "fixed in 4.2.000",
                                "kb_support": ["AbCd"],
                            }
                        ],
                    },
                    {"case_number": "00000003", "match_type": "none", "rationale": "hardware swap"},
                ]
            )
        )
        + "\n```",
        encoding="utf-8",
    )
    assert mapping.publish_response(out_dir, request_id, path) == []
    assert mapping.check(out_dir)["ready"] is True

    summary = mapping.finalize(out_dir)
    assert summary["case_count"] == 3
    assert summary["direct_links"] == 1
    assert summary["cases_with_a_match"] == 2
    assert summary["cases_without_a_match"] == 1
    assert summary["top_issues"] == [{"issue_key": "NC-1", "cases": 1}]
    rows = [json.loads(line) for line in (out_dir / mapping.MAPPINGS_FILE).read_text().splitlines()]
    assert {row["case_number"]: row["match_type"] for row in rows} == {
        "00000001": "defect",
        "00000002": "direct_link",
        "00000003": "none",
    }
    book = load_workbook(out_dir / mapping.WORKBOOK_FILE)
    assert book.sheetnames == ["Mappings", "Batches"]


def test_finalize_refuses_while_requests_are_open(tmp_path: Path) -> None:
    out_dir, _ = _prepared(tmp_path)
    with pytest.raises(mapping.MappingError, match="lack a valid response"):
        mapping.finalize(out_dir)


# -- CLI and prompts --------------------------------------------------------------------


def test_cli_reports_errors_with_exit_code_2(tmp_path: Path, capsys) -> None:
    assert main(["jira-normalize", "--dir", str(tmp_path / "missing")]) == 2
    assert "error:" in capsys.readouterr().err


def test_cli_check_exits_1_while_incomplete(tmp_path: Path, capsys) -> None:
    out_dir, _ = _prepared(tmp_path)
    assert main(["check", "--dir", str(out_dir)]) == 1
    assert json.loads(capsys.readouterr().out)["ready"] is False


def test_defect_mapping_prompt_carries_glossary_and_examples() -> None:
    text = get_analysis("defect_mapping").prompt("v1")
    assert "{{GLOSSARY}}" not in text and "{{EXAMPLES}}" not in text
    assert "NC-24878: 00622535" in text
    assert "RCA-1823: 01034871" in text
    assert "Nurse Call" in text


def test_known_mappings_are_unique_and_normalized() -> None:
    from complaint_analyzer.prompts.defect_mapping import EXAMPLES_PATH

    rows = [line.split(",") for line in EXAMPLES_PATH.read_text(encoding="utf-8").splitlines()[1:]]
    keys = [row[0] for row in rows]
    assert len(keys) == len(set(keys))
    for row in rows:
        assert row[0].startswith(("NC-", "RCA-"))
        assert all(len(number) == 8 and number.isdigit() for number in row[1].split())
