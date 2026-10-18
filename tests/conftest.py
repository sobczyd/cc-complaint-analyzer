"""Shared fixtures: a generated workbook, settings, and a scripted fake transport."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence

import pytest
from openpyxl import Workbook as OpenpyxlWorkbook
from pydantic import SecretStr

from complaint_analyzer.llm_client.protocol import ChatRequest, ChatResponse, LlmError
from complaint_analyzer.core.settings import Credentials, Settings

SHEET = "Service Records"

HEADERS = ["Record ID", "Status", "Priority", "Service description", "Troubleshooting", "Resolution"]

ROWS: list[list[Any]] = [
    [
        "R-001",
        "Closed",
        3,
        "Customer reports the router drops connection every evening.",
        "Checked the line, found heavy packet loss on the uplink port.",
        "Replaced the uplink cable and the connection stabilized.",
    ],
    [
        "R-002",
        "Closed",
        1,
        "No dial tone on the main line.",
        "Rebooted the gateway.",
        "Service restored after reboot.",
    ],
    [
        "R-003",
        "Open",
        2,
        "Intermittent slowness during business hours.",
        "",
        "",
    ],
    [
        "R-004",
        "Closed",
        5,
        "Handset will not power on.",
        "Swapped the power supply.",
        "Handset powered on with the replacement supply.",
    ],
    [
        "",
        "Closed",
        2,
        "Row with no identifier.",
        "",
        "",
    ],
    [
        "R-001",
        "Closed",
        "",
        "Duplicate identifier row.",
        "Logged against the same identifier.",
        "",
    ],
]


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    """A small xlsx workbook covering blanks, duplicates, and mixed statuses."""
    path = tmp_path / "service_records.xlsx"
    book = OpenpyxlWorkbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(row)
    book.save(path)
    book.close()
    return path


def make_settings(workbook_path: Path, runs_dir: Path, **overrides: Any) -> Settings:
    """Build a valid Settings object, with keyword overrides merged per section."""
    raw: dict[str, Any] = {
        "dataset": {
            "path": str(workbook_path),
            "sheet": SHEET,
            "id_field": "Record ID",
            "fields": ["Service description", "Troubleshooting", "Resolution"],
            "filters": [{"field": "Status", "operator": "eq", "value": "Closed"}],
            "limit": None,
        },
        "analysis": {"name": "rca", "version": "v1", "concurrency": 2},
        "rate_limits": {"rpm": None, "tpm": None, "max_attempts": 3},
        "cache": {"mode": "off"},
        "output": {"runs_dir": str(runs_dir)},
    }
    for section, values in overrides.items():
        if isinstance(values, dict) and isinstance(raw.get(section), dict):
            raw[section].update(values)
        else:
            raw[section] = values
    return Settings.model_validate(raw)


@pytest.fixture
def settings(workbook_path: Path, tmp_path: Path) -> Settings:
    return make_settings(workbook_path, tmp_path / "runs")


@pytest.fixture
def credentials() -> Credentials:
    return Credentials(
        base_url="https://example.invalid/v1",
        api_key=SecretStr("sk-test-secret-value"),
        model="test-model",
    )


def rca_payload(**overrides: Any) -> dict[str, Any]:
    """A schema-valid RCA reply that also passes the validators."""
    payload: dict[str, Any] = {
        "symptom": "The router drops its connection every evening.",
        "restoration_action": "Replaced the uplink cable.",
        "documented_cause": None,
        "hypothesis": "Degraded uplink cabling caused the packet loss.",
        "cause_status": "hypothesized",
        "evidence": [],
        "missing_evidence": ["Line test results after the replacement."],
        "needs_review": True,
    }
    payload.update(overrides)
    return payload


class FakeTransport:
    """A scripted ``ChatClient``. Records every request it is given."""

    def __init__(
        self,
        responder: Callable[[ChatRequest], Any] | None = None,
        *,
        replies: Iterable[Any] = (),
    ) -> None:
        self.requests: list[ChatRequest] = []
        self._responder = responder
        self._replies = list(replies)
        self.closed = False

    async def complete(self, request: ChatRequest) -> ChatResponse:
        self.requests.append(request)
        if self._responder is not None:
            outcome = self._responder(request)
        elif self._replies:
            outcome = self._replies.pop(0)
        else:
            outcome = rca_payload()
        return self._materialize(outcome, request)

    async def aclose(self) -> None:
        self.closed = True

    @property
    def call_count(self) -> int:
        return len(self.requests)

    @staticmethod
    def _materialize(outcome: Any, request: ChatRequest) -> ChatResponse:
        if isinstance(outcome, BaseException):
            raise outcome
        if isinstance(outcome, ChatResponse):
            return outcome
        content = outcome if isinstance(outcome, str) else json.dumps(outcome)
        return ChatResponse(
            content=content,
            model=request.model,
            prompt_tokens=120,
            completion_tokens=60,
        )


def transient(message: str = "temporarily unavailable") -> LlmError:
    return LlmError(message, status_code=503, retryable=True, kind="http_error")


def permanent(message: str = "bad request") -> LlmError:
    return LlmError(message, status_code=400, retryable=False, kind="invalid_request")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def sheet_rows(path: Path) -> list[Sequence[Any]]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [tuple(row) for row in book[book.sheetnames[0]].iter_rows(values_only=True)]
    finally:
        book.close()
