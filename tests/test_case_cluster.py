"""Batch clustering and merge workflow."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from complaint_analyzer.prompts import get_analysis
from complaint_analyzer.workflows.case_cluster import clustering as cl
from complaint_analyzer.workflows.case_cluster.__main__ import main

HEADER = ["CaseNumber", *cl.SUMMARY_COLUMNS, "case_summary.analysis_status"]


def _workbook(path: Path, count: int) -> Path:
    book = Workbook()
    ws = book.active
    ws.title = "Cases"
    ws.append(HEADER)
    for i in range(count):
        ws.append([1000 + i, f"summary {i}", f"issue {i}", "RCB", None, "fixed", "resolved", "kp", "ok"])
    ws.append([9999, None, None, None, None, None, None, None, "not_run"])
    book.save(path)
    return path


def _batch(request: dict) -> dict:
    cases = request["case_numbers"]
    half = max(1, len(cases) // 2)
    groups = [cases[:half], cases[half:]] if cases[half:] else [cases]
    return {
        "clusters": [
            {
                "cluster_id": f"C{n}",
                "label": f"Cluster {n}",
                "description": "d",
                "rationale": "r",
                "subclusters": [
                    {"subcluster_id": f"C{n}.S1", "label": f"Sub {n}", "case_numbers": group}
                ],
            }
            for n, group in enumerate(groups, 1)
        ]
    }


def test_glossary_is_injected_into_both_prompts() -> None:
    for name in ("case_cluster", "case_cluster_merge"):
        prompt = get_analysis(name).prompt("v1")
        assert "{{GLOSSARY}}" not in prompt
        assert "Clustering rules" in prompt and "GRS10" in prompt


def test_context_uses_column_name_convention(tmp_path: Path) -> None:
    out = tmp_path / "run"
    cl.prepare(_workbook(tmp_path / "w.xlsx", 3), out, limit=None)
    request = json.loads((out / "requests" / "batch_01.json").read_text())
    assert request["case_numbers"] == ["1000", "1001", "1002"]  # not_run row excluded
    assert "case_summary.summary: summary 0" in request["user"]
    assert "case_summary.actions_taken" not in request["user"]  # blank column skipped
    assert request["user"].count(cl.CASE_SEPARATOR) == 3


def test_single_batch_finalizes_without_merge(tmp_path: Path) -> None:
    out = tmp_path / "run"
    manifest = cl.prepare(_workbook(tmp_path / "w.xlsx", 4), out, limit=None)
    assert manifest["merge_required"] is False
    candidate = tmp_path / "c.json"
    request = json.loads((out / "requests" / "batch_01.json").read_text())
    candidate.write_text("```json\n" + json.dumps(_batch(request)) + "\n```")
    assert cl.publish_response(out, "batch_01", candidate) == []
    assert cl.check(out)["ready"] is True
    summary = cl.finalize(out)
    assert summary["case_count"] == 4
    assert "Case Assignments" in load_workbook(out / "clusters.xlsx").sheetnames


def test_invalid_batch_is_rejected(tmp_path: Path) -> None:
    out = tmp_path / "run"
    cl.prepare(_workbook(tmp_path / "w.xlsx", 3), out, limit=None)
    request = json.loads((out / "requests" / "batch_01.json").read_text())
    bad = _batch(request)
    bad["clusters"][0]["subclusters"][0]["case_numbers"][0] = "424242"
    bad["clusters"][0]["subclusters"][0]["label"] = "one two three four five six seven eight nine"
    candidate = tmp_path / "bad.json"
    candidate.write_text(json.dumps(bad))
    errors = cl.publish_response(out, "batch_01", candidate)
    assert errors
    assert cl.check(out)["incomplete_request_ids"] == ["batch_01"]
    missing = _batch(request)
    missing["clusters"][1]["subclusters"][0]["case_numbers"].pop()
    candidate.write_text(json.dumps(missing))
    assert any("missing" in e or "unknown" in e for e in cl.publish_response(out, "batch_01", candidate))


def test_two_batches_merge_and_finalize(tmp_path: Path) -> None:
    out = tmp_path / "run"
    manifest = cl.prepare(_workbook(tmp_path / "w.xlsx", 6), out, limit=None, batch_size=3)
    assert manifest["merge_required"] is True
    candidate = tmp_path / "c.json"
    for batch in manifest["batches"]:
        request = json.loads((out / "requests" / f"{batch['batch_id']}.json").read_text())
        candidate.write_text(json.dumps(_batch(request)))
        assert cl.publish_response(out, batch["batch_id"], candidate) == []
    with pytest.raises(cl.ClusterError):
        cl.finalize(out)

    assert cl.prepare_merge(out)["subclusters"] == 4
    merge_request = json.loads((out / "requests" / "merge.json").read_text())
    mappings = [
        {
            "batch_id": s["batch_id"],
            "subcluster_id": s["subcluster_id"],
            "final_cluster": "Lost audio",
            "final_subcluster": "Sub A" if s["subcluster_id"].endswith("1") and s["subcluster_id"].startswith("C1") else "Sub B",
        }
        for s in merge_request["subclusters"]
    ]
    merge = {
        "final_clusters": [
            {"label": "Lost audio", "subclusters": [{"label": "Sub A"}, {"label": "Sub B"}]}
        ],
        "mappings": mappings,
    }
    candidate.write_text(json.dumps(merge))
    assert cl.publish_response(out, "merge", candidate) == []

    summary = cl.finalize(out)
    assert summary["final_clusters"] == 1 and summary["case_count"] == 6
    rows = [json.loads(line) for line in (out / "case_assignments.jsonl").read_text().splitlines()]
    assert {r["final_cluster"] for r in rows} == {"Lost audio"}
    assert [r["case_number"] for r in rows] == [str(1000 + i) for i in range(6)]


def test_merge_rejects_unmapped_subcluster(tmp_path: Path) -> None:
    out = tmp_path / "run"
    manifest = cl.prepare(_workbook(tmp_path / "w.xlsx", 4), out, limit=None, batch_size=2)
    candidate = tmp_path / "c.json"
    for batch in manifest["batches"]:
        request = json.loads((out / "requests" / f"{batch['batch_id']}.json").read_text())
        candidate.write_text(json.dumps(_batch(request)))
        cl.publish_response(out, batch["batch_id"], candidate)
    cl.prepare_merge(out)
    first = json.loads((out / "requests" / "merge.json").read_text())["subclusters"][0]
    candidate.write_text(
        json.dumps(
            {
                "final_clusters": [{"label": "A", "subclusters": [{"label": "B"}]}],
                "mappings": [{**first, "final_cluster": "A", "final_subcluster": "B"}],
            }
        )
    )
    errors = cl.publish_response(out, "merge", candidate)
    assert any("unmapped" in e for e in errors)


def test_cli_prepare_and_check(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    workbook = _workbook(tmp_path / "w.xlsx", 2)
    out = tmp_path / "run"
    assert main(["prepare", "--workbook", str(workbook), "--out-dir", str(out), "--limit", "0"]) == 0
    capsys.readouterr()
    assert main(["check", "--dir", str(out)]) == 1
    assert json.loads(capsys.readouterr().out)["incomplete_request_ids"] == ["batch_01"]


def _publish_merge(out: Path, request_id: str, tmp_path: Path, labels: tuple[str, ...]) -> None:
    request = json.loads((out / "requests" / f"{request_id}.json").read_text())
    mappings = [
        {
            "batch_id": s["batch_id"],
            "subcluster_id": s["subcluster_id"],
            "final_cluster": f"{labels[i % len(labels)]} problems",
            "final_subcluster": f"{labels[i % len(labels)]} detail",
        }
        for i, s in enumerate(request["subclusters"])
    ]
    merge = {
        "final_clusters": [
            {"label": f"{name} problems", "subclusters": [{"label": f"{name} detail"}]} for name in labels
        ],
        "mappings": mappings,
    }
    candidate = tmp_path / f"{request_id}.candidate.json"
    candidate.write_text(json.dumps(merge))
    assert cl.publish_response(out, request_id, candidate) == []


def test_merge_plan_tiers() -> None:
    assert cl.build_merge_plan(["b1"], 5) == []
    single = cl.build_merge_plan([f"b{i}" for i in range(5)], 5)
    assert [r["request_id"] for t in single for r in t] == ["merge"]
    plan = cl.build_merge_plan([f"b{i}" for i in range(26)], 5)
    assert [len(t) for t in plan] == [6, 2, 1]
    assert plan[0][0]["prompt"] == "case_cluster_merge" and plan[0][0]["request_id"] == "merge_01_01"
    assert plan[1][0]["prompt"] == "case_cluster_meta_merge" and plan[1][0]["inputs"][0] == "merge_01_01"
    assert plan[2][0]["request_id"] == "meta_03_01"


def test_tiered_merge_and_finalize(tmp_path: Path) -> None:
    out = tmp_path / "run"
    manifest = cl.prepare(_workbook(tmp_path / "w.xlsx", 8), out, limit=None, batch_size=1, merge_fanin=3)
    candidate = tmp_path / "c.json"
    for batch in manifest["batches"]:
        request = json.loads((out / "requests" / f"{batch['batch_id']}.json").read_text())
        candidate.write_text(json.dumps(_batch(request)))
        assert cl.publish_response(out, batch["batch_id"], candidate) == []

    tier1 = cl.prepare_merge(out)
    assert tier1["tier"] == 1 and tier1["final"] is False
    assert tier1["request_ids"] == ["merge_01_01", "merge_01_02", "merge_01_03"]
    for request_id in tier1["request_ids"]:
        _publish_merge(out, request_id, tmp_path, ("Audio", "Light"))
    with pytest.raises(cl.ClusterError):
        cl.finalize(out)

    tier2 = cl.prepare_merge(out)
    assert tier2["request_ids"] == ["meta_02_01"] and tier2["final"] is True
    meta = json.loads((out / "requests" / "meta_02_01.json").read_text())
    assert meta["prompt"] == "case_cluster_meta_merge"
    assert "=== Group merge_01_01 (3 cases) ===" in meta["user"]
    assert len(meta["subclusters"]) == 6
    _publish_merge(out, "meta_02_01", tmp_path, ("Audio", "Light", "Other"))
    assert cl.prepare_merge(out)["complete"] is True

    summary = cl.finalize(out)
    assert summary["case_count"] == 8 and summary["final_clusters"] == 3
    rows = [json.loads(line) for line in (out / "case_assignments.jsonl").read_text().splitlines()]
    assert [r["case_number"] for r in rows] == [str(1000 + i) for i in range(8)]
    merge_map = load_workbook(out / "clusters.xlsx")["Merge Map"]
    assert merge_map.max_row == 9 and "merge_01_" in merge_map["F2"].value
