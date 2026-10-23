from openpyxl import load_workbook

from complaint_analyzer.workflows.powerbi_export.export import DEFAULT_TEMPLATE
from complaint_analyzer.workflows.powerbi_export.package import build_package
from complaint_analyzer.workflows.powerbi_export.trackwise import (
    CODES_FILE,
    FACT_COMPLAINT_COLUMNS,
    PRODUCT_FILE,
    QUERY_FILE,
    apply_trackwise,
    complaint_rows,
    load_trackwise,
    read_csv,
)
from tests.test_powerbi_package import _run

QUERY_HEADERS = ["PR ID", "Source Identifier", "Project", "PR State", "Date Opened", "Event Description"]


def _csv(path, headers, rows, encoding="utf-16"):
    lines = ["\t".join(headers)] + ["\t".join(row) for row in rows]
    path.write_text("\n".join(lines) + "\n", encoding=encoding)


def _trackwise(tmp_path):
    folder = tmp_path / "trackwise"
    folder.mkdir()
    _csv(
        folder / QUERY_FILE,
        QUERY_HEADERS,
        [
            ["100", "01000001", "Complaint", "Open", "01/01/2026 01:00 PM", "first"],
            ["101", "1000001", "Complaint", "Closed", "01/02/2026 01:00 PM", "second"],
            ["102", "01000001", "Complaint", "Pending", "01/03/2026 01:00 PM", "third"],
            ["200", "01999999", "Complaint", "Open", "", "no case"],
            ["300", "01000002", "Other", "Open", "", "ignored project"],
        ],
    )
    _csv(
        folder / PRODUCT_FILE,
        ["PR ID", "Product Family^", "Product Code/Part #^", "Product Lot/Serial #^"],
        [["101", "Old", "P-1", "S1"], ["101", "Family", "P-product", "S2"], ["200", "Beds", "P-2", "S3"]],
    )
    _csv(
        folder / CODES_FILE,
        ["PR ID", "Product Code/Part #^", "As Determined Code^", "As Determined Component Code^"],
        [["101", "P-codes", "Code A", ""]],
        encoding="utf-8-sig",
    )
    return folder


def test_read_csv_handles_utf16_and_utf8(tmp_path):
    folder = _trackwise(tmp_path)
    assert read_csv(folder / QUERY_FILE)[0]["Source Identifier"] == "01000001"
    assert read_csv(folder / CODES_FILE)[0]["As Determined Code^"] == "Code A"


def test_load_filters_project_and_flags_duplicates(tmp_path):
    tw = load_trackwise(_trackwise(tmp_path))
    assert set(tw.query) == {"100", "101", "102", "200"}
    assert tw.product_duplicates == {"101"}
    assert tw.product["101"]["Product Family^"] == "Family"
    assert tw.part_number("101") == "P-codes"
    assert tw.part_number("200") == "P-2"


def test_match_ignores_leading_zeros_prefers_closed_and_falls_back_to_pr(tmp_path):
    tw = load_trackwise(_trackwise(tmp_path))
    assert tw.match("01000001") == ("101", 3)
    assert tw.match("nope", 200.0) == ("200", 1)
    assert tw.match("nope", None) == (None, 0)


def test_apply_fills_columns_and_merged_prefers_trackwise():
    tw = load_trackwise_from(
        {"101": {"Source Identifier": "01000001", "PR State": "Closed", "Date Opened": "d"}}
    )
    headers = [
        "Source Identifier", "PR ID", "PR State", "Date Opened", "kept_file", "As Determined Code^",
        "As Determined Problem Codes", "As Determined Problem Codes_Merged", "Component Code", "Component Code_Merged",
    ]
    row = ["01000001", None, None, None, None, None, "SF code", None, "SF comp", None]
    pr_key, count = apply_trackwise(tw, headers, row, "cases.xlsx")
    assert (pr_key, count) == ("101", 1)
    assert row[1:5] == [101, "Closed", "d", "cases.xlsx"]
    assert row[7] == "SF code" and row[9] == "SF comp"
    row = ["x", None, None, None, None, None, "SF code", None, "SF comp", None]
    assert apply_trackwise(None, headers, row, "cases.xlsx") == (None, 0)
    assert row[7] == "SF code" and row[4] is None


def load_trackwise_from(query):
    from complaint_analyzer.workflows.powerbi_export.trackwise import TrackwiseData

    return TrackwiseData(
        directory=None, query=query, product={}, codes={}, product_duplicates=set(), codes_duplicates=set(),
        by_source={"1000001": list(query)},
    )


def test_complaint_rows_link_to_cases(tmp_path):
    tw = load_trackwise(_trackwise(tmp_path))
    rows = complaint_rows(tw, [{"Source Identifier Raw": "01000001", "Source Identifier": "01000001", "PR ID": None}])
    by_pr = {row["PR ID"]: row for row in rows}
    assert set(by_pr) == {100, 101, 102, 200}
    assert list(rows[0]) == FACT_COMPLAINT_COLUMNS
    assert by_pr[101]["Source Identifier Raw"] == "01000001" and by_pr[100]["Source Identifier Raw"] == "01000001"
    assert by_pr[200]["Source Identifier Raw"] == "01999999"
    assert by_pr[101]["Product Lot/Serial #^"] == "S2" and by_pr[101]["product_IsDuplicated"] == "Yes"


def test_build_package_writes_fact_complaint(tmp_path):
    run = _run(tmp_path)
    result = build_package(run, template=DEFAULT_TEMPLATE, trackwise_dir=_trackwise(tmp_path))
    out = run / "powerbi_export_package"
    assert result["rows"]["FactComplaint"] == 4
    assert (out / "FactComplaint.xlsx").is_file()
    sheet = load_workbook(out / "FactSource.xlsx").active
    headers = [cell.value for cell in sheet[1]]
    first = dict(zip(headers, [cell.value for cell in next(sheet.iter_rows(min_row=2))]))
    assert first["PR State"] == "Closed" and first["kept_file"] == "cases.xlsx"
    assert first["Product Code/Part #^"] == "P-codes"
    dim = load_workbook(out / "DimID.xlsx").active
    assert [cell.value for cell in next(dim.iter_rows(min_row=2))][-1] == "Closed"
