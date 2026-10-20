"""Tests for the Salesforce export and the Excel merge."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from complaint_analyzer.workflows.report_merge import MergeError
from complaint_analyzer.workflows.report_merge.__main__ import main as report_main
from complaint_analyzer.workflows.report_merge.merge import GENERATED_FILL_COLOR, LEGEND_SHEET, merge_results
from complaint_analyzer.workflows.salesforce_export.__main__ import main as export_main
from complaint_analyzer.workflows.salesforce_export.export import (
    ExportError,
    export_query_set,
    flatten_records,
    parse_response,
    parse_select,
)

SOQL = """SELECT Id, CaseNumber, Account.Name,
       (SELECT Id, Cause_Code__r.Name FROM Problem_Codes__r WHERE X = 'a, b')
FROM Case WHERE CreatedDate = LAST_N_DAYS:7"""


def _record(number: str, *, codes: list[str] | None = None, subject: str = "Subj") -> dict:
    problem = None
    if codes:
        problem = {
            "totalSize": len(codes),
            "done": True,
            "records": [
                {"attributes": {"type": "Problem_Code__c"}, "Id": f"p{i}", "Cause_Code__r": {"Name": c}}
                for i, c in enumerate(codes)
            ],
        }
    return {
        "attributes": {"type": "Case"},
        "Id": f"id{number}",
        "CaseNumber": number,
        "Subject": subject,
        "Account": {"attributes": {"type": "Account"}, "Name": "Acme"},
        "Problem_Codes__r": problem,
        "Complaint_Qualifications__r": None,
    }


def _write_raw(path: Path, records: list[dict], *, done: bool = True) -> Path:
    payload = {"status": 0, "result": {"records": records, "totalSize": len(records), "done": done}}
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def _sheet_rows(path: Path, title: str) -> list[list]:
    book = load_workbook(path)
    return [list(row) for row in book[title].iter_rows(values_only=True)]


# ------------------------------------------------------------------ export


def test_parse_select_reads_fields_and_children():
    shape = parse_select(SOQL)
    assert shape.fields == ("Id", "CaseNumber", "Account.Name")
    assert shape.children == {"Problem_Codes__r": ("Id", "Cause_Code__r.Name")}


def test_parse_response_rejects_errors():
    with pytest.raises(ExportError, match="did not return JSON"):
        parse_response("not json")
    with pytest.raises(ExportError, match="failed: boom"):
        parse_response(json.dumps({"status": 1, "message": "boom"}))


def test_flatten_records_joins_children():
    export = flatten_records(
        [_record("001", codes=["A", "B"]), _record("002")], shape=parse_select(SOQL)
    )
    main = export.main
    assert main.columns[:3] == ["Id", "CaseNumber", "Account.Name"]
    assert "Problem Codes" in main.columns and "Problem Codes (count)" in main.columns
    first, second = main.rows
    assert first["Account.Name"] == "Acme"
    assert first["Problem Codes (count)"] == 2
    assert "Cause_Code__r.Name: A" in first["Problem Codes"]
    assert second["Problem Codes (count)"] == 0
    assert [row["CaseNumber"] for row in export.children["Problem Codes"].rows] == ["001", "001"]


def test_export_query_set_from_json(tmp_path):
    raw = _write_raw(tmp_path / "saved.json", [_record("001", codes=["A"])], done=False)
    out = tmp_path / "out"
    result = export_query_set("cq_crm_codes_nurse_call_no_comments", out, from_json=raw)
    assert result.row_count == 1 and result.done is False
    for name in ("query.soql", "raw.json", "export.json", "cases.xlsx"):
        assert (out / name).is_file()
    meta = json.loads((out / "export.json").read_text(encoding="utf-8"))
    assert meta["version"] == "v3" and meta["row_count"] == 1
    header = _sheet_rows(out / "cases.xlsx", "Cases")[0]
    assert header[:2] == ["Id", "CaseNumber"]
    assert "Problem Codes" in load_workbook(out / "cases.xlsx").sheetnames


def test_export_cli_reports_errors(tmp_path, capsys):
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps({"status": 1, "message": "nope"}), encoding="utf-8")
    code = export_main(
        ["export", "--query-set", "cq_crm_codes_nurse_call_no_comments", "--out-dir", str(tmp_path / "o"), "--from-json", str(bad)]
    )
    assert code == 1
    assert "nope" in capsys.readouterr().err


# ------------------------------------------------------------------ merge


def _cases(tmp_path: Path, numbers: list[str]) -> Path:
    out = tmp_path / "export"
    raw = _write_raw(tmp_path / "raw_in.json", [_record(n, codes=["A"]) for n in numbers])
    export_query_set("cq_crm_codes_nurse_call_no_comments", out, from_json=raw)
    return out / "cases.xlsx"


def _write_results(path: Path, entries: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(e) + "\n" for e in entries), encoding="utf-8")


def test_merge_results_adds_filled_columns_and_legend(tmp_path):
    cases = _cases(tmp_path, ["001", "002", "003"])
    results = tmp_path / "run" / "results.jsonl"
    _write_results(
        results,
        [
            {"record_id": "001", "status": "ok", "values": {"summary": "S1", "key_points": ["a", "b"]}},
            {"record_id": "002", "status": "failed", "error": "bad json", "issues": ["missing summary"]},
            {"record_id": "999", "status": "ok", "values": {"summary": "orphan"}},
        ],
    )
    out = tmp_path / "merged.xlsx"
    result = merge_results(
        cases, results, out, output_keys=["summary", "key_points"], analysis="case_summary"
    )
    assert result.row_count == 3 and result.matched == 2 and result.unmatched_results == 1
    assert result.status_counts == {"failed": 1, "not_run": 1, "ok": 1}

    book = load_workbook(out)
    assert LEGEND_SHEET in book.sheetnames and "Problem Codes" in book.sheetnames
    sheet = book["Cases"]
    header = [cell.value for cell in sheet[1]]
    col = {name: i for i, name in enumerate(header)}
    for name in ("case_summary.summary", "case_summary.analysis_status", "case_summary.analysis_error"):
        assert name in col
    rows = list(sheet.iter_rows(min_row=2))
    by_id = {row[col["CaseNumber"]].value: row for row in rows}
    assert by_id["001"][col["case_summary.key_points"]].value == "a\nb"
    assert by_id["002"][col["case_summary.analysis_error"]].value == "bad json\nmissing summary"
    assert by_id["003"][col["case_summary.analysis_status"]].value == "not_run"
    generated_cell = by_id["001"][col["case_summary.summary"]]
    assert generated_cell.fill.fgColor.rgb.endswith(GENERATED_FILL_COLOR)
    assert by_id["001"][col["CaseNumber"]].fill.fgColor.rgb in (None, "00000000")


def test_merge_results_missing_file(tmp_path):
    cases = _cases(tmp_path, ["001"])
    missing = tmp_path / "run" / "results.jsonl"
    with pytest.raises(MergeError, match="results file not found"):
        merge_results(cases, missing, tmp_path / "m.xlsx", output_keys=["summary"], analysis="x")
    result = merge_results(
        cases, missing, tmp_path / "m.xlsx", output_keys=["summary"], analysis="x",
        allow_missing_results=True,
    )
    assert result.status_counts == {"not_run": 1}


def test_merge_results_rejects_column_clash(tmp_path):
    cases = _cases(tmp_path, ["001"])
    results = tmp_path / "run" / "results.jsonl"
    _write_results(results, [])
    with pytest.raises(MergeError, match="already exist"):
        merge_results(cases, results, tmp_path / "m.xlsx", output_keys=["Name"], analysis="Account")


def test_report_cli(tmp_path, capsys):
    cases = _cases(tmp_path, ["001"])
    run_dir = tmp_path / "run"
    _write_results(run_dir / "results.jsonl", [{"record_id": "001", "status": "ok", "values": {"summary": "S"}}])
    out = tmp_path / "merged.xlsx"
    code = report_main(["merge", "--cases", str(cases), "--run-dir", str(run_dir), "--out", str(out)])
    assert code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["matched"] == 1 and payload["status_counts"] == {"ok": 1}
    assert "case_summary.summary" in payload["generated_columns"]

    code = report_main(
        ["merge", "--cases", str(cases), "--run-dir", str(tmp_path / "none"), "--out", str(out)]
    )
    assert code == 2
