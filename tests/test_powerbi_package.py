import json

from openpyxl import Workbook, load_workbook

from complaint_analyzer.workflows.powerbi_export.export import DEFAULT_TEMPLATE
from complaint_analyzer.workflows.powerbi_export.package import build_package


def _book(path, sheets):
    book = Workbook()
    book.remove(book.active)
    for name, (headers, rows) in sheets.items():
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
    book.save(path)


def _rows(path):
    sheet = load_workbook(path).active
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, [cell.value for cell in row])) for row in sheet.iter_rows(min_row=2)]


def _run(tmp_path):
    run = tmp_path / "run"
    (run / "defects" / "d1").mkdir(parents=True)
    (run / "clusters" / "c1").mkdir(parents=True)
    _book(
        run / "cases.xlsx",
        {
            "Cases": (["CaseNumber", "Id"], [["01000001", "500AAAAAAAAAAAAAAA"], ["01000002", "500BBBBBBBBBBBBBBB"]]),
            "Problem Codes": (["CaseNumber"], []),
            "Complaint Qualifications": (["CaseNumber"], []),
        },
    )
    _book(
        run / "defects" / "d1" / "mappings.xlsx",
        {
            "Mappings": (
                ["case_number", "cluster", "subcluster", "match_type", "issue_key", "issue_type", "confidence",
                 "rationale", "version_justification", "kb_support"],
                [
                    ["01000001", "c", "s", "match", "NC-1", "Defect", "high", "r1", None, None],
                    ["01000001", "c", "s", "match", "RCA-2", "RCA", "low", "r2", "v", "KB1"],
                    ["01000002", "c", "s", "none", None, None, None, "no defect", None, None],
                    ["09999999", "c", "s", "match", "NC-9", "Defect", "high", "orphan", None, None],
                ],
            ),
            "Batches": (["request_id"], []),
        },
    )
    (run / "clusters" / "c1" / "case_assignments.jsonl").write_text(
        "\n".join(
            json.dumps({"case_number": number, "final_cluster": "Cl", "final_subcluster": "Sub"})
            for number in ("01000001", "08888888")
        ),
        encoding="utf-8",
    )
    return run


def test_package_with_case_assignments_fallback(tmp_path):
    run = _run(tmp_path)
    result = build_package(run, template=DEFAULT_TEMPLATE)
    out = run / "powerbi_export_package"
    assert sorted(path.name for path in out.iterdir()) == [
        "DimCluster.xlsx", "DimDefect.xlsx", "DimID.xlsx", "FactSource.xlsx", "README.md", "package_manifest.json",
    ]
    assert result["rows"] == {"FactSource": 2, "DimID": 2, "DimCluster": 1, "DimDefect": 2}
    assert [row["Source Identifier Raw"] for row in _rows(out / "DimID.xlsx")] == ["01000001", "01000002"]
    defects = _rows(out / "DimDefect.xlsx")
    assert [row["issue_key"] for row in defects] == ["NC-1", "RCA-2"]
    assert defects[1]["kb_support"] == "KB1"
    cluster = _rows(out / "DimCluster.xlsx")
    assert cluster[0]["Case Number"] == "01000001" and cluster[0]["Resolved cluster"] == "Cl"
    assert result["coverage"]["cases_without_cluster"] == 1
    assert result["coverage"]["orphan_cluster_cases_dropped"] == 1
    assert result["coverage"]["orphan_defect_rows_dropped"] == 1
    assert result["coverage"]["cases_without_defect"] == 1


def test_package_prefers_comparison_and_warns_without_sources(tmp_path):
    run = _run(tmp_path)
    comparison = tmp_path / "comparison.xlsx"
    _book(
        comparison,
        {
            "Case Assignments": (
                ["Period", "CaseNumber", "Created", "Original cluster", "Original subcluster", "Resolved cluster",
                 "Resolved subcluster", "Mapping status", "Mapping confidence"],
                [["B", "01000002", "2026-10-01", "o", "o", "Resolved", "RS", "mapped", "high"]],
            )
        },
    )
    result = build_package(run, comparison=comparison, template=DEFAULT_TEMPLATE)
    cluster = _rows(run / "powerbi_export_package" / "DimCluster.xlsx")
    assert cluster == [
        {"Case Number": "01000002", "Period": "B", "Resolved cluster": "Resolved", "Resolved subcluster": "RS",
         "Mapping status": "mapped", "Mapping confidence": "high"}
    ]
    assert result["cluster_source"] == "comparison"

    (run / "clusters" / "c1" / "case_assignments.jsonl").unlink()
    (run / "defects" / "d1" / "mappings.xlsx").unlink()
    bare = build_package(run, template=DEFAULT_TEMPLATE, out_dir=tmp_path / "bare")
    assert bare["rows"]["DimCluster"] == 0 and bare["rows"]["DimDefect"] == 0
    assert len(bare["warnings"]) == 2
