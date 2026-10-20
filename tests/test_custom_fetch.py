"""Tests for ad-hoc (run-specific) Salesforce fetches."""

from __future__ import annotations

import json
from datetime import date

import pytest
import yaml
from openpyxl import load_workbook

from complaint_analyzer.workflows.analysis_run import AnalysisRunError
from complaint_analyzer.workflows.custom_fetch import (
    finalize_analysis,
    init_run,
    new_run_dir,
    prepare_analysis,
    slugify,
)
from complaint_analyzer.workflows.custom_fetch.__main__ import main as fetch_main
from complaint_analyzer.workflows.salesforce_export.__main__ import main as export_main
from complaint_analyzer.workflows.salesforce_export.export import (
    ExportError,
    export_soql,
    select_clause,
)

SOQL = "SELECT Id, CaseNumber, Account.Name FROM Case WHERE Status = 'Closed' ORDER BY CreatedDate"


def _raw(path, numbers):
    records = [
        {"attributes": {"type": "Case"}, "Id": f"id{n}", "CaseNumber": n,
         "Account": {"attributes": {"type": "Account"}, "Name": "Acme"}}
        for n in numbers
    ]
    payload = {"status": 0, "result": {"records": records, "totalSize": len(records), "done": True}}
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def test_export_soql_writes_run_artifacts(tmp_path):
    raw = _raw(tmp_path / "in.json", ["1", "2"])
    result = export_soql(SOQL, tmp_path / "out", from_json=raw)

    assert result.row_count == 2
    assert result.query_set is None
    assert (tmp_path / "out" / "query.soql").read_text(encoding="utf-8") == SOQL
    rows = list(load_workbook(result.workbook)["Cases"].iter_rows(values_only=True))
    assert rows[0][:2] == ("Id", "CaseNumber")


def test_export_soql_rejects_non_select(tmp_path):
    with pytest.raises(ExportError):
        export_soql("DELETE FROM Case", tmp_path / "out", from_json=_raw(tmp_path / "r.json", []))


def test_select_clause_strips_top_level_filters_only():
    clause = select_clause("cq_crm_codes_nurse_call_no_comments")

    assert clause.rstrip().endswith("FROM Case")
    assert "LAST_N_DAYS" not in clause
    assert "ORDER BY" not in clause
    assert "WHERE Product_Family__r.Name = 'CareComm - Nurse Call'" in clause


def test_run_folder_naming_avoids_collisions(tmp_path):
    day = date(2026, 10, 19)
    first = new_run_dir(tmp_path, "Nurse Call: May!", day)
    assert first.name == "nurse_call_may_2026-10-19"
    first.mkdir()
    assert new_run_dir(tmp_path, "Nurse Call: May!", day).name == "nurse_call_may_2026-10-19_2"
    assert slugify("!!!") == "custom_fetch"


def test_init_run_records_request_and_session(tmp_path):
    raw = _raw(tmp_path / "in.json", ["1"])
    result = init_run(
        SOQL, slug="closed cases", prompt="closed cases please", based_on="some_set",
        runs_root=tmp_path / "runs", day=date(2026, 10, 19), from_json=raw,
    )

    run = result.run_dir
    assert run.name == "closed_cases_2026-10-19"
    assert (run / "request.txt").read_text(encoding="utf-8") == "closed cases please"
    session = json.loads((run / "session.json").read_text(encoding="utf-8"))
    assert session["kind"] == "custom_fetch"
    assert session["fields_based_on"] == "some_set"
    assert session["export"]["row_count"] == 1


def test_clis(tmp_path, capsys):
    soql = tmp_path / "q.soql"
    soql.write_text(SOQL, encoding="utf-8")
    raw = _raw(tmp_path / "in.json", ["1"])

    assert fetch_main([
        "init", "--soql-file", str(soql), "--slug", "x", "--runs-root", str(tmp_path / "runs"),
        "--from-json", str(raw),
    ]) == 0
    assert "exported 1 record(s)" in capsys.readouterr().out

    assert export_main(["export", "--soql-file", str(soql), "--out-dir", str(tmp_path / "o"),
                        "--from-json", str(raw)]) == 0
    assert export_main(["fields", "--query-set", "cq_crm_codes_nurse_call_no_comments"]) == 0
    assert "FROM Case" in capsys.readouterr().out


def _run(tmp_path, numbers=("1", "2")):
    raw = _raw(tmp_path / "in.json", list(numbers))
    return init_run(
        SOQL, slug="closed", prompt="closed cases", runs_root=tmp_path / "runs",
        day=date(2026, 10, 19), from_json=raw,
    ).run_dir


def test_failed_export_leaves_no_run_folder(tmp_path):
    with pytest.raises(ExportError):
        init_run("DELETE FROM Case", slug="bad", runs_root=tmp_path / "runs", day=date(2026, 10, 19))
    assert list((tmp_path / "runs").iterdir()) == []


def test_prepare_writes_config_and_commands(tmp_path):
    run = _run(tmp_path)
    result = prepare_analysis(run, batch_size=5)

    assert result.needs_batch_run and result.run_id == "case_summary"
    assert result.fields == ("Account.Name",)
    config = yaml.safe_load((run / "case_summary.yaml").read_text(encoding="utf-8"))
    assert config["dataset"]["id_field"] == "CaseNumber"
    assert config["dataset"]["fields"] == ["Account.Name"]
    assert config["output"]["runs_dir"] == str(run.resolve())
    assert "--batch-size 5" in result.commands["prepare"]
    session = json.loads((run / "session.json").read_text(encoding="utf-8"))
    assert session["analysis"]["run_id"] == "case_summary"


def test_prepare_rejects_unknown_field_and_skips_empty_export(tmp_path):
    run = _run(tmp_path)
    with pytest.raises(AnalysisRunError):
        prepare_analysis(run, fields=["Nope"])
    (tmp_path / "e").mkdir()
    empty = _run(tmp_path / "e", numbers=())
    result = prepare_analysis(empty)
    assert not result.needs_batch_run and not (empty / "case_summary.yaml").exists()


def test_finalize_merges_results_and_requires_manifest(tmp_path, capsys):
    run = _run(tmp_path)
    prepare_analysis(run)
    results = run / "case_summary" / "results.jsonl"
    results.parent.mkdir()
    results.write_text(
        "".join(json.dumps({"record_id": n, "status": "ok", "values": {"summary": f"S{n}"}}) + "\n"
                for n in ("1", "2")),
        encoding="utf-8",
    )
    with pytest.raises(AnalysisRunError, match="not finalized"):
        finalize_analysis(run)

    (run / "case_summary" / "manifest.json").write_text("{}", encoding="utf-8")
    assert fetch_main(["finalize", "--run", str(run)]) == 0
    assert '"matched": 2' in capsys.readouterr().out
    book = load_workbook(run / "cases_with_summaries.xlsx")
    assert "Legend" in book.sheetnames
    session = json.loads((run / "session.json").read_text(encoding="utf-8"))
    assert session["merge"]["matched"] == 2 and session["finalized_at"]


def test_prepare_context_fields_builds_case_context(tmp_path):
    run = _run(tmp_path)
    result = prepare_analysis(run, context_fields=["CaseNumber", "Account.Name"], version="v2")

    assert result.fields == ("Case Context",)
    config = yaml.safe_load((run / "case_summary.yaml").read_text(encoding="utf-8"))
    assert config["dataset"]["path"].endswith("case_context.xlsx")
    assert config["analysis"]["version"] == "v2"
    rows = list(load_workbook(run / "case_context.xlsx")["Cases"].iter_rows(values_only=True))
    assert rows[0] == ("CaseNumber", "Case Context")
    assert rows[1] == ("1", "CaseNumber: 1\nAccount.Name: Acme")
    with pytest.raises(AnalysisRunError):
        prepare_analysis(run, context_fields=["Missing__c"])


def test_prepare_cli_and_errors(tmp_path, capsys):
    run = _run(tmp_path)
    assert fetch_main(["prepare", "--run", str(run), "--fields", "CaseNumber,Account.Name"]) == 0
    assert '"needs_batch_run": true' in capsys.readouterr().out
    assert fetch_main(["finalize", "--run", str(tmp_path)]) == 1
    assert fetch_main(["prepare", "--run", str(run), "--analysis", "nope"]) == 1